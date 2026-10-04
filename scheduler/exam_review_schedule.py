# -*- coding: utf-8 -*-
from __future__ import annotations

import argparse
import json
import logging
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl.styles import Font, PatternFill

from ai_orchestrated_optimization import (
    CpSatSolveConfig,
    LinearConstraintSpec,
    LinearExpression,
    LinearTerm,
    ObjectiveSpec,
    OptimizationProblemSpec,
    RuleSpec,
    SolverParameterSpec,
    VariableSpec,
    solve_cp_sat_problem,
)
from scheduler.data.day_rules_reader import load_teacher_positioning
from scheduler.output.excel_writer import create_excel_writer


logger = logging.getLogger(__name__)

ALLOWED_SUBJECTS = ("语文", "数学", "外语", "物理", "化学", "生物", "历史", "政治", "地理")
CORE_SUBJECTS = frozenset(("语文", "数学"))


@dataclass(frozen=True)
class ReviewBlock:
    code: str
    date_label: str
    block_label: str
    periods: str
    sort_index: int

    @property
    def column_label(self) -> str:
        return f"{self.date_label} {self.block_label}"


REVIEW_BLOCKS = (
    ReviewBlock("0521_1", "5.21", "第1-2节", "1-2", 1),
    ReviewBlock("0521_2", "5.21", "第3-4节", "3-4", 2),
    ReviewBlock("0522_1", "5.22", "第1-2节", "1-2", 3),
    ReviewBlock("0522_2", "5.22", "第3-4节", "3-4", 4),
    ReviewBlock("0522_3", "5.22", "第5-6节", "5-6", 5),
    ReviewBlock("0522_4", "5.22", "第7-8节", "7-8", 6),
)
BLOCK_BY_CODE = {block.code: block for block in REVIEW_BLOCKS}


@dataclass(frozen=True)
class ReviewAssignment:
    class_name: str
    subject: str
    teacher: str
    block_code: str


@dataclass(frozen=True)
class ReviewScheduleResult:
    classes: tuple[str, ...]
    courses_by_class: dict[str, dict[str, str]]
    assignments: tuple[ReviewAssignment, ...]
    solver_status: str
    objective_value: float | None
    validation_errors: tuple[str, ...]


@dataclass(frozen=True)
class CompiledReviewScheduleProblem:
    problem: OptimizationProblemSpec
    assignment_by_variable: dict[str, ReviewAssignment]


@dataclass(frozen=True)
class OutputPaths:
    run_dir: Path
    excel_path: Path
    long_csv_path: Path
    metadata_path: Path


def _class_sort_key(class_name: str) -> tuple[int, str]:
    match = re.match(r"^(\d+)", class_name)
    if match:
        return int(match.group(1)), class_name
    return 10_000, class_name


def _allowed_blocks_for_subject(subject: str) -> tuple[ReviewBlock, ...]:
    if subject in CORE_SUBJECTS:
        return REVIEW_BLOCKS[:2]
    return REVIEW_BLOCKS[2:]


def _sum_expression(variable_names: Iterable[str]) -> LinearExpression:
    return LinearExpression(tuple(LinearTerm(variable, 1) for variable in variable_names))


def load_review_courses(teacher_positioning_path: str | Path) -> tuple[tuple[str, ...], dict[str, dict[str, str]]]:
    classes, cls_subj_teacher = load_teacher_positioning(str(teacher_positioning_path))
    ordered_classes = tuple(sorted(classes, key=_class_sort_key))
    courses_by_class: dict[str, dict[str, str]] = {}
    errors: list[str] = []

    for class_name in ordered_classes:
        courses = {
            subject: teacher
            for subject in ALLOWED_SUBJECTS
            if (teacher := cls_subj_teacher.get((class_name, subject), ""))
        }
        courses_by_class[class_name] = courses

        missing_core = sorted(CORE_SUBJECTS - set(courses))
        if len(courses) != 6:
            errors.append(
                f"{class_name} 在允许学科中有 {len(courses)} 门：{','.join(courses) or '无'}；本任务要求正好 6 门"
            )
        if missing_core:
            errors.append(f"{class_name} 缺少 5.21 必排学科：{','.join(missing_core)}")
        if len(set(courses) - CORE_SUBJECTS) != 4:
            errors.append(f"{class_name} 的 5.22 剩余学科不是 4 门")

    if errors:
        raise ValueError("教师定位表不满足临时考前复习排课要求：\n- " + "\n- ".join(errors))

    return ordered_classes, courses_by_class


def build_review_schedule_problem(
    classes: Iterable[str],
    courses_by_class: dict[str, dict[str, str]],
) -> CompiledReviewScheduleProblem:
    ordered_classes = tuple(classes)
    subject_indexes = {subject: index for index, subject in enumerate(ALLOWED_SUBJECTS)}
    block_indexes = {block.code: index for index, block in enumerate(REVIEW_BLOCKS)}
    unknown_subjects = sorted(
        {
            subject
            for class_name in ordered_classes
            for subject in courses_by_class.get(class_name, {})
            if subject not in subject_indexes
        }
    )
    if unknown_subjects:
        raise ValueError(f"考前复习排课包含不支持学科：{','.join(unknown_subjects)}")

    variables: list[VariableSpec] = []
    assignment_by_variable: dict[str, ReviewAssignment] = {}
    by_class_subject: dict[tuple[str, str], list[str]] = defaultdict(list)
    by_class_block: dict[tuple[str, str], list[str]] = defaultdict(list)
    by_teacher_block: dict[tuple[str, str], list[str]] = defaultdict(list)
    by_subject_block: dict[tuple[str, str], list[str]] = defaultdict(list)

    for class_index, class_name in enumerate(ordered_classes):
        for subject, teacher in courses_by_class[class_name].items():
            for block in _allowed_blocks_for_subject(subject):
                var = f"x_c{class_index}_s{subject_indexes[subject]}_b{block_indexes[block.code]}"
                key = (class_name, subject, block.code)
                variables.append(VariableSpec.bool(var))
                assignment_by_variable[var] = ReviewAssignment(
                    class_name=class_name,
                    subject=subject,
                    teacher=teacher,
                    block_code=block.code,
                )
                by_class_subject[(class_name, subject)].append(var)
                by_class_block[(class_name, block.code)].append(var)
                by_teacher_block[(teacher, block.code)].append(var)
                by_subject_block[(subject, block.code)].append(var)

    constraints: list[LinearConstraintSpec] = []
    for class_index, class_name in enumerate(ordered_classes):
        for subject in courses_by_class[class_name]:
            constraints.append(
                LinearConstraintSpec(
                    name=f"class_subject_once_c{class_index}_s{subject_indexes[subject]}",
                    expression=_sum_expression(by_class_subject[(class_name, subject)]),
                    sense="==",
                    rhs=1,
                    rule_id="review.class_subject_once",
                )
            )
        for block in REVIEW_BLOCKS:
            constraints.append(
                LinearConstraintSpec(
                    name=f"class_block_once_c{class_index}_b{block_indexes[block.code]}",
                    expression=_sum_expression(by_class_block[(class_name, block.code)]),
                    sense="==",
                    rhs=1,
                    rule_id="review.class_block_once",
                )
            )

    for teacher_block_index, ((_teacher, block_code), vars_for_teacher) in enumerate(
        sorted(by_teacher_block.items(), key=lambda item: (item[0][0], item[0][1]))
    ):
        constraints.append(
            LinearConstraintSpec(
                name=f"teacher_block_capacity_t{teacher_block_index}_b{block_indexes[block_code]}",
                expression=_sum_expression(vars_for_teacher),
                sense="<=",
                rhs=1,
                rule_id="review.teacher_block_capacity",
            )
        )

    subject_sync_terms: list[str] = []
    for subject in ALLOWED_SUBJECTS:
        for block in _allowed_blocks_for_subject(subject):
            vars_for_subject = by_subject_block.get((subject, block.code), [])
            if not vars_for_subject:
                continue
            active = f"active_s{subject_indexes[subject]}_b{block_indexes[block.code]}"
            variables.append(VariableSpec.bool(active))
            for var_index, var in enumerate(vars_for_subject):
                constraints.append(
                    LinearConstraintSpec(
                        name=f"subject_sync_var_requires_active_s{subject_indexes[subject]}_b{block_indexes[block.code]}_v{var_index}",
                        expression=LinearExpression((LinearTerm(var, 1), LinearTerm(active, -1))),
                        sense="<=",
                        rhs=0,
                        rule_id="review.subject_sync",
                    )
                )
            constraints.append(
                LinearConstraintSpec(
                    name=f"subject_sync_active_requires_var_s{subject_indexes[subject]}_b{block_indexes[block.code]}",
                    expression=LinearExpression(
                        (LinearTerm(active, 1),)
                        + tuple(LinearTerm(var, -1) for var in vars_for_subject)
                    ),
                    sense="<=",
                    rhs=0,
                    rule_id="review.subject_sync",
                )
            )
            subject_sync_terms.append(active)

    return CompiledReviewScheduleProblem(
        problem=OptimizationProblemSpec(
            problem_id="exam_review_schedule",
            variables=tuple(variables),
            rules=(
                RuleSpec(
                    "review.class_subject_once",
                    "each class-subject review requirement is assigned to one allowed block",
                    "hard",
                    10,
                    "scheduler.exam_review",
                ),
                RuleSpec(
                    "review.class_block_once",
                    "each class has exactly one review assignment in every review block",
                    "hard",
                    20,
                    "scheduler.exam_review",
                ),
                RuleSpec(
                    "review.teacher_block_capacity",
                    "a teacher can teach at most one class in the same review block",
                    "hard",
                    30,
                    "scheduler.exam_review",
                ),
                RuleSpec(
                    "review.subject_sync",
                    "minimize the number of active subject-block pairs to synchronize same-subject reviews",
                    "soft",
                    100,
                    "scheduler.exam_review",
                ),
            ),
            constraints=tuple(constraints),
            objective=ObjectiveSpec(
                "review.subject_sync",
                "minimize",
                _sum_expression(subject_sync_terms),
            ),
            solve_config=CpSatSolveConfig(parameters=(SolverParameterSpec("num_search_workers", 1),)),
        ),
        assignment_by_variable=assignment_by_variable,
    )


def solve_review_schedule(
    classes: Iterable[str],
    courses_by_class: dict[str, dict[str, str]],
    *,
    time_limit_seconds: float = 30.0,
) -> ReviewScheduleResult:
    ordered_classes = tuple(classes)
    compiled = build_review_schedule_problem(ordered_classes, courses_by_class)
    solution = solve_cp_sat_problem(compiled.problem, time_limit_seconds=time_limit_seconds)
    if solution.status_name not in {"OPTIMAL", "FEASIBLE"}:
        return ReviewScheduleResult(
            classes=ordered_classes,
            courses_by_class=courses_by_class,
            assignments=(),
            solver_status=solution.status_name,
            objective_value=None,
            validation_errors=(f"CP-SAT 未找到可行解：{solution.status_name}",),
        )

    assignments: list[ReviewAssignment] = []
    for variable_name, assignment in compiled.assignment_by_variable.items():
        if solution.values.get(variable_name) == 1:
            assignments.append(assignment)
    assignments.sort(key=lambda item: (_class_sort_key(item.class_name), BLOCK_BY_CODE[item.block_code].sort_index))
    validation_errors = tuple(validate_review_schedule(ordered_classes, courses_by_class, assignments))
    return ReviewScheduleResult(
        classes=ordered_classes,
        courses_by_class=courses_by_class,
        assignments=tuple(assignments),
        solver_status=solution.status_name,
        objective_value=solution.objective_value,
        validation_errors=validation_errors,
    )


def validate_review_schedule(
    classes: Iterable[str],
    courses_by_class: dict[str, dict[str, str]],
    assignments: Iterable[ReviewAssignment],
) -> list[str]:
    errors: list[str] = []
    by_class_block: dict[tuple[str, str], list[ReviewAssignment]] = defaultdict(list)
    by_class_subject: dict[tuple[str, str], list[ReviewAssignment]] = defaultdict(list)
    by_teacher_block: dict[tuple[str, str], list[ReviewAssignment]] = defaultdict(list)
    by_class_date: dict[tuple[str, str], list[ReviewAssignment]] = defaultdict(list)

    for assignment in assignments:
        block = BLOCK_BY_CODE[assignment.block_code]
        by_class_block[(assignment.class_name, assignment.block_code)].append(assignment)
        by_class_subject[(assignment.class_name, assignment.subject)].append(assignment)
        by_teacher_block[(assignment.teacher, assignment.block_code)].append(assignment)
        by_class_date[(assignment.class_name, block.date_label)].append(assignment)

    for class_name in classes:
        subjects = set(courses_by_class[class_name])
        for block in REVIEW_BLOCKS:
            current = by_class_block.get((class_name, block.code), [])
            if len(current) != 1:
                errors.append(f"{class_name} {block.column_label} 排课数为 {len(current)}，应为 1")
        for subject in subjects:
            current = by_class_subject.get((class_name, subject), [])
            if len(current) != 1:
                errors.append(f"{class_name} {subject} 排课数为 {len(current)}，应为 1")

        day_0521_subjects = {item.subject for item in by_class_date.get((class_name, "5.21"), [])}
        if day_0521_subjects != CORE_SUBJECTS:
            errors.append(f"{class_name} 5.21 学科为 {','.join(sorted(day_0521_subjects))}，应为 语文,数学")

        day_0522_subjects = {item.subject for item in by_class_date.get((class_name, "5.22"), [])}
        expected_remaining = subjects - CORE_SUBJECTS
        if day_0522_subjects != expected_remaining:
            errors.append(
                f"{class_name} 5.22 学科为 {','.join(sorted(day_0522_subjects))}，"
                f"应为 {','.join(sorted(expected_remaining))}"
            )

    for (teacher, block_code), current in by_teacher_block.items():
        if len(current) > 1:
            block = BLOCK_BY_CODE[block_code]
            classes_text = ",".join(item.class_name for item in current)
            errors.append(f"{teacher} 在 {block.column_label} 同时授课多个班：{classes_text}")

    return errors


def build_long_table(result: ReviewScheduleResult) -> pd.DataFrame:
    rows = []
    for assignment in result.assignments:
        block = BLOCK_BY_CODE[assignment.block_code]
        rows.append(
            {
                "日期": block.date_label,
                "双节块": block.block_label,
                "连续节次": block.periods,
                "班级": assignment.class_name,
                "学科": assignment.subject,
                "教师": assignment.teacher,
            }
        )
    return pd.DataFrame(rows, columns=["日期", "双节块", "连续节次", "班级", "学科", "教师"])


def build_class_wide_table(result: ReviewScheduleResult) -> pd.DataFrame:
    rows = []
    by_class_block = {(item.class_name, item.block_code): item for item in result.assignments}
    for class_name in result.classes:
        row = {"班级": class_name}
        for block in REVIEW_BLOCKS:
            assignment = by_class_block.get((class_name, block.code))
            row[block.column_label] = "" if assignment is None else f"{assignment.subject}（{assignment.teacher}）"
        rows.append(row)
    return pd.DataFrame(rows, columns=["班级"] + [block.column_label for block in REVIEW_BLOCKS])


def build_teacher_wide_table(result: ReviewScheduleResult) -> pd.DataFrame:
    by_teacher: dict[str, dict[str, str]] = defaultdict(dict)
    for assignment in result.assignments:
        value = f"{assignment.class_name}-{assignment.subject}"
        by_teacher[assignment.teacher][assignment.block_code] = value

    rows = []
    for teacher in sorted(by_teacher):
        row = {"教师": teacher}
        for block in REVIEW_BLOCKS:
            row[block.column_label] = by_teacher[teacher].get(block.code, "")
        rows.append(row)
    return pd.DataFrame(rows, columns=["教师"] + [block.column_label for block in REVIEW_BLOCKS])


def build_subject_sync_table(result: ReviewScheduleResult) -> pd.DataFrame:
    rows = []
    grouped: dict[tuple[str, str], list[ReviewAssignment]] = defaultdict(list)
    for assignment in result.assignments:
        grouped[(assignment.subject, assignment.block_code)].append(assignment)

    for subject in ALLOWED_SUBJECTS:
        for block in _allowed_blocks_for_subject(subject):
            items = grouped.get((subject, block.code), [])
            if not items:
                continue
            rows.append(
                {
                    "学科": subject,
                    "日期": block.date_label,
                    "双节块": block.block_label,
                    "班级数": len(items),
                    "班级": "、".join(item.class_name for item in sorted(items, key=lambda x: _class_sort_key(x.class_name))),
                    "教师": "、".join(sorted({item.teacher for item in items})),
                }
            )
    return pd.DataFrame(rows, columns=["学科", "日期", "双节块", "班级数", "班级", "教师"])


def _style_workbook(excel_path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(excel_path)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(value))
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 10), 42)
    workbook.save(excel_path)


def write_review_outputs(result: ReviewScheduleResult, output_dir: str | Path) -> OutputPaths:
    run_dir = Path(output_dir) / f"run_{datetime.now():%Y%m%d_%H%M%S}"
    run_dir.mkdir(parents=True, exist_ok=False)

    excel_path = run_dir / "考前复习课表.xlsx"
    long_csv_path = run_dir / "考前复习课表_长表.csv"
    metadata_path = run_dir / "metadata.json"

    class_wide = build_class_wide_table(result)
    teacher_wide = build_teacher_wide_table(result)
    long_table = build_long_table(result)
    subject_sync = build_subject_sync_table(result)
    validation = pd.DataFrame(
        [{"校验项": "总体", "结果": "OK" if not result.validation_errors else "ERROR"}]
        + [{"校验项": "明细", "结果": error} for error in result.validation_errors]
    )
    metadata = {
        "solver_status": result.solver_status,
        "objective_value": result.objective_value,
        "class_count": len(result.classes),
        "assignment_count": len(result.assignments),
        "validation_ok": not result.validation_errors,
        "rules": [
            "5.21 only Chinese and Math, one two-period block each",
            "5.22 schedules the remaining four subjects",
            "each assignment is a continuous two-period block",
            "teacher cannot teach multiple classes in the same block",
            "class cannot have multiple teachers in the same block",
            "same-subject classes are synchronized as a soft objective",
        ],
    }

    with create_excel_writer(excel_path) as writer:
        class_wide.to_excel(writer, sheet_name="班级课表", index=False)
        teacher_wide.to_excel(writer, sheet_name="教师课表", index=False)
        long_table.to_excel(writer, sheet_name="长表", index=False)
        subject_sync.to_excel(writer, sheet_name="同学科同步", index=False)
        validation.to_excel(writer, sheet_name="校验", index=False)

    _style_workbook(excel_path)
    long_table.to_csv(long_csv_path, index=False, encoding="utf-8-sig")
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    return OutputPaths(run_dir=run_dir, excel_path=excel_path, long_csv_path=long_csv_path, metadata_path=metadata_path)


def run_exam_review_schedule(
    teacher_positioning_path: str | Path = "scheduler/教师定位表.xlsx",
    output_dir: str | Path = "outputs/exam_review",
    *,
    time_limit_seconds: float = 30.0,
) -> tuple[ReviewScheduleResult, OutputPaths]:
    classes, courses_by_class = load_review_courses(teacher_positioning_path)
    result = solve_review_schedule(classes, courses_by_class, time_limit_seconds=time_limit_seconds)
    paths = write_review_outputs(result, output_dir)
    return result, paths


def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description="生成 5.21/5.22 临时考前复习课表")
    parser.add_argument("--teacher-positioning", default="scheduler/教师定位表.xlsx", help="教师定位表 Excel 路径")
    parser.add_argument("--output-dir", default="outputs/exam_review", help="输出目录")
    parser.add_argument("--time-limit", type=float, default=30.0, help="求解时间上限，单位秒")
    args = parser.parse_args(argv)

    result, paths = run_exam_review_schedule(
        teacher_positioning_path=args.teacher_positioning,
        output_dir=args.output_dir,
        time_limit_seconds=args.time_limit,
    )
    logger.info("solver_status=%s", result.solver_status)
    logger.info("validation_ok=%s", not result.validation_errors)
    logger.info("class_count=%s", len(result.classes))
    logger.info("assignment_count=%s", len(result.assignments))
    logger.info("excel=%s", paths.excel_path)
    logger.info("long_csv=%s", paths.long_csv_path)
    if result.validation_errors:
        for error in result.validation_errors:
            logger.error("validation_error=%s", error)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
