# -*- coding: utf-8 -*-
"""User-facing timetable exports for the formal publish workflow."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, Mapping
from xml.sax.saxutils import escape

from scheduler.app.config_service import PROJECT_ROOT
from scheduler.app.result_preview import build_result_preview
from scheduler.output.constants import DAYS


def resolve_result_excel(
    status: Mapping[str, Any],
    *,
    project_root: Path | None = PROJECT_ROOT,
) -> Path:
    preview = build_result_preview(dict(status), project_root=project_root)
    for workbook in preview.get("workbooks", []):
        if not isinstance(workbook, Mapping):
            continue
        path = Path(str(workbook.get("path") or ""))
        if path.exists() and path.is_file() and path.suffix.lower() in {".xlsx", ".xlsm"}:
            return path.resolve()
    files = status.get("files") if isinstance(status.get("files"), list) else []
    for item in files:
        if not isinstance(item, Mapping):
            continue
        path = Path(str(item.get("path") or ""))
        label = f"{item.get('label') or ''} {item.get('name') or ''} {path.name}"
        if (
            path.exists()
            and path.is_file()
            and path.suffix.lower() in {".xlsx", ".xlsm"}
            and ("课表" in label or "schedule" in label.lower() or "timetable" in label.lower())
        ):
            return path.resolve()
    raise ValueError("当前候选没有可下载的 Excel 课表，请先完成求解并生成结果")


def build_result_pdf(
    status: Mapping[str, Any],
    *,
    project_root: Path | None = PROJECT_ROOT,
    view: str = "all",
) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:  # pragma: no cover - deployment dependency guard
        raise RuntimeError("PDF 导出组件尚未安装") from exc

    preview = build_result_preview(dict(status), project_root=project_root)
    class_views, teacher_views = _selected_views(preview, view)
    if not class_views and not teacher_views:
        raise ValueError("当前候选没有可导出的课表，请先完成求解并生成结果")

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=15 * mm,
        bottomMargin=14 * mm,
        title="课有序课表",
        author="课有序 CourseOrder AI",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CourseOrderTitle",
        parent=styles["Title"],
        fontName="STSong-Light",
        fontSize=18,
        leading=24,
        textColor=colors.HexColor("#0B1739"),
        alignment=TA_CENTER,
        spaceAfter=5 * mm,
    )
    section_style = ParagraphStyle(
        "CourseOrderSection",
        parent=styles["Heading2"],
        fontName="STSong-Light",
        fontSize=13,
        leading=18,
        textColor=colors.HexColor("#163B78"),
        spaceAfter=3 * mm,
    )
    cell_style = ParagraphStyle(
        "CourseOrderCell",
        parent=styles["BodyText"],
        fontName="STSong-Light",
        fontSize=6.4,
        leading=8,
        textColor=colors.HexColor("#213552"),
        alignment=TA_CENTER,
    )
    header_style = ParagraphStyle(
        "CourseOrderHeader",
        parent=cell_style,
        fontSize=7,
        leading=9,
        textColor=colors.white,
    )
    story: list[Any] = [
        Paragraph("课有序 · 正式课表", title_style),
        Paragraph(
            escape(str((preview.get("summary") or {}).get("message") or "由当前用户确认并导出的课表。")),
            ParagraphStyle(
                "CourseOrderSummary",
                parent=cell_style,
                fontSize=8,
                leading=11,
                textColor=colors.HexColor("#52617C"),
            ),
        ),
        Spacer(1, 5 * mm),
    ]
    view_groups = [("班级课表", class_views), ("教师课表", teacher_views)]
    first_view = True
    for group_label, views in view_groups:
        if not views:
            continue
        if not first_view:
            story.append(PageBreak())
        story.append(Paragraph(group_label, title_style))
        for view_index, view in enumerate(views[:60]):
            if view_index:
                story.append(PageBreak())
            name = str(view.get("name") or f"{group_label}{view_index + 1}")
            story.append(Paragraph(escape(name), section_style))
            table_rows: list[list[Any]] = [
                [Paragraph("节次", header_style), *[Paragraph(escape(day), header_style) for day in DAYS]]
            ]
            for row in view.get("rows", [])[:30]:
                if not isinstance(row, Mapping):
                    continue
                values = row.get("values") if isinstance(row.get("values"), Mapping) else {}
                table_rows.append([
                    Paragraph(escape(str(row.get("slot") or "")), cell_style),
                    *[Paragraph(escape(str(values.get(day) or "")), cell_style) for day in DAYS],
                ])
            table = Table(
                table_rows,
                colWidths=[25 * mm, *([34.5 * mm] * len(DAYS))],
                repeatRows=1,
                hAlign="CENTER",
            )
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1464F6")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CFD9E8")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F9FE")]),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(table)
        first_view = False

    def draw_page_footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("STSong-Light", 7)
        canvas.setFillColor(colors.HexColor("#78869B"))
        canvas.drawString(12 * mm, 7 * mm, "课有序 CourseOrder AI · 用户确认导出")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 7 * mm, f"第 {doc.page} 页")
        canvas.restoreState()

    document.build(story, onFirstPage=draw_page_footer, onLaterPages=draw_page_footer)
    return buffer.getvalue()


def build_result_xlsx(
    status: Mapping[str, Any],
    *,
    project_root: Path | None = PROJECT_ROOT,
    view: str = "all",
) -> bytes:
    """Build focused class/teacher workbooks from the original result preview."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:  # pragma: no cover - deployment dependency guard
        raise RuntimeError("Excel 导出组件尚未安装") from exc

    preview = build_result_preview(dict(status), project_root=project_root)
    class_views, teacher_views = _selected_views(preview, view)
    groups = [("班级", class_views), ("教师", teacher_views)]
    if not class_views and not teacher_views:
        raise ValueError("当前候选没有可导出的课表，请先完成求解并生成结果")

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()
    header_fill = PatternFill("solid", fgColor="1464F6")
    header_font = Font(color="FFFFFF", bold=True)
    title_font = Font(color="0B1739", bold=True, size=16)
    thin = Side(style="thin", color="D6DFEC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for group_label, views in groups:
        for index, schedule in enumerate(views, start=1):
            raw_name = str(schedule.get("name") or f"{group_label}课表{index}")
            sheet_name = _unique_sheet_name(raw_name, used_names)
            sheet = workbook.create_sheet(sheet_name)
            sheet.sheet_view.showGridLines = False
            sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(DAYS) + 1)
            title_cell = sheet.cell(1, 1, f"{raw_name} · {group_label}课表")
            title_cell.font = title_font
            title_cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.row_dimensions[1].height = 30
            headers = ["节次", *DAYS]
            for column, label in enumerate(headers, start=1):
                cell = sheet.cell(2, column, label)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border
            for row_index, row in enumerate(schedule.get("rows", []), start=3):
                if not isinstance(row, Mapping):
                    continue
                values = row.get("values") if isinstance(row.get("values"), Mapping) else {}
                row_values = [str(row.get("slot") or ""), *[str(values.get(day) or "") for day in DAYS]]
                for column, value in enumerate(row_values, start=1):
                    cell = sheet.cell(row_index, column, value)
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = border
                    if row_index % 2:
                        cell.fill = PatternFill("solid", fgColor="F6F9FE")
                sheet.row_dimensions[row_index].height = 34
            sheet.freeze_panes = "B3"
            sheet.column_dimensions["A"].width = 15
            for column in range(2, len(DAYS) + 2):
                sheet.column_dimensions[sheet.cell(2, column).column_letter].width = 23
            sheet.print_title_rows = "1:2"
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.fitToWidth = 1
            sheet.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _selected_views(preview: Mapping[str, Any], view: str) -> tuple[list[Mapping[str, Any]], list[Mapping[str, Any]]]:
    normalized = str(view or "all").strip().lower()
    class_views = [item for item in preview.get("class_views", []) if isinstance(item, Mapping)]
    teacher_views = [item for item in preview.get("teacher_views", []) if isinstance(item, Mapping)]
    if normalized == "class":
        teacher_views = []
    elif normalized == "teacher":
        class_views = []
    elif normalized not in {"all", "academic"}:
        raise ValueError("课表类型仅支持班级、教师或教务总表")
    return class_views, teacher_views


def _unique_sheet_name(value: str, used: set[str]) -> str:
    cleaned = "".join("_" if char in "[]:*?/\\" else char for char in str(value)).strip() or "课表"
    base = cleaned[:31]
    name = base
    suffix = 2
    while name in used:
        tail = f"_{suffix}"
        name = f"{base[:31-len(tail)]}{tail}"
        suffix += 1
    used.add(name)
    return name


def result_export_name(export_format: str, view: str = "all") -> str:
    label = {"class": "班级课表", "teacher": "教师课表", "academic": "教务总表", "all": "正式课表"}.get(str(view).lower(), "正式课表")
    suffix = "pdf" if str(export_format).lower() == "pdf" else "xlsx"
    return f"课有序_{label}.{suffix}"
