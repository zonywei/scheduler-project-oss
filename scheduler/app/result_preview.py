# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from collections import Counter, defaultdict

from scheduler.app.config_service import PROJECT_ROOT
from scheduler.app.rule_explain import build_rule_explain, empty_rule_explain
from scheduler.output.constants import DAYS, DAY_COLS


MAX_WORKBOOKS = 1
MAX_SHEETS = 24
MAX_TABLE_ROWS = 32
MAX_TABLE_COLS = 18
MAX_VIEWS = 60


def build_result_preview(status: dict[str, Any], project_root: Path | None = PROJECT_ROOT) -> dict[str, Any]:
    """Build a lightweight, browser-safe preview from latest schedule workbooks."""
    status = status if isinstance(status, dict) else {}
    root = project_root.resolve() if project_root is not None else None
    status, preview_source = _select_preview_status(status, root)
    files = status.get("files") if isinstance(status.get("files"), list) else []
    active_files = []
    for key, label in (
        ("local_timetable_adjustment", "active_manual_adjustment"),
        ("local_timetable_repair", "active_local_repair"),
    ):
        local = status.get(key) if isinstance(status.get(key), dict) else {}
        active = str(local.get("active_file") or "").strip()
        if active:
            active_files.append({"path": active, "label": label})
    availability = status.get("result_availability") if isinstance(status.get("result_availability"), dict) else {}
    primary = availability.get("primary_schedule_files") if isinstance(availability.get("primary_schedule_files"), list) else []
    candidates = [_safe_file_item(item, root) for item in [*active_files, *primary, *files] if isinstance(item, dict)]
    schedule_files = sorted(
        [item for item in candidates if item and _is_schedule_workbook(item)],
        key=_schedule_preview_rank,
    )

    if not schedule_files:
        return _with_preview_source(
            _empty_preview("未发现可在线预览的课表 Excel。完成可行求解并导出课表后，这里会出现班级和教师视图。"),
            preview_source,
        )

    workbooks: list[dict[str, Any]] = []
    class_views: list[dict[str, Any]] = []
    teacher_views: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []

    for item in schedule_files[:MAX_WORKBOOKS]:
        try:
            parsed = _parse_workbook(item)
        except Exception as exc:  # pragma: no cover - exact parser exceptions vary by workbook corruption
            errors.append({"label": str(item.get("label") or item.get("path") or ""), "message": str(exc)})
            continue
        workbooks.append(parsed["workbook"])
        class_views.extend(parsed["class_views"])
        teacher_views.extend(parsed["teacher_views"])

    if class_views and not teacher_views:
        teacher_views = _derive_teacher_views(class_views)

    class_views = _dedupe_views(class_views)[:MAX_VIEWS]
    teacher_views = _dedupe_views(teacher_views)[:MAX_VIEWS]
    sheet_count = sum(len(workbook.get("sheets") or []) for workbook in workbooks)
    quality = _build_quality_summary(class_views)
    statistics = _build_schedule_statistics(status, class_views, teacher_views, root)
    versions = _build_schedule_versions(status, schedule_files)

    if class_views or teacher_views:
        preview_status = "ready"
        message = f"已解析 {len(class_views)} 个班级视图、{len(teacher_views)} 个教师视图，可在线抽查课表。"
    elif workbooks:
        preview_status = "generic"
        message = "已读取课表工作簿，但未识别出标准班级/教师网格；可先查看原始表预览。"
    else:
        preview_status = "error" if errors else "empty"
        message = "课表文件读取失败，建议下载结果包后检查 Excel 是否损坏。"

    return _with_preview_source({
        "summary": {
            "status": preview_status,
            "message": message,
            "schedule_files": len(schedule_files),
            "workbooks": len(workbooks),
            "sheets": sheet_count,
            "class_views": len(class_views),
            "teacher_views": len(teacher_views),
            "errors": len(errors),
        },
        "quality": quality,
        "statistics": statistics,
        "versions": versions,
        "workbooks": workbooks,
        "class_views": class_views,
        "teacher_views": teacher_views,
        "errors": errors,
    }, preview_source)


def _empty_preview(message: str) -> dict[str, Any]:
    return {
        "summary": {
            "status": "empty",
            "message": message,
            "schedule_files": 0,
            "workbooks": 0,
            "sheets": 0,
            "class_views": 0,
            "teacher_views": 0,
            "errors": 0,
        },
        "quality": _empty_quality(),
        "statistics": _empty_statistics(),
        "versions": _empty_versions(),
        "workbooks": [],
        "class_views": [],
        "teacher_views": [],
        "errors": [],
    }


def _select_preview_status(status: dict[str, Any], root: Path | None) -> tuple[dict[str, Any], dict[str, str]]:
    current_run_id = _status_run_id(status)
    source = {
        "source_kind": "latest_run",
        "source_run_id": current_run_id,
        "latest_run_id": current_run_id,
        "source_message": "预览使用最新求解批次。",
    }
    recommendation = status.get("recommended_formal_candidate")
    if not isinstance(recommendation, dict):
        return status, source
    summary = recommendation.get("summary") if isinstance(recommendation.get("summary"), dict) else {}
    selected_run_id = str(summary.get("selected_run_id") or "").strip()
    if not selected_run_id or selected_run_id == current_run_id:
        return status, source
    candidate = recommendation.get("candidate") if isinstance(recommendation.get("candidate"), dict) else {}
    candidate_status_path = _candidate_status_path(candidate, selected_run_id=selected_run_id, root=root)
    if candidate_status_path is None:
        source.update(
            {
                "source_kind": "recommended_candidate_missing",
                "source_run_id": selected_run_id,
                "source_message": "已推荐历史正式候选，但未定位到候选批次 status.json；预览暂用最新批次。",
            }
        )
        return status, source
    try:
        candidate_status = json.loads(candidate_status_path.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError, UnicodeDecodeError):
        source.update(
            {
                "source_kind": "recommended_candidate_unreadable",
                "source_run_id": selected_run_id,
                "source_message": "已推荐历史正式候选，但候选批次 status.json 无法读取；预览暂用最新批次。",
            }
        )
        return status, source
    if not isinstance(candidate_status, dict):
        return status, source
    candidate_status.setdefault("status_file", str(candidate_status_path))
    source.update(
        {
            "source_kind": "recommended_formal_candidate",
            "source_run_id": selected_run_id,
            "source_message": "最新正式批次退化；预览已切换到推荐的历史最好正式候选。",
        }
    )
    return candidate_status, source


def _with_preview_source(preview: dict[str, Any], source: dict[str, str]) -> dict[str, Any]:
    summary = preview.get("summary") if isinstance(preview.get("summary"), dict) else {}
    summary.update(source)
    preview["summary"] = summary
    return preview


def _status_run_id(status: dict[str, Any]) -> str:
    run_id = str(status.get("run_id") or "").strip()
    if run_id:
        return run_id
    run_dir = str(status.get("run_dir") or "").strip()
    return Path(run_dir).name if run_dir else ""


def _candidate_status_path(candidate: dict[str, Any], *, selected_run_id: str, root: Path | None) -> Path | None:
    raw_status = str(candidate.get("status_file") or "").strip()
    if raw_status:
        path = Path(raw_status)
        if path.exists() and path.is_file():
            return path
    raw_dir = str(candidate.get("run_dir") or "").strip()
    if raw_dir:
        path = Path(raw_dir) / "status.json"
        if path.exists() and path.is_file():
            return path
    if root is not None and selected_run_id:
        path = root / "outputs" / "web_runs" / selected_run_id / "status.json"
        if path.exists() and path.is_file():
            return path
    return None


def _safe_file_item(item: dict[str, Any], root: Path | None) -> dict[str, Any] | None:
    path = Path(str(item.get("path") or ""))
    if not path.exists() or not path.is_file():
        return None
    target = path.resolve()
    if root is not None and root not in target.parents and target != root:
        return None
    return {
        "label": str(item.get("label") or target.name),
        "path": str(target),
        "size": item.get("size"),
        "modified_at": str(item.get("modified_at") or ""),
    }


def _is_schedule_workbook(item: dict[str, Any]) -> bool:
    source = f"{item.get('label') or ''} {item.get('path') or ''}".replace("\\", "/")
    lower = source.lower()
    filename = Path(str(item.get("path") or "")).name.lower()
    if not filename.endswith((".xlsx", ".xlsm")):
        return False
    if any(token in lower for token in ("诊断", "审计", "diagnostic", "audit", "penalty", "conflict", "violation")):
        return False
    return (
        "课表" in source
        or "最优解" in source
        or any(token in filename for token in ("schedule", "timetable", "class", "teacher"))
    )


def _schedule_preview_rank(item: dict[str, Any]) -> tuple[int, int, str]:
    source = f"{item.get('label') or ''} {item.get('path') or ''}".replace("\\", "/")
    lower = source.lower()
    rank = 50
    if "课表微调" in source or "微调" in source:
        rank -= 60
    if "局部调课" in source or "调整后" in source:
        rank -= 50
    if "/best/" in lower or "最终" in source or "最优解" in source:
        rank -= 30
    if "正式版" in source:
        rank -= 20
    if "显示节次调整版" in source:
        rank += 10
    if "top50" in lower:
        rank += 20
    return (rank, int(item.get("size") or 0), source)


def _parse_workbook(item: dict[str, Any]) -> dict[str, Any]:
    path = Path(str(item["path"]))
    wb = load_workbook(path, read_only=False, data_only=True)
    workbook = {
        "label": item.get("label") or path.name,
        "path": str(path),
        "sheets": [],
    }
    class_views: list[dict[str, Any]] = []
    teacher_views: list[dict[str, Any]] = []

    try:
        for ws in wb.worksheets[:MAX_SHEETS]:
            workbook["sheets"].append(_sheet_preview(ws, workbook["label"]))
            sheet_class_views = _extract_class_views(ws, workbook["label"])
            sheet_teacher_views = _extract_teacher_views(ws, workbook["label"])
            class_views.extend(sheet_class_views)
            teacher_views.extend(sheet_teacher_views)
    finally:
        wb.close()

    return {"workbook": workbook, "class_views": class_views, "teacher_views": teacher_views}


def _sheet_preview(ws: Any, workbook_label: str) -> dict[str, Any]:
    rows: list[list[str]] = []
    max_row = min(int(ws.max_row or 0), MAX_TABLE_ROWS)
    max_col = min(int(ws.max_column or 0), MAX_TABLE_COLS)
    if max_row <= 0 or max_col <= 0:
        return {"name": ws.title, "source": workbook_label, "rows": [], "row_count": 0, "col_count": 0}
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True):
        values = [_cell_text(value, max_len=80) for value in row]
        if any(values):
            rows.append(values)
    return {
        "name": ws.title,
        "source": workbook_label,
        "rows": rows,
        "row_count": int(ws.max_row or 0),
        "col_count": int(ws.max_column or 0),
    }


def _extract_class_views(ws: Any, workbook_label: str) -> list[dict[str, Any]]:
    views: list[dict[str, Any]] = []
    title = str(ws.title or "")
    if "班级课表" in title or "汇总" in title:
        views.extend(_extract_block_views(ws, workbook_label, kind="class"))
    elif "老师课表" not in title and "教师课表" not in title:
        views.extend(_extract_simple_grid_views(ws, workbook_label, kind="class"))
        views.extend(_extract_wide_class_views(ws, workbook_label))
    return views


def _extract_teacher_views(ws: Any, workbook_label: str) -> list[dict[str, Any]]:
    title = str(ws.title or "")
    if "老师课表" in title or "教师课表" in title or "班主任课表" in title:
        return _extract_block_views(ws, workbook_label, kind="teacher")
    return []


def _extract_block_views(ws: Any, workbook_label: str, *, kind: str) -> list[dict[str, Any]]:
    views: list[dict[str, Any]] = []
    max_row = int(ws.max_row or 0)
    for top in range(1, max_row + 1):
        title = _cell_text(ws.cell(top, 2).value) or _cell_text(ws.cell(top, 1).value)
        if "课表" not in title or not _has_day_header(ws, top + 2):
            continue
        rows: list[dict[str, Any]] = []
        row_idx = top + 4
        end = min(max_row, top + 42)
        while row_idx <= end:
            slot = _cell_text(ws.cell(row_idx, 2).value) or _cell_text(ws.cell(row_idx, 1).value)
            if "课表" in slot and row_idx > top + 4:
                break
            values = {day: _cell_text(ws.cell(row_idx, cols[0]).value) for day, cols in DAY_COLS.items()}
            if slot or any(values.values()):
                rows.append({"slot": slot or f"第{len(rows) + 1}行", "values": values})
            row_idx += 2
        if rows:
            views.append(_view(kind, _clean_view_name(title), workbook_label, ws.title, rows))
    return views


def _has_day_header(ws: Any, row_idx: int) -> bool:
    if row_idx <= 0 or row_idx > int(ws.max_row or 0):
        return False
    hits = 0
    for day, cols in DAY_COLS.items():
        text = _normalize_day(_cell_text(ws.cell(row_idx, cols[0]).value))
        if text == day:
            hits += 1
    return hits >= 3


def _extract_simple_grid_views(ws: Any, workbook_label: str, *, kind: str) -> list[dict[str, Any]]:
    header = _find_day_header_row(ws)
    if not header:
        return []
    header_row, day_cols = header
    if _cell_text(ws.cell(header_row, 1).value) == "班级":
        return []
    rows: list[dict[str, Any]] = []
    for row_idx in range(header_row + 1, min(int(ws.max_row or 0), header_row + MAX_TABLE_ROWS) + 1):
        slot = _cell_text(ws.cell(row_idx, 1).value)
        values = {day: _cell_text(ws.cell(row_idx, col).value) for day, col in day_cols}
        if slot or any(values.values()):
            rows.append({"slot": slot or f"第{len(rows) + 1}行", "values": values})
    return [_view(kind, ws.title, workbook_label, ws.title, rows)] if rows else []


def _extract_wide_class_views(ws: Any, workbook_label: str) -> list[dict[str, Any]]:
    max_col = int(ws.max_column or 0)
    if _cell_text(ws.cell(1, 1).value) != "班级" or max_col <= 1:
        return []
    columns: list[tuple[int, str, str]] = []
    for col in range(2, max_col + 1):
        parsed = _parse_day_slot_header(_cell_text(ws.cell(1, col).value))
        if parsed:
            day, slot = parsed
            columns.append((col, day, slot))
    if len(columns) < 2:
        return []

    views: list[dict[str, Any]] = []
    for row_idx in range(2, min(int(ws.max_row or 0), MAX_VIEWS + 1) + 1):
        class_name = _cell_text(ws.cell(row_idx, 1).value)
        if not class_name:
            continue
        by_slot: dict[str, dict[str, str]] = {}
        for col, day, slot in columns:
            by_slot.setdefault(slot, {d: "" for d in DAYS})[day] = _cell_text(ws.cell(row_idx, col).value)
        rows = [{"slot": slot, "values": values} for slot, values in by_slot.items()]
        views.append(_view("class", class_name, workbook_label, ws.title, rows))
    return views


def _find_day_header_row(ws: Any) -> tuple[int, list[tuple[str, int]]] | None:
    for row_idx in range(1, min(int(ws.max_row or 0), 6) + 1):
        day_cols: list[tuple[str, int]] = []
        for col in range(1, min(int(ws.max_column or 0), MAX_TABLE_COLS) + 1):
            day = _normalize_day(_cell_text(ws.cell(row_idx, col).value))
            if day in DAYS:
                day_cols.append((day, col))
        if len(day_cols) >= 3:
            return row_idx, day_cols
    return None


def _parse_day_slot_header(text: str) -> tuple[str, str] | None:
    if not text:
        return None
    normalized = text.replace("周", "星期").replace("_", "-")
    match = re.match(r"^(星期[一二三四五六日])(?:[-.](.*))?$", normalized)
    if not match:
        return None
    day = match.group(1)
    raw_slot = (match.group(2) or "").strip()
    if raw_slot == "1":
        slot = "晚自习2"
    elif raw_slot:
        slot = raw_slot
    else:
        slot = "晚自习1"
    return day, slot


def _derive_teacher_views(class_views: list[dict[str, Any]]) -> list[dict[str, Any]]:
    teacher_rows: dict[str, dict[str, dict[str, list[str]]]] = {}
    for class_view in class_views:
        class_name = str(class_view.get("name") or "")
        for row in class_view.get("rows") or []:
            slot = str(row.get("slot") or "")
            values = row.get("values") if isinstance(row.get("values"), dict) else {}
            for day, cell in values.items():
                text = str(cell or "").strip()
                for teacher in _extract_teachers(text):
                    teacher_rows.setdefault(teacher, {}).setdefault(slot, {d: [] for d in DAYS})
                    subject = _strip_teacher_marker(text, teacher)
                    label = f"{class_name} {subject}".strip()
                    teacher_rows[teacher][slot].setdefault(day, []).append(label)

    views: list[dict[str, Any]] = []
    for teacher, slot_map in teacher_rows.items():
        rows = []
        for slot, day_map in slot_map.items():
            values = {day: "；".join(day_map.get(day, [])) for day in DAYS}
            rows.append({"slot": slot, "values": values})
        views.append(_view("teacher", teacher, "由班级课表派生", "班级课表", rows))
    return sorted(views, key=lambda item: str(item.get("name") or ""))


def _build_quality_summary(class_views: list[dict[str, Any]]) -> dict[str, Any]:
    if not class_views:
        return _empty_quality()

    cell_rows: list[dict[str, str]] = []
    active_day_slots: set[tuple[str, str]] = set()
    for view in class_views:
        class_name = str(view.get("name") or "")
        for row in view.get("rows") or []:
            slot = str(row.get("slot") or "")
            if not _is_teaching_slot(slot):
                continue
            values = row.get("values") if isinstance(row.get("values"), dict) else {}
            for day, cell in values.items():
                text = str(cell or "").strip()
                day_slot = (str(day), slot)
                cell_rows.append({"class_name": class_name, "day": str(day), "slot": slot, "text": text})
                if text:
                    active_day_slots.add(day_slot)

    ignore_global_empty_slots = len(class_views) > 1
    inactive_day_slots = (
        {
            (row["day"], row["slot"])
            for row in cell_rows
            if (row["day"], row["slot"]) not in active_day_slots
        }
        if ignore_global_empty_slots
        else set()
    )
    inactive_cells_ignored = sum(
        1 for row in cell_rows if (row["day"], row["slot"]) in inactive_day_slots
    )
    total_cells = 0
    empty_cells = 0
    teacher_slot_map: dict[tuple[str, str, str], list[dict[str, str]]] = defaultdict(list)
    teacher_load: Counter[str] = Counter()
    class_load: Counter[str] = Counter()

    for row in cell_rows:
        day = row["day"]
        slot = row["slot"]
        if (day, slot) in inactive_day_slots:
            continue
        text = row["text"]
        total_cells += 1
        if not text:
            empty_cells += 1
            continue
        class_name = row["class_name"]
        class_load[class_name] += 1
        teachers = _extract_teachers(text)
        if not teachers:
            continue
        subject = text
        for teacher in teachers:
            teacher_load[teacher] += 1
            teacher_slot_map[(teacher, day, slot)].append(
                {
                    "class_name": class_name,
                    "subject": _strip_teacher_marker(subject, teacher),
                }
            )

    issues: list[dict[str, Any]] = []
    for (teacher, day, slot), events in sorted(teacher_slot_map.items()):
        classes = sorted({event["class_name"] for event in events if event.get("class_name")})
        if len(classes) <= 1:
            continue
        details = "、".join(f"{event['class_name']} {event['subject']}".strip() for event in events[:5])
        issues.append(
            {
                "severity": "error",
                "title": "教师同节次出现在多个班级",
                "detail": f"{teacher} 在 {day} {slot} 同时出现在：{details}。",
                "suggestion": "回到调课单或求解规则中调整其中一个班级的教师或节次。",
            }
        )

    if empty_cells:
        issues.append(
            {
                "severity": "warning",
                "title": "班级课表存在空课格",
                "detail": f"抽查到 {empty_cells}/{total_cells} 个教学课格为空。",
                "suggestion": "确认这些空格是合法空档；否则回到求解前校验补齐课时或学科教师。",
            }
        )

    if inactive_day_slots:
        issues.append(
            {
                "severity": "info",
                "title": "已忽略全校非开课空白时段",
                "detail": f"有 {len(inactive_day_slots)} 个日-节次组合全校均为空，共 {inactive_cells_ignored} 个空白格未计入漏排。",
                "suggestion": "这些空白按非开课时段处理；如果学校实际要求这些时段上课，请先回到时间格子配置确认。",
            }
        )

    load_summary = _teacher_load_summary(teacher_load)
    if teacher_load:
        if load_summary["teacher_load_spread"] >= 6:
            high = "、".join(
                f"{item['teacher']} {item['count']}节" for item in load_summary["high_load_teachers"][:3]
            )
            low = "、".join(
                f"{item['teacher']} {item['count']}节" for item in load_summary["low_load_teachers"][:3]
            )
            issues.append(
                {
                    "severity": "info",
                    "title": "教师课时负荷需按任课量复核",
                    "detail": (
                        "预览仅按课表出现次数统计，不等同于学校正式工作量口径。"
                        f"高负荷样本：{high}；低负荷样本：{low}；"
                        f"中位数 {load_summary['teacher_load_median']} 节，跨度 {load_summary['teacher_load_spread']} 节。"
                    ),
                    "suggestion": "发布前结合学科周课时、跨班兼课、值班和学校标准工作量确认；未配置目标课时前不要把该项视为漏排或硬阻断。",
                }
            )
    else:
        issues.append(
            {
                "severity": "warning",
                "title": "未识别到教师姓名",
                "detail": "班级课表有内容，但单元格未包含教师姓名或无法解析教师标记。",
                "suggestion": "建议导出包含“学科(教师)”格式的课表，方便发布前自动抽查教师冲突。",
            }
        )

    errors = sum(1 for issue in issues if issue["severity"] == "error")
    warnings = sum(1 for issue in issues if issue["severity"] == "warning")
    if errors:
        status = "blocked"
        label = "有冲突"
    elif warnings:
        status = "review"
        label = "需复核"
    else:
        status = "ok"
        label = "抽查通过"

    return {
        "summary": {
            "status": status,
            "status_label": label,
            "classes": len(class_views),
            "teaching_cells": total_cells,
            "empty_cells": empty_cells,
            "inactive_cells_ignored": inactive_cells_ignored,
            "inactive_slot_count": len(inactive_day_slots),
            "inactive_slots": [
                {"day": day, "slot": slot}
                for day, slot in sorted(inactive_day_slots, key=lambda item: (DAYS.index(item[0]) if item[0] in DAYS else 99, item[1]))
            ][:12],
            "filled_cells": max(0, total_cells - empty_cells),
            "teacher_conflicts": errors,
            "teacher_count": load_summary["teacher_count"],
            "max_teacher_load": load_summary["max_teacher_load"],
            "min_teacher_load": load_summary["min_teacher_load"],
            "teacher_load_avg": load_summary["teacher_load_avg"],
            "teacher_load_median": load_summary["teacher_load_median"],
            "teacher_load_spread": load_summary["teacher_load_spread"],
            "warnings": warnings,
            "errors": errors,
        },
        "issues": issues[:12],
        "teacher_load_top": [{"teacher": teacher, "count": count} for teacher, count in teacher_load.most_common(8)],
        "teacher_load_low": load_summary["low_load_teachers"][:8],
        "class_load_top": [{"class_name": cls, "count": count} for cls, count in class_load.most_common(8)],
    }


def _teacher_load_summary(teacher_load: Counter[str]) -> dict[str, Any]:
    if not teacher_load:
        return {
            "teacher_count": 0,
            "max_teacher_load": 0,
            "min_teacher_load": 0,
            "teacher_load_avg": 0,
            "teacher_load_median": 0,
            "teacher_load_spread": 0,
            "high_load_teachers": [],
            "low_load_teachers": [],
        }
    counts = sorted(int(value) for value in teacher_load.values())
    count_len = len(counts)
    middle = count_len // 2
    if count_len % 2:
        median = counts[middle]
    else:
        median = round((counts[middle - 1] + counts[middle]) / 2, 1)
    max_load = counts[-1]
    min_load = counts[0]
    return {
        "teacher_count": count_len,
        "max_teacher_load": max_load,
        "min_teacher_load": min_load,
        "teacher_load_avg": round(sum(counts) / count_len, 1),
        "teacher_load_median": median,
        "teacher_load_spread": max_load - min_load,
        "high_load_teachers": [
            {"teacher": teacher, "count": count}
            for teacher, count in teacher_load.most_common(8)
        ],
        "low_load_teachers": [
            {"teacher": teacher, "count": count}
            for teacher, count in sorted(teacher_load.items(), key=lambda item: (item[1], item[0]))[:8]
        ],
    }


def _teacher_day_load_summary(
    teacher_day_load: Counter[tuple[str, str]],
    teacher_day_slots: dict[tuple[str, str], list[dict[str, str]]] | None = None,
    class_slot_values: dict[tuple[str, str, str], str] | None = None,
    teacher_busy_slots: set[tuple[str, str, str]] | None = None,
) -> dict[str, Any]:
    rows = sorted(
        (
            {
                "teacher": str(teacher),
                "day": str(day),
                "count": int(count),
                "suggestion": _teacher_day_balance_suggestion(
                    str(teacher),
                    str(day),
                    int(count),
                    teacher_day_load,
                    teacher_day_slots or {},
                    class_slot_values or {},
                    teacher_busy_slots or set(),
                ),
            }
            for (teacher, day), count in teacher_day_load.items()
            if str(teacher).strip() and str(day).strip()
        ),
        key=lambda item: (
            -int(item["count"]),
            str(item["teacher"]),
            DAYS.index(str(item["day"])) if str(item["day"]) in DAYS else 99,
        ),
    )
    return {
        "top": rows[:12],
        "max": int(rows[0]["count"]) if rows else 0,
        "teacher_day_count": len(rows),
    }


def _teacher_day_balance_suggestion(
    teacher: str,
    source_day: str,
    source_count: int,
    teacher_day_load: Counter[tuple[str, str]],
    teacher_day_slots: dict[tuple[str, str], list[dict[str, str]]],
    class_slot_values: dict[tuple[str, str, str], str],
    teacher_busy_slots: set[tuple[str, str, str]],
) -> dict[str, Any] | None:
    if source_count < 4:
        return None
    sources = _preferred_teacher_day_sources(teacher, source_day, teacher_day_slots)
    for source in sources:
        class_name = str(source.get("class_name") or "").strip()
        source_slot = str(source.get("slot") or "").strip()
        source_text = str(source.get("text") or "").strip()
        if not class_name or not source_slot or not source_text:
            continue
        targets = _teacher_day_balance_targets(
            teacher,
            class_name,
            source_day,
            source_slot,
            source_count,
            teacher_day_load,
            class_slot_values,
            teacher_busy_slots,
        )
        for target in targets:
            return {
                "kind": "teacher_day_balance",
                "title": f"均衡教师日负荷：{teacher} {source_day}",
                "teacher": teacher,
                "class_name": class_name,
                "from_day": source_day,
                "from_slot": source_slot,
                "to_day": str(target.get("day") or ""),
                "to_slot": str(target.get("slot") or ""),
                "original_subject": _extract_subject(source_text),
                "original_teacher": teacher,
                "source_text": source_text,
                "target_before": target.get("before") or "",
                "target_kind": target.get("kind") or "",
                "reason": target.get("reason") or "",
            }
    return None


def _preferred_teacher_day_sources(
    teacher: str,
    source_day: str,
    teacher_day_slots: dict[tuple[str, str], list[dict[str, str]]],
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for entry in teacher_day_slots.get((teacher, source_day), []):
        class_name = str(entry.get("class_name") or "").strip()
        slot = str(entry.get("slot") or "").strip()
        text = str(entry.get("text") or "").strip()
        if not class_name or not slot or not text:
            continue
        if not _is_regular_lesson_slot(slot) or _is_self_study_cell(text):
            continue
        rows.append({"class_name": class_name, "slot": slot, "text": text})
    return sorted(
        rows,
        key=lambda item: (
            -_slot_phase_index(str(item.get("slot") or "")),
            -_slot_number_index(str(item.get("slot") or "")),
            str(item.get("class_name") or ""),
        ),
    )


def _teacher_day_balance_targets(
    teacher: str,
    class_name: str,
    source_day: str,
    source_slot: str,
    source_count: int,
    teacher_day_load: Counter[tuple[str, str]],
    class_slot_values: dict[tuple[str, str, str], str],
    teacher_busy_slots: set[tuple[str, str, str]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for (candidate_class, day, slot), value in class_slot_values.items():
        if candidate_class != class_name or day == source_day:
            continue
        if not _is_regular_lesson_slot(str(slot)):
            continue
        target_day_count = int(teacher_day_load.get((teacher, str(day)), 0))
        if target_day_count >= source_count - 1:
            continue
        if (teacher, str(day), str(slot)) in teacher_busy_slots:
            continue
        text = str(value or "").strip()
        if not text:
            target_kind = "empty"
            penalty = 0
            reason = "移动到教师低负荷日同班常规空课格"
        elif _is_self_study_cell(text):
            target_kind = "self_study"
            penalty = 4
            reason = "移动到教师低负荷日同班常规自习课格"
        else:
            continue
        balance_penalty = abs((source_count - 1) - (target_day_count + 1))
        rows.append(
            {
                "day": str(day),
                "slot": str(slot),
                "before": text,
                "kind": target_kind,
                "reason": reason,
                "score": (
                    balance_penalty * 20
                    + target_day_count * 8
                    + penalty
                    + _slot_phase_distance(source_slot, str(slot)) * 3
                    + _slot_number_distance(source_slot, str(slot))
                ),
            }
        )
    return sorted(
        rows,
        key=lambda item: (
            int(item["score"]),
            DAYS.index(str(item["day"])) if str(item["day"]) in DAYS else 99,
            _slot_sort_key(str(item["slot"])),
        ),
    )


def _class_day_load_summary(
    class_day_load: dict[str, Counter[str]],
    class_slot_values: dict[tuple[str, str, str], str],
    teacher_busy_slots: set[tuple[str, str, str]],
) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    for class_name, day_counts in class_day_load.items():
        class_name_text = str(class_name or "").strip()
        if not class_name_text:
            continue
        positive_counts = [int(count) for count in day_counts.values() if int(count) > 0]
        if not positive_counts:
            continue
        day_spread = max(positive_counts) - min(positive_counts)
        for day, count in day_counts.items():
            count_int = int(count)
            if count_int <= 0:
                continue
            suggestion = _class_day_balance_suggestion(
                class_name_text,
                str(day),
                count_int,
                day_spread,
                day_counts,
                class_slot_values,
                teacher_busy_slots,
            )
            rows.append(
                {
                    "class_name": class_name_text,
                    "day": str(day),
                    "count": count_int,
                    "day_load_spread": day_spread,
                    "suggestion": suggestion,
                }
            )

    rows = sorted(
        rows,
        key=lambda item: (
            -int(item["count"]),
            -int(item["day_load_spread"]),
            str(item["class_name"]),
            DAYS.index(str(item["day"])) if str(item["day"]) in DAYS else 99,
        ),
    )
    return {
        "top": rows[:12],
        "max": int(rows[0]["count"]) if rows else 0,
        "class_day_count": len(rows),
    }


def _class_day_balance_suggestion(
    class_name: str,
    source_day: str,
    source_count: int,
    day_spread: int,
    day_counts: Counter[str],
    class_slot_values: dict[tuple[str, str, str], str],
    teacher_busy_slots: set[tuple[str, str, str]],
) -> dict[str, Any] | None:
    if source_count < 4 or day_spread < 2:
        return None
    sources = _preferred_class_day_sources(class_name, source_day, class_slot_values)
    for source in sources:
        source_text = str(source.get("text") or "").strip()
        source_slot = str(source.get("slot") or "").strip()
        if not source_text or _is_self_study_cell(source_text):
            continue
        source_teachers = _extract_teachers(source_text)
        if not source_teachers:
            continue
        targets = _class_day_balance_targets(
            class_name,
            source_day,
            source_slot,
            source_text,
            source_count,
            day_counts,
            class_slot_values,
        )
        for target in targets:
            target_day = str(target.get("day") or "")
            target_slot = str(target.get("slot") or "")
            if any((teacher, target_day, target_slot) in teacher_busy_slots for teacher in source_teachers):
                continue
            return {
                "kind": "class_day_balance",
                "title": f"均衡班级日课量：{class_name} {source_day}",
                "class_name": class_name,
                "from_day": source_day,
                "from_slot": source_slot,
                "to_day": target_day,
                "to_slot": target_slot,
                "original_subject": _extract_subject(source_text),
                "original_teacher": source_teachers[0],
                "source_text": source_text,
                "target_before": target.get("before") or "",
                "target_kind": target.get("kind") or "",
                "reason": target.get("reason") or "",
            }
    return None


def _preferred_class_day_sources(
    class_name: str,
    source_day: str,
    class_slot_values: dict[tuple[str, str, str], str],
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for (candidate_class, day, slot), value in class_slot_values.items():
        if candidate_class != class_name or day != source_day:
            continue
        text = str(value or "").strip()
        if not text or _is_self_study_cell(text):
            continue
        rows.append({"slot": str(slot), "text": text})
    return sorted(
        rows,
        key=lambda item: (
            -_slot_phase_index(str(item.get("slot") or "")),
            -_slot_number_index(str(item.get("slot") or "")),
            str(item.get("slot") or ""),
        ),
    )


def _class_day_balance_targets(
    class_name: str,
    source_day: str,
    source_slot: str,
    source_text: str,
    source_count: int,
    day_counts: Counter[str],
    class_slot_values: dict[tuple[str, str, str], str],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    source_is_regular = _is_regular_lesson_slot(source_slot)
    for (candidate_class, day, slot), value in class_slot_values.items():
        if candidate_class != class_name or day == source_day:
            continue
        day_count = int(day_counts.get(day, 0))
        if day_count <= 0 or day_count >= source_count - 1:
            continue
        if source_is_regular and not _is_regular_lesson_slot(str(slot)):
            continue
        text = str(value or "").strip()
        if not text:
            target_kind = "empty"
            penalty = 0
            reason = "移动到低负荷日常规空课格"
        elif _is_self_study_cell(text):
            target_kind = "self_study"
            penalty = 4
            reason = "移动到低负荷日常规自习课格" if source_is_regular else "移动到低负荷日自习课格"
        else:
            continue
        phase_penalty = _slot_phase_distance(source_slot, str(slot)) * 3
        distance_penalty = _slot_number_distance(source_slot, str(slot))
        balance_penalty = _class_day_balance_penalty(source_day, str(day), day_counts)
        rows.append(
            {
                "day": str(day),
                "slot": str(slot),
                "before": text,
                "kind": target_kind,
                "reason": reason,
                "score": balance_penalty * 20 + day_count * 8 + penalty + phase_penalty + distance_penalty,
            }
        )
    return sorted(
        rows,
        key=lambda item: (
            int(item["score"]),
            DAYS.index(str(item["day"])) if str(item["day"]) in DAYS else 99,
            _slot_sort_key(str(item["slot"])),
        ),
    )


def _teacher_consecutive_summary(
    teacher_day_slots: dict[tuple[str, str], list[dict[str, str]]],
    class_slot_values: dict[tuple[str, str, str], str],
    teacher_busy_slots: set[tuple[str, str, str]],
) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []

    def push_run(teacher: str, day: str, run: list[dict[str, Any]]) -> None:
        if len(run) < 3:
            return
        slots = [str(item.get("slot") or "") for item in run if str(item.get("slot") or "").strip()]
        classes = sorted({
            str(class_name)
            for item in run
            for class_name in item.get("classes", set())
            if str(class_name).strip()
        })
        subjects = sorted({
            str(subject)
            for item in run
            for subject in item.get("subjects", set())
            if str(subject).strip()
        })
        suggestion = _consecutive_break_suggestion(
            teacher,
            day,
            run,
            class_slot_values,
            teacher_busy_slots,
        )
        rows.append(
            {
                "teacher": teacher,
                "day": day,
                "streak": len(run),
                "slots": slots,
                "slot_text": _format_slot_run(slots),
                "classes": classes[:6],
                "class_count": len(classes),
                "subjects": subjects[:6],
                "suggestion": suggestion,
            }
        )

    for (teacher, day), entries in teacher_day_slots.items():
        teacher_name = str(teacher or "").strip()
        day_name = str(day or "").strip()
        if not teacher_name or not day_name:
            continue
        slot_map: dict[str, dict[str, Any]] = {}
        for entry in entries:
            slot = str(entry.get("slot") or "").strip()
            seq = _slot_sequence_key(slot)
            if not slot or seq is None:
                continue
            item = slot_map.setdefault(
                slot,
                {
                    "slot": slot,
                    "seq": seq,
                    "classes": set(),
                    "subjects": set(),
                    "entries": [],
                },
            )
            class_name = str(entry.get("class_name") or "").strip()
            subject = str(entry.get("subject") or "").strip()
            if class_name:
                item["classes"].add(class_name)
            if subject and subject != "未识别":
                item["subjects"].add(subject)
            item["entries"].append(
                {
                    "class_name": class_name,
                    "subject": subject,
                    "slot": slot,
                    "text": str(entry.get("text") or "").strip(),
                }
            )

        ordered = sorted(
            slot_map.values(),
            key=lambda item: (item["seq"][0], item["seq"][1], _slot_sort_key(str(item.get("slot") or ""))),
        )
        run: list[dict[str, Any]] = []
        previous_seq: tuple[int, int] | None = None
        for item in ordered:
            seq = item["seq"]
            if run and previous_seq is not None and seq[0] == previous_seq[0] and seq[1] == previous_seq[1] + 1:
                run.append(item)
            else:
                push_run(teacher_name, day_name, run)
                run = [item]
            previous_seq = seq
        push_run(teacher_name, day_name, run)

    rows = sorted(
        rows,
        key=lambda item: (
            -int(item.get("streak") or 0),
            str(item.get("teacher") or ""),
            DAYS.index(str(item.get("day") or "")) if str(item.get("day") or "") in DAYS else 99,
            _slot_sort_key(str((item.get("slots") or [""])[0] if item.get("slots") else "")),
        ),
    )
    return {
        "top": rows[:12],
        "max_streak": int(rows[0]["streak"]) if rows else 0,
        "streak_count": len(rows),
    }


def _consecutive_break_suggestion(
    teacher: str,
    day: str,
    run: list[dict[str, Any]],
    class_slot_values: dict[tuple[str, str, str], str],
    teacher_busy_slots: set[tuple[str, str, str]],
) -> dict[str, Any] | None:
    source_items = _preferred_break_sources(run)
    for source_item in source_items:
        source_slot = str(source_item.get("slot") or "")
        for entry in source_item.get("entries") or []:
            class_name = str(entry.get("class_name") or "").strip()
            if not class_name:
                continue
            target = _best_consecutive_break_target(
                teacher,
                class_name,
                day,
                source_slot,
                class_slot_values,
                teacher_busy_slots,
            )
            if not target:
                continue
            return {
                "kind": "consecutive_break",
                "title": f"拆解连续课：{teacher} {day} {source_slot}",
                "teacher": teacher,
                "class_name": class_name,
                "from_day": day,
                "from_slot": source_slot,
                "to_day": target["day"],
                "to_slot": target["slot"],
                "original_subject": entry.get("subject") or "",
                "original_teacher": teacher,
                "source_text": entry.get("text") or "",
                "target_before": target["before"],
                "target_kind": target["kind"],
                "reason": target["reason"],
            }
    return None


def _preferred_break_sources(run: list[dict[str, Any]]) -> list[dict[str, Any]]:
    total = len(run)
    indexed = list(enumerate(run))
    if total <= 3:
        midpoint = total // 2
        return [item for _, item in sorted(indexed, key=lambda pair: (abs(pair[0] - midpoint), pair[0]))]
    return [
        item
        for index, item in sorted(
            indexed,
            key=lambda pair: (
                0 if 0 < pair[0] < total - 1 else 1,
                abs(pair[0] - (total - 1) / 2),
                pair[0],
            ),
        )
    ]


def _best_consecutive_break_target(
    teacher: str,
    class_name: str,
    source_day: str,
    source_slot: str,
    class_slot_values: dict[tuple[str, str, str], str],
    teacher_busy_slots: set[tuple[str, str, str]],
) -> dict[str, Any] | None:
    candidates: list[dict[str, Any]] = []
    source_seq = _slot_sequence_key(source_slot)
    source_is_regular = _is_regular_lesson_slot(source_slot)
    for (candidate_class, day, slot), value in class_slot_values.items():
        if candidate_class != class_name:
            continue
        if day == source_day and slot == source_slot:
            continue
        if source_is_regular and not _is_regular_lesson_slot(str(slot)):
            continue
        if (teacher, day, slot) in teacher_busy_slots:
            continue
        destination = _break_destination_type(
            value,
            source_day,
            source_slot,
            day,
            slot,
            teacher_busy_slots,
            source_is_regular=source_is_regular,
        )
        if not destination:
            continue
        target_seq = _slot_sequence_key(slot)
        same_day = day == source_day
        distance = (
            abs(target_seq[1] - source_seq[1])
            if source_seq is not None and target_seq is not None and source_seq[0] == target_seq[0]
            else 8
        )
        candidates.append(
            {
                "day": day,
                "slot": slot,
                "before": value,
                "kind": destination["kind"],
                "reason": destination["reason"],
                "score": int(destination["penalty"]) + (0 if same_day else 12) + distance,
            }
        )
    return sorted(
        candidates,
        key=lambda item: (
            int(item["score"]),
            DAYS.index(str(item["day"])) if str(item["day"]) in DAYS else 99,
            _slot_sort_key(str(item["slot"])),
        ),
    )[0] if candidates else None


def _break_destination_type(
    value: str,
    source_day: str,
    source_slot: str,
    target_day: str,
    target_slot: str,
    teacher_busy_slots: set[tuple[str, str, str]],
    *,
    source_is_regular: bool = False,
) -> dict[str, Any] | None:
    text = str(value or "").strip()
    if not text:
        reason = "移动到同班常规空课格" if source_is_regular else "移动到同班空课格"
        return {"kind": "empty", "penalty": 0, "reason": reason}
    if _is_self_study_cell(text):
        reason = "移动到同班常规自习课格" if source_is_regular else "移动到同班自习课格"
        return {"kind": "self_study", "penalty": 4, "reason": reason}
    target_teachers = _extract_teachers(text)
    if target_teachers and all((teacher, source_day, source_slot) not in teacher_busy_slots for teacher in target_teachers):
        return {"kind": "swap", "penalty": 12, "reason": "与同班低冲突课位互换"}
    return None


def _build_schedule_statistics(
    status: dict[str, Any],
    class_views: list[dict[str, Any]],
    teacher_views: list[dict[str, Any]],
    project_root: Path | None,
) -> dict[str, Any]:
    if not class_views:
        return _empty_statistics()

    total_cells = 0
    filled_cells = 0
    empty_cells = 0
    self_study_cells = 0
    teacher_load: Counter[str] = Counter()
    teacher_day_load: Counter[tuple[str, str]] = Counter()
    teacher_day_slots: defaultdict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    teacher_busy_slots: set[tuple[str, str, str]] = set()
    class_slot_values: dict[tuple[str, str, str], str] = {}
    subject_load: Counter[str] = Counter()
    day_load: Counter[str] = Counter({day: 0 for day in DAYS})
    slot_load: Counter[str] = Counter()
    class_load: Counter[str] = Counter()
    class_day_load: dict[str, Counter[str]] = defaultdict(Counter)

    for view in class_views:
        class_name = str(view.get("name") or "")
        if class_name:
            class_load.setdefault(class_name, 0)
        for row in view.get("rows") or []:
            slot = str(row.get("slot") or "")
            if not _is_teaching_slot(slot):
                continue
            values = row.get("values") if isinstance(row.get("values"), dict) else {}
            for day in DAYS:
                text = str(values.get(day) or "").strip()
                if class_name:
                    class_slot_values[(class_name, day, slot)] = text
                total_cells += 1
                if not text:
                    empty_cells += 1
                    continue
                filled_cells += 1
                if class_name:
                    class_load[class_name] += 1
                    class_day_load[class_name][day] += 1
                day_load[day] += 1
                slot_load[slot] += 1
                subject = _extract_subject(text)
                if subject:
                    subject_load[subject] += 1
                if _is_self_study_cell(text):
                    self_study_cells += 1
                for teacher in _extract_teachers(text):
                    teacher_load[teacher] += 1
                    teacher_day_load[(teacher, day)] += 1
                    teacher_busy_slots.add((teacher, day, slot))
                    teacher_day_slots[(teacher, day)].append(
                        {
                            "slot": slot,
                            "class_name": class_name,
                            "subject": subject,
                            "text": text,
                        }
                    )

    teacher_summary = _teacher_load_summary(teacher_load)
    teacher_day_summary = _teacher_day_load_summary(
        teacher_day_load,
        teacher_day_slots,
        class_slot_values,
        teacher_busy_slots,
    )
    teacher_consecutive_summary = _teacher_consecutive_summary(
        teacher_day_slots,
        class_slot_values,
        teacher_busy_slots,
    )
    class_day_summary = _class_day_load_summary(
        class_day_load,
        class_slot_values,
        teacher_busy_slots,
    )
    class_counts = list(class_load.values())
    class_max = max(class_counts) if class_counts else 0
    class_min = min(class_counts) if class_counts else 0
    class_avg = round(sum(class_counts) / len(class_counts), 1) if class_counts else 0
    day_distribution = _counter_distribution(day_load, "day", total=filled_cells, order=DAYS)
    slot_distribution = _counter_distribution(slot_load, "slot", total=filled_cells, sort_key=lambda item: _slot_sort_key(item[0]))
    subject_distribution = _counter_distribution(subject_load, "subject", total=filled_cells, limit=12)
    repair_history = status.get("local_timetable_repairs") if isinstance(status.get("local_timetable_repairs"), list) else []
    adjustment_history = status.get("local_timetable_adjustments") if isinstance(status.get("local_timetable_adjustments"), list) else []
    active_repair = status.get("local_timetable_repair") if isinstance(status.get("local_timetable_repair"), dict) else {}
    active_adjustment = status.get("local_timetable_adjustment") if isinstance(status.get("local_timetable_adjustment"), dict) else {}
    adjustment_records = [item for item in [*adjustment_history, *repair_history] if isinstance(item, dict)]
    latest_adjustment_summary = (
        adjustment_records[0].get("summary")
        if adjustment_records and isinstance(adjustment_records[0].get("summary"), dict)
        else None
    )
    insights = _build_schedule_insights(
        total_cells=total_cells,
        filled_cells=filled_cells,
        empty_cells=empty_cells,
        self_study_cells=self_study_cells,
        teacher_summary=teacher_summary,
        class_load=class_load,
        day_load=day_load,
        slot_load=slot_load,
        subject_load=subject_load,
        teacher_day_summary=teacher_day_summary,
        teacher_consecutive_summary=teacher_consecutive_summary,
        class_day_summary=class_day_summary,
        adjustment_records=adjustment_records,
    )
    adjustment_suggestions = _build_adjustment_suggestion_queue(
        teacher_day_summary,
        teacher_consecutive_summary,
        class_day_summary,
    )
    rule_explain = build_rule_explain(status, project_root)

    return {
        "schema_version": "scheduler.result_schedule_statistics.v1",
        "summary": {
            "status": "ready" if filled_cells else "empty",
            "status_label": "已统计" if filled_cells else "暂无统计",
            "classes": len(class_views),
            "teachers": teacher_summary["teacher_count"] or len(teacher_views),
            "total_cells": total_cells,
            "filled_cells": filled_cells,
            "empty_cells": empty_cells,
            "fill_rate": round(filled_cells / total_cells * 100, 1) if total_cells else 0,
            "self_study_cells": self_study_cells,
            "avg_class_lessons": class_avg,
            "class_load_spread": class_max - class_min if class_counts else 0,
            "teacher_load_spread": teacher_summary["teacher_load_spread"],
            "max_teacher_consecutive": teacher_consecutive_summary["max_streak"],
            "max_class_day_load": class_day_summary["max"],
            "busiest_day": _top_counter_label(day_load),
            "busiest_slot": _top_counter_label(slot_load),
            "top_subject": _top_counter_label(subject_load),
        },
        "day_distribution": day_distribution,
        "slot_distribution": slot_distribution[:12],
        "subject_distribution": subject_distribution,
        "insights": insights,
        "adjustment_suggestions": adjustment_suggestions,
        "teacher_load": {
            "top": teacher_summary["high_load_teachers"][:10],
            "low": teacher_summary["low_load_teachers"][:10],
            "avg": teacher_summary["teacher_load_avg"],
            "median": teacher_summary["teacher_load_median"],
            "max": teacher_summary["max_teacher_load"],
            "min": teacher_summary["min_teacher_load"],
        },
        "teacher_day_load": teacher_day_summary,
        "teacher_consecutive_load": teacher_consecutive_summary,
        "class_day_load": class_day_summary,
        "class_load": {
            "top": [{"class_name": name, "count": count} for name, count in class_load.most_common(8)],
            "low": [{"class_name": name, "count": count} for name, count in sorted(class_load.items(), key=lambda item: (item[1], item[0]))[:8]],
            "avg": class_avg,
            "max": class_max,
            "min": class_min,
        },
        "micro_adjustment": {
            "active_adjusted_schedule": str(active_adjustment.get("active_file") or active_repair.get("active_file") or ""),
            "repair_count": len(adjustment_records),
            "last_summary": latest_adjustment_summary,
            "affected_teachers": sorted({
                str(teacher)
                for item in adjustment_records
                if isinstance(item, dict)
                for teacher in (item.get("summary", {}).get("involved_teachers") if isinstance(item.get("summary"), dict) else []) or []
                if str(teacher).strip()
            }),
        },
        "rule_explain": rule_explain,
    }


def _empty_statistics() -> dict[str, Any]:
    return {
        "schema_version": "scheduler.result_schedule_statistics.v1",
        "summary": {
            "status": "empty",
            "status_label": "暂无统计",
            "classes": 0,
            "teachers": 0,
            "total_cells": 0,
            "filled_cells": 0,
            "empty_cells": 0,
            "fill_rate": 0,
            "self_study_cells": 0,
            "avg_class_lessons": 0,
            "class_load_spread": 0,
            "teacher_load_spread": 0,
            "max_teacher_consecutive": 0,
            "max_class_day_load": 0,
            "busiest_day": "",
            "busiest_slot": "",
            "top_subject": "",
        },
        "day_distribution": [],
        "slot_distribution": [],
        "subject_distribution": [],
        "insights": [],
        "adjustment_suggestions": {"top": [], "count": 0},
        "teacher_load": {"top": [], "low": [], "avg": 0, "median": 0, "max": 0, "min": 0},
        "teacher_day_load": {"top": [], "max": 0, "teacher_day_count": 0},
        "teacher_consecutive_load": {"top": [], "max_streak": 0, "streak_count": 0},
        "class_day_load": {"top": [], "max": 0, "class_day_count": 0},
        "class_load": {"top": [], "low": [], "avg": 0, "max": 0, "min": 0},
        "micro_adjustment": {"active_adjusted_schedule": "", "repair_count": 0, "last_summary": None, "affected_teachers": []},
        "rule_explain": empty_rule_explain(),
    }


def _build_schedule_insights(
    *,
    total_cells: int,
    filled_cells: int,
    empty_cells: int,
    self_study_cells: int,
    teacher_summary: dict[str, Any],
    class_load: Counter[str],
    day_load: Counter[str],
    slot_load: Counter[str],
    subject_load: Counter[str],
    teacher_day_summary: dict[str, Any],
    teacher_consecutive_summary: dict[str, Any],
    class_day_summary: dict[str, Any],
    adjustment_records: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    insights: list[dict[str, Any]] = []

    def add(
        severity: str,
        title: str,
        detail: str,
        suggestion: str,
        *,
        metric_label: str = "",
        metric_value: str | int | float = "",
        related_teachers: list[str] | None = None,
        related_classes: list[str] | None = None,
    ) -> None:
        insights.append(
            {
                "severity": severity,
                "title": title,
                "detail": detail,
                "suggestion": suggestion,
                "metric_label": metric_label,
                "metric_value": metric_value,
                "related_teachers": related_teachers or [],
                "related_classes": related_classes or [],
            }
        )

    fill_rate = round(filled_cells / total_cells * 100, 1) if total_cells else 0
    if total_cells and empty_cells:
        severity = "warning" if fill_rate < 95 else "info"
        add(
            severity,
            "存在空课格需要确认",
            f"教学课格填充率 {fill_rate}%，仍有 {empty_cells} 个课格为空。",
            "先确认这些空课格是否属于学校允许的非开课时段；如果不是，应回到课时、固定课位或调课单补齐。",
            metric_label="空课格",
            metric_value=empty_cells,
        )

    spread = int(teacher_summary.get("teacher_load_spread") or 0)
    if spread >= 8:
        high = teacher_summary.get("high_load_teachers", [])[:3]
        low = teacher_summary.get("low_load_teachers", [])[:3]
        high_text = "、".join(f"{item['teacher']} {item['count']}节" for item in high if item.get("teacher"))
        low_text = "、".join(f"{item['teacher']} {item['count']}节" for item in low if item.get("teacher"))
        add(
            "warning",
            "教师负荷跨度偏大",
            f"教师课表出现次数跨度 {spread} 节。高负荷样本：{high_text or '无'}；低负荷样本：{low_text or '无'}。",
            "结合学校正式工作量口径确认是否需要调整跨班兼课、课时分配或后续人工微调。",
            metric_label="负荷跨度",
            metric_value=f"{spread} 节",
            related_teachers=[str(item.get("teacher")) for item in high if str(item.get("teacher") or "").strip()],
        )
    elif spread >= 5:
        add(
            "info",
            "教师负荷存在可复查差异",
            f"教师课表出现次数跨度 {spread} 节，暂未达到高风险阈值。",
            "可在教师视图中抽查最高和最低负荷教师，确认是否符合年级组实际分工。",
            metric_label="负荷跨度",
            metric_value=f"{spread} 节",
        )

    class_counts = list(class_load.values())
    class_spread = (max(class_counts) - min(class_counts)) if class_counts else 0
    if class_spread >= 3:
        high_classes = [name for name, _ in class_load.most_common(3)]
        low_classes = [name for name, _ in sorted(class_load.items(), key=lambda item: (item[1], item[0]))[:3]]
        add(
            "warning",
            "班级课量不均衡",
            f"班级已排课量跨度 {class_spread} 节，高低班级样本分别为：{'、'.join(high_classes)} / {'、'.join(low_classes)}。",
            "核对班级差异表、走班安排和临时调课记录，确认是否存在某些班级漏排或重复排课。",
            metric_label="班级跨度",
            metric_value=f"{class_spread} 节",
            related_classes=high_classes + low_classes,
        )

    positive_days = {day: count for day, count in day_load.items() if count > 0}
    if len(positive_days) >= 2:
        max_day, max_count = max(positive_days.items(), key=lambda item: (item[1], item[0]))
        min_day, min_count = min(positive_days.items(), key=lambda item: (item[1], item[0]))
        day_spread = max_count - min_count
        class_count = max(1, len(class_load))
        if day_spread >= class_count * 2:
            add(
                "info",
                "星期分布差异较明显",
                f"{max_day} 已排 {max_count} 个课格，{min_day} 已排 {min_count} 个课格，差值 {day_spread}。",
                "如果学校希望周内负荷更均衡，可复查固定课位、周末课时和软约束权重。",
                metric_label="日分布差",
                metric_value=day_spread,
            )

    if filled_cells and self_study_cells:
        self_study_rate = round(self_study_cells / filled_cells * 100, 1)
        if self_study_rate >= 8:
            add(
                "info",
                "自习课占比较高",
                f"课表中识别到 {self_study_cells} 个自习课格，占已填课格 {self_study_rate}%。",
                "确认这些自习是否符合年级教学计划；如果用于兜底排课，建议继续补足学科教师或课时规则。",
                metric_label="自习占比",
                metric_value=f"{self_study_rate}%",
            )

    teacher_day_top = teacher_day_summary.get("top") if isinstance(teacher_day_summary.get("top"), list) else []
    max_teacher_day_load = int(teacher_day_summary.get("max") or 0)
    if max_teacher_day_load >= 4 and teacher_day_top:
        samples = "、".join(
            f"{item.get('teacher')} {item.get('day')} {item.get('count')}节"
            for item in teacher_day_top[:3]
            if item.get("teacher") and item.get("day")
        )
        add(
            "warning" if max_teacher_day_load >= 6 else "info",
            "教师单日课量偏高",
            f"最高单日教师课量为 {max_teacher_day_load} 节。样本：{samples or '无'}。",
            "建议抽查对应教师当天是否存在连续课、跨楼栋奔波或备课压力过高；必要时通过课表微调沙盘分散到其他日期。",
            metric_label="单日最高",
            metric_value=f"{max_teacher_day_load} 节",
            related_teachers=[
                str(item.get("teacher"))
                for item in teacher_day_top[:4]
                if str(item.get("teacher") or "").strip()
            ],
        )

    consecutive_top = teacher_consecutive_summary.get("top") if isinstance(teacher_consecutive_summary.get("top"), list) else []
    max_consecutive = int(teacher_consecutive_summary.get("max_streak") or 0)
    if max_consecutive >= 3 and consecutive_top:
        samples = "、".join(
            f"{item.get('teacher')} {item.get('day')} {item.get('slot_text') or '连续课'} {item.get('streak')}连"
            for item in consecutive_top[:3]
            if item.get("teacher") and item.get("day")
        )
        add(
            "warning" if max_consecutive >= 4 else "info",
            "教师连续课风险",
            f"最长连续课达到 {max_consecutive} 节。样本：{samples or '无'}。",
            "优先抽查对应教师当天课表；如不是学校明确允许的连堂课，建议在课表微调沙盘中与同学科或同年级低冲突课位交换。",
            metric_label="最长连续",
            metric_value=f"{max_consecutive} 连",
            related_teachers=[
                str(item.get("teacher"))
                for item in consecutive_top[:4]
                if str(item.get("teacher") or "").strip()
            ],
        )

    class_day_top = class_day_summary.get("top") if isinstance(class_day_summary.get("top"), list) else []
    max_class_day_load = int(class_day_summary.get("max") or 0)
    if max_class_day_load >= 5 and class_day_top:
        samples = "、".join(
            f"{item.get('class_name')} {item.get('day')} {item.get('count')}节"
            for item in class_day_top[:3]
            if item.get("class_name") and item.get("day")
        )
        add(
            "warning" if max_class_day_load >= 7 else "info",
            "班级单日课量偏高",
            f"最高班级单日课量为 {max_class_day_load} 节。样本：{samples or '无'}。",
            "建议抽查对应班级当天课表；如某天明显过满，可优先把非固定课移动到同班低负荷日的空课或自习课格。",
            metric_label="班级单日最高",
            metric_value=f"{max_class_day_load} 节",
            related_classes=[
                str(item.get("class_name"))
                for item in class_day_top[:4]
                if str(item.get("class_name") or "").strip()
            ],
        )

    if adjustment_records:
        teacher_names = sorted({
            str(teacher)
            for record in adjustment_records
            if isinstance(record.get("summary"), dict)
            for teacher in (record["summary"].get("involved_teachers") or [])
            if str(teacher).strip()
        })
        add(
            "info",
            "当前课表已有局部调整历史",
            f"已记录 {len(adjustment_records)} 次课表微调或请假代课局部调整。",
            "下载或继续调整前，先在版本时间线中确认当前全局课表是否为最新调整版。",
            metric_label="调整批次",
            metric_value=len(adjustment_records),
            related_teachers=teacher_names[:8],
        )

    if not insights and filled_cells:
        top_subject, top_count = subject_load.most_common(1)[0] if subject_load else ("", 0)
        busiest_slot, busiest_count = slot_load.most_common(1)[0] if slot_load else ("", 0)
        add(
            "ok",
            "未发现明显统计异常",
            f"课格填充率 {fill_rate}%，最高频学科 {top_subject or '无'} {top_count} 格，最高频节次 {busiest_slot or '无'} {busiest_count} 格。",
            "仍建议按班级、教师两个视角抽查重点班级和高负荷教师后再使用。",
            metric_label="填充率",
            metric_value=f"{fill_rate}%",
        )

    severity_rank = {"warning": 0, "info": 1, "ok": 2}
    return sorted(insights, key=lambda item: (severity_rank.get(str(item.get("severity")), 9), str(item.get("title") or "")))[:8]


def _build_adjustment_suggestion_queue(
    teacher_day_summary: dict[str, Any],
    teacher_consecutive_summary: dict[str, Any],
    class_day_summary: dict[str, Any],
) -> dict[str, Any]:
    suggestions: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str, str, str, str]] = set()
    source_labels = {
        "teacher_day_load": "教师日负荷均衡",
        "teacher_consecutive_load": "连续课拆解",
        "class_day_load": "班级日课量均衡",
    }

    def add(
        *,
        source: str,
        title: str,
        detail: str,
        reason: str,
        metric_label: str,
        metric_value: str,
        priority: int,
        suggestion: dict[str, Any] | None,
        related_teacher: str = "",
        related_class: str = "",
    ) -> None:
        if not suggestion:
            return
        key = (
            str(suggestion.get("class_name") or ""),
            str(suggestion.get("from_day") or ""),
            str(suggestion.get("from_slot") or ""),
            str(suggestion.get("to_day") or ""),
            str(suggestion.get("to_slot") or ""),
            str(suggestion.get("kind") or source),
        )
        if key in seen:
            return
        seen.add(key)
        suggestions.append(
            {
                "source": source,
                "source_label": source_labels.get(source, "课表微调建议"),
                "title": title,
                "detail": detail,
                "reason": reason,
                "metric_label": metric_label,
                "metric_value": metric_value,
                "priority": priority,
                "priority_label": "优先处理" if priority >= 80 else "建议处理",
                "related_teacher": related_teacher,
                "related_class": related_class or str(suggestion.get("class_name") or ""),
                "suggestion": suggestion,
            }
        )

    teacher_day_rows = teacher_day_summary.get("top") if isinstance(teacher_day_summary.get("top"), list) else []
    for item in teacher_day_rows:
        suggestion = item.get("suggestion") if isinstance(item, dict) and isinstance(item.get("suggestion"), dict) else None
        teacher = str(item.get("teacher") or "").strip()
        day = str(item.get("day") or "").strip()
        count = int(item.get("count") or 0)
        add(
            source="teacher_day_load",
            title=f"均衡教师日负荷：{teacher} {day}",
            detail=f"{teacher} {day} 已排 {count} 节，建议先把 {suggestion.get('class_name') if suggestion else ''} {suggestion.get('from_slot') if suggestion else ''} 移到 {suggestion.get('to_day') if suggestion else ''}{suggestion.get('to_slot') if suggestion else ''}。",
            reason=str(suggestion.get("reason") or "降低教师单日课量集中度") if suggestion else "",
            metric_label="教师日课量",
            metric_value=f"{count} 节",
            priority=88 if count >= 6 else 68,
            suggestion=suggestion,
            related_teacher=teacher,
            related_class=str((suggestion or {}).get("class_name") or ""),
        )

    consecutive_rows = teacher_consecutive_summary.get("top") if isinstance(teacher_consecutive_summary.get("top"), list) else []
    for item in consecutive_rows:
        suggestion = item.get("suggestion") if isinstance(item, dict) and isinstance(item.get("suggestion"), dict) else None
        teacher = str(item.get("teacher") or "").strip()
        day = str(item.get("day") or "").strip()
        streak = int(item.get("streak") or 0)
        add(
            source="teacher_consecutive_load",
            title=f"拆解连续课：{teacher} {day}",
            detail=f"{teacher} {day} 出现 {streak} 连课，建议先把 {suggestion.get('class_name') if suggestion else ''} {suggestion.get('from_slot') if suggestion else ''} 移到 {suggestion.get('to_day') if suggestion else ''}{suggestion.get('to_slot') if suggestion else ''}。",
            reason=str(suggestion.get("reason") or "降低教师连续授课压力") if suggestion else "",
            metric_label="连续课",
            metric_value=f"{streak} 连",
            priority=90 if streak >= 4 else 70,
            suggestion=suggestion,
            related_teacher=teacher,
            related_class=str((suggestion or {}).get("class_name") or ""),
        )

    class_day_rows = class_day_summary.get("top") if isinstance(class_day_summary.get("top"), list) else []
    for item in class_day_rows:
        suggestion = item.get("suggestion") if isinstance(item, dict) and isinstance(item.get("suggestion"), dict) else None
        class_name = str(item.get("class_name") or "").strip()
        day = str(item.get("day") or "").strip()
        count = int(item.get("count") or 0)
        add(
            source="class_day_load",
            title=f"均衡班级日课量：{class_name} {day}",
            detail=f"{class_name} {day} 已排 {count} 节，建议把 {suggestion.get('from_slot') if suggestion else ''} 移到 {suggestion.get('to_day') if suggestion else ''}{suggestion.get('to_slot') if suggestion else ''}。",
            reason=str(suggestion.get("reason") or "降低班级单日课量集中度") if suggestion else "",
            metric_label="班级日课量",
            metric_value=f"{count} 节",
            priority=85 if count >= 7 else 65,
            suggestion=suggestion,
            related_teacher=str((suggestion or {}).get("original_teacher") or ""),
            related_class=class_name,
        )

    suggestions = sorted(
        suggestions,
        key=lambda item: (
            -int(item.get("priority") or 0),
            str(item.get("source_label") or ""),
            str(item.get("title") or ""),
        ),
    )
    visible: list[dict[str, Any]] = []
    visible_ids: set[int] = set()

    def push_visible(item: dict[str, Any] | None) -> None:
        if not item:
            return
        marker = id(item)
        if marker in visible_ids:
            return
        visible_ids.add(marker)
        visible.append(item)

    for item in suggestions[:3]:
        push_visible(item)
    for source in ("teacher_day_load", "teacher_consecutive_load", "class_day_load"):
        push_visible(next((item for item in suggestions if item.get("source") == source), None))
    for item in suggestions:
        push_visible(item)
    return {
        "top": visible[:10],
        "count": len(suggestions),
    }


def _build_schedule_versions(status: dict[str, Any], schedule_files: list[dict[str, Any]]) -> dict[str, Any]:
    active_file = _active_schedule_file(status)
    active_resolved = _resolve_path_text(active_file)
    base_file = _base_schedule_file(schedule_files, active_resolved)
    items: list[dict[str, Any]] = []
    if base_file:
        base_path = str(base_file.get("path") or "")
        items.append(
            {
                "id": "base",
                "kind": "base",
                "kind_label": "原始求解课表",
                "label": _version_label(base_file.get("label") or base_file.get("path") or "原始课表"),
                "path": base_path,
                "active": not active_resolved or _resolve_path_text(base_path) == active_resolved,
                "applied_at": "",
                "summary": {
                    "message": "求解器生成的原始课表。",
                    "changed_cell_count": 0,
                    "involved_teacher_count": 0,
                    "involved_teachers": [],
                    "affected_classes": [],
                },
                "changes": [],
            }
        )

    for index, entry in enumerate(_schedule_version_history(status), start=1):
        record = entry["record"]
        changes = _record_changes(record, entry["kind"])
        summary = record.get("summary") if isinstance(record.get("summary"), dict) else {}
        active_path = str(record.get("active_file") or "")
        involved_teachers = _record_involved_teachers(summary, changes)
        affected_classes = _record_affected_classes(summary, changes)
        changed_cell_count = int(summary.get("changed_cell_count") or len(changes))
        items.append(
            {
                "id": f"{entry['kind']}_{index}",
                "kind": entry["kind"],
                "kind_label": entry["kind_label"],
                "label": _version_label(active_path or entry["kind_label"]),
                "path": active_path,
                "active": bool(active_resolved and _resolve_path_text(active_path) == active_resolved),
                "applied_at": str(record.get("applied_at") or ""),
                "summary": {
                    "message": str(summary.get("message") or _record_title(record, entry["kind_label"])),
                    "changed_cell_count": changed_cell_count,
                    "involved_teacher_count": len(involved_teachers),
                    "involved_teachers": involved_teachers,
                    "affected_classes": affected_classes,
                },
                "changes": changes[:16],
            }
        )

    if items and not any(item.get("active") for item in items):
        items[-1]["active"] = True
    active_item = next((item for item in items if item.get("active")), items[-1] if items else {})
    changed_total = sum(int(item.get("summary", {}).get("changed_cell_count") or 0) for item in items if item.get("kind") != "base")
    return {
        "schema_version": "scheduler.result_schedule_versions.v1",
        "summary": {
            "status": "ready" if items else "empty",
            "version_count": len(items),
            "adjustment_count": max(0, len(items) - 1),
            "changed_cell_count": changed_total,
            "active_version": active_item.get("label", ""),
            "active_kind": active_item.get("kind_label", ""),
        },
        "items": items,
    }


def _empty_versions() -> dict[str, Any]:
    return {
        "schema_version": "scheduler.result_schedule_versions.v1",
        "summary": {
            "status": "empty",
            "version_count": 0,
            "adjustment_count": 0,
            "changed_cell_count": 0,
            "active_version": "",
            "active_kind": "",
        },
        "items": [],
    }


def _active_schedule_file(status: dict[str, Any]) -> str:
    for key in ("local_timetable_adjustment", "local_timetable_repair"):
        local = status.get(key) if isinstance(status.get(key), dict) else {}
        active = str(local.get("active_file") or "").strip()
        if active:
            return active
    return ""


def _base_schedule_file(schedule_files: list[dict[str, Any]], active_resolved: str) -> dict[str, Any] | None:
    if not schedule_files:
        return None
    if active_resolved:
        for item in schedule_files:
            if _resolve_path_text(str(item.get("path") or "")) != active_resolved:
                return item
    return schedule_files[0]


def _schedule_version_history(status: dict[str, Any]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for kind, label, key in (
        ("manual_adjustment", "课表微调", "local_timetable_adjustments"),
        ("leave_repair", "请假代课调课", "local_timetable_repairs"),
    ):
        records = status.get(key) if isinstance(status.get(key), list) else []
        for order, record in enumerate(records):
            if isinstance(record, dict):
                entries.append({"kind": kind, "kind_label": label, "record": record, "order": order})
    return sorted(entries, key=lambda item: (str(item["record"].get("applied_at") or ""), -int(item["order"])))


def _record_changes(record: dict[str, Any], kind: str) -> list[dict[str, Any]]:
    if kind == "manual_adjustment":
        plan = record.get("plan") if isinstance(record.get("plan"), dict) else {}
        raw_changes = plan.get("changes") if isinstance(plan.get("changes"), list) else []
    else:
        plans = record.get("plans") if isinstance(record.get("plans"), list) else []
        raw_changes = [
            change
            for plan in plans
            if isinstance(plan, dict)
            for change in (plan.get("changes") if isinstance(plan.get("changes"), list) else [])
        ]
    changes: list[dict[str, Any]] = []
    for change in raw_changes:
        if not isinstance(change, dict):
            continue
        changes.append(
            {
                "class_name": str(change.get("class_name") or ""),
                "day": str(change.get("day") or ""),
                "slot": str(change.get("slot") or ""),
                "before": str(change.get("before") or ""),
                "after": str(change.get("after") or ""),
                "reason": str(change.get("reason") or ""),
            }
        )
    return changes


def _record_involved_teachers(summary: dict[str, Any], changes: list[dict[str, Any]]) -> list[str]:
    teachers = [str(item) for item in summary.get("involved_teachers", []) if str(item).strip()] if isinstance(summary.get("involved_teachers"), list) else []
    if not teachers:
        found: set[str] = set()
        for change in changes:
            found.update(_extract_teachers(change.get("before") or ""))
            found.update(_extract_teachers(change.get("after") or ""))
        teachers = sorted(found)
    return list(dict.fromkeys(teachers))


def _record_affected_classes(summary: dict[str, Any], changes: list[dict[str, Any]]) -> list[str]:
    classes = [str(item) for item in summary.get("affected_classes", []) if str(item).strip()] if isinstance(summary.get("affected_classes"), list) else []
    if not classes:
        classes = sorted({str(change.get("class_name") or "") for change in changes if str(change.get("class_name") or "").strip()})
    return list(dict.fromkeys(classes))


def _record_title(record: dict[str, Any], fallback: str) -> str:
    adjustment = record.get("adjustment") if isinstance(record.get("adjustment"), dict) else {}
    leave = record.get("leave") if isinstance(record.get("leave"), dict) else {}
    return str(adjustment.get("title") or leave.get("reason") or fallback)


def _version_label(value: Any) -> str:
    text = str(value or "").replace("\\", "/").strip()
    if not text:
        return ""
    return text.rsplit("/", 1)[-1]


def _resolve_path_text(value: str) -> str:
    if not value:
        return ""
    try:
        return str(Path(value).resolve())
    except OSError:
        return str(value)


def _counter_distribution(
    counter: Counter[str],
    key_name: str,
    *,
    total: int,
    limit: int | None = None,
    order: list[str] | None = None,
    sort_key: Any | None = None,
) -> list[dict[str, Any]]:
    if order:
        pairs = [(key, int(counter.get(key, 0))) for key in order]
    elif sort_key:
        pairs = sorted(counter.items(), key=sort_key)
    else:
        pairs = counter.most_common()
    if limit is not None:
        pairs = pairs[:limit]
    return [
        {
            key_name: str(key),
            "count": int(count),
            "share": round(int(count) / total * 100, 1) if total else 0,
        }
        for key, count in pairs
    ]


def _top_counter_label(counter: Counter[str]) -> str:
    if not counter:
        return ""
    key, count = counter.most_common(1)[0]
    return f"{key} {count}"


def _empty_quality() -> dict[str, Any]:
    return {
        "summary": {
            "status": "empty",
            "status_label": "暂无抽查",
            "classes": 0,
            "teaching_cells": 0,
            "empty_cells": 0,
            "inactive_cells_ignored": 0,
            "inactive_slot_count": 0,
            "inactive_slots": [],
            "filled_cells": 0,
            "teacher_conflicts": 0,
            "teacher_count": 0,
            "max_teacher_load": 0,
            "min_teacher_load": 0,
            "teacher_load_avg": 0,
            "teacher_load_median": 0,
            "teacher_load_spread": 0,
            "warnings": 0,
            "errors": 0,
        },
        "issues": [],
        "teacher_load_top": [],
        "teacher_load_low": [],
        "class_load_top": [],
    }


def _is_teaching_slot(slot: str) -> bool:
    text = str(slot or "").strip()
    if not text:
        return True
    return not any(token in text for token in ("早餐", "午休", "晚餐", "晚休", "休息"))


def _extract_teachers(text: str) -> list[str]:
    teachers = [match.strip() for match in re.findall(r"[（(]([^()（）]+)[）)]", text) if match.strip()]
    if teachers:
        return list(dict.fromkeys(teachers))
    if "-" in text:
        tail = text.rsplit("-", 1)[-1].strip()
        if tail and len(tail) <= 8:
            return [tail]
    return []


def _strip_teacher_marker(text: str, teacher: str) -> str:
    value = re.sub(rf"[（(]\s*{re.escape(teacher)}\s*[）)]", "", text).strip()
    if value.endswith(f"-{teacher}"):
        value = value[: -len(teacher) - 1].strip()
    return value


def _extract_subject(text: str) -> str:
    value = re.sub(r"[（(][^()（）]+[）)]", "", str(text or "")).strip()
    value = re.split(r"[；;,/]", value, maxsplit=1)[0].strip()
    return value or "未识别"


def _is_self_study_cell(text: str) -> bool:
    value = str(text or "").strip()
    return value in {"自习", "自主学习", "空课"} or value.startswith("自习")


def _slot_sort_key(slot: str) -> tuple[int, str]:
    text = str(slot or "")
    order = ["早自习", "上午", "下午", "晚自习"]
    for index, prefix in enumerate(order):
        if text.startswith(prefix):
            match = re.search(r"(\d+)", text)
            return (index, match.group(1).zfill(2) if match else text)
    return (99, text)


def _is_regular_lesson_slot(slot: str) -> bool:
    text = str(slot or "").strip()
    return text.startswith("上午") or text.startswith("下午")


def _slot_phase_distance(left: str, right: str) -> int:
    left_seq = _slot_sequence_key(left)
    right_seq = _slot_sequence_key(right)
    if left_seq is None or right_seq is None:
        return 4
    return abs(int(left_seq[0]) - int(right_seq[0]))


def _slot_number_distance(left: str, right: str) -> int:
    left_seq = _slot_sequence_key(left)
    right_seq = _slot_sequence_key(right)
    if left_seq is None or right_seq is None:
        return 6
    return abs(int(left_seq[1]) - int(right_seq[1]))


def _class_day_balance_penalty(source_day: str, target_day: str, day_counts: Counter[str]) -> int:
    simulated = {str(day): int(count) for day, count in day_counts.items() if int(count) > 0}
    if not simulated:
        return 99
    simulated[source_day] = max(0, int(simulated.get(source_day, 0)) - 1)
    simulated[target_day] = int(simulated.get(target_day, 0)) + 1
    positive = [count for count in simulated.values() if count > 0]
    if not positive:
        return 99
    return max(positive) - min(positive)


def _slot_phase_index(slot: str) -> int:
    seq = _slot_sequence_key(slot)
    if seq is not None:
        return int(seq[0])
    return int(_slot_sort_key(slot)[0])


def _slot_number_index(slot: str) -> int:
    seq = _slot_sequence_key(slot)
    if seq is not None:
        return int(seq[1])
    try:
        return int(str(_slot_sort_key(slot)[1]))
    except ValueError:
        return 0


def _slot_sequence_key(slot: str) -> tuple[int, int] | None:
    text = re.sub(r"\s+", "", str(slot or ""))
    if not text:
        return None
    for phase, prefix in enumerate(("早自习", "上午", "下午", "晚自习")):
        if text.startswith(prefix):
            match = re.search(r"(\d+)", text)
            if match:
                return (phase, int(match.group(1)))
            return (phase, 1) if prefix == "早自习" else None
    generic = re.fullmatch(r"(?:第)?(\d+)(?:节|课)?", text)
    if generic:
        return (10, int(generic.group(1)))
    return None


def _format_slot_run(slots: list[str]) -> str:
    clean = [str(slot or "").strip() for slot in slots if str(slot or "").strip()]
    if not clean:
        return ""
    if len(clean) == 1:
        return clean[0]
    return f"{clean[0]} 至 {clean[-1]}"


def _view(kind: str, name: str, source: str, sheet: str, rows: list[dict[str, Any]]) -> dict[str, Any]:
    non_empty = sum(1 for row in rows for value in (row.get("values") or {}).values() if str(value or "").strip())
    return {
        "kind": kind,
        "name": name,
        "source": source,
        "sheet": sheet,
        "columns": DAYS,
        "rows": rows,
        "non_empty": non_empty,
    }


def _dedupe_views(views: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[tuple[str, str, str]] = set()
    out = []
    for view in views:
        key = (str(view.get("kind") or ""), str(view.get("name") or ""), str(view.get("sheet") or ""))
        if key in seen:
            continue
        seen.add(key)
        out.append(view)
    return out


def _clean_view_name(title: str) -> str:
    text = re.sub(r"课表.*$", "", title).strip()
    return text or title.strip() or "未命名"


def _normalize_day(text: str) -> str:
    value = str(text or "").strip().replace("周", "星期")
    aliases = {"星期天": "星期日"}
    return aliases.get(value, value)


def _cell_text(value: Any, *, max_len: int = 120) -> str:
    text = "" if value is None else str(value).strip()
    text = re.sub(r"\s+", " ", text)
    return text[:max_len]
