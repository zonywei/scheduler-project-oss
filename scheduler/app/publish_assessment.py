# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import load_workbook

from scheduler.calendar import DAY_ORDER, day_rank, ordered_days
from scheduler.solver_quality import solver_quality_fields


SEVERITY_ORDER = {"ok": 0, "info": 1, "warning": 2, "error": 3}


def build_publish_assessment(status: dict[str, Any], readiness: dict[str, Any] | None = None) -> dict[str, Any]:
    status = status if isinstance(status, dict) else {}
    readiness = readiness if isinstance(readiness, dict) else {}
    files = status.get("files") if isinstance(status.get("files"), list) else []
    outcome = _solve_outcome(status, files)
    categories = _category_counts(files, status.get("package_file"))
    gates: list[dict[str, Any]] = []

    gates.append(_run_status_gate(status))
    diagnostic_gate = _diagnostic_trial_gate(status)
    if diagnostic_gate:
        gates.append(diagnostic_gate)
    config_gate = _config_freshness_gate(status)
    if config_gate:
        gates.append(config_gate)
    gates.append(_solution_gate(status, outcome, categories))
    quality_gate = _solver_quality_gate(outcome, status)
    if quality_gate:
        gates.append(quality_gate)
    gates.append(_schedule_file_gate(status, categories))
    gates.append(_diagnostic_file_gate(status, categories))
    checkin_policy_gate = _checkin_policy_gate(status, files)
    if checkin_policy_gate:
        gates.append(checkin_policy_gate)
    gates.append(_package_gate(status, categories))
    readiness_gate = _readiness_gate(readiness)
    if readiness_gate:
        gates.append(readiness_gate)

    gates = [_with_gate_key(gate) for gate in gates if gate]
    summary = _summary(gates)
    return {
        "summary": summary,
        "gates": gates,
        "outcome": outcome,
        "file_categories": categories,
        "next_actions": _next_actions(gates),
    }


def _run_status_gate(status: dict[str, Any]) -> dict[str, Any]:
    value = str(status.get("status") or "idle")
    labels = {
        "idle": "暂无求解结果",
        "starting": "求解启动中",
        "running": "求解运行中",
        "pause_requested": "等待暂停导出",
        "completed": "求解已完成",
        "failed": "求解失败",
        "stopped": "求解已停止",
        "unknown": "求解状态未知",
    }
    if value == "completed":
        return _gate("ok", "求解状态", labels[value], "后台任务正常结束。", "继续检查可行解、诊断材料和发布风险。")
    if value in {"starting", "running", "pause_requested"}:
        return _gate("warning", "求解状态", labels.get(value, value), "任务尚未结束，当前文件可能不是最终交付版本。", "等待求解完成，或暂停并导出当前最优解后再评估。")
    if value == "idle":
        return _gate("warning", "求解状态", labels[value], "还没有可发布的求解批次。", "先完成一次求解。")
    return _gate("error", "求解状态", labels.get(value, value), str(status.get("message") or "求解批次没有正常完成。"), "先处理失败原因并重新求解。", blocking=True)


def _diagnostic_trial_gate(status: dict[str, Any]) -> dict[str, Any] | None:
    if str(status.get("run_purpose") or "") != "diagnostic_trial":
        return None
    title = str(status.get("diagnostic_plan_title") or status.get("diagnostic_plan_id") or "诊断试跑")
    return _gate(
        "error",
        "发布门禁",
        "诊断试跑结果不可发布",
        f"本批次使用“{title}”临时放宽规则，只能用于定位无解来源。",
        "撤销临时规则后，用正式规则重新求解；只有正式批次通过发布评估后才能交付。",
        blocking=True,
    )


def _config_freshness_gate(status: dict[str, Any]) -> dict[str, Any] | None:
    lifecycle_status = str(status.get("status") or "")
    if lifecycle_status not in {"completed", "stopped"}:
        return None

    match = status.get("config_fingerprint_match")
    freshness = status.get("config_freshness") if isinstance(status.get("config_freshness"), dict) else {}
    if match is False or bool(status.get("config_changed_after_run")):
        run_hash = _short_hash(status.get("config_fingerprint"))
        current_hash = _short_hash(status.get("current_config_fingerprint"))
        detail = "当前基础数据或规则配置已不同于本批次求解时的版本。"
        if run_hash or current_hash:
            detail += f" 求解版本 {run_hash or '未知'}，当前版本 {current_hash or '未知'}。"
        return _gate(
            "error",
            "配置版本",
            "当前配置已变更",
            detail,
            "用当前配置重新运行正式求解；旧批次只能留作历史记录，不能直接发布。",
            blocking=True,
        )
    if match is None:
        detail = str(freshness.get("message") or "本批次缺少配置版本记录，无法确认结果是否对应当前规则。")
        return _gate(
            "warning",
            "配置版本",
            "配置版本未校验",
            detail,
            "发布前建议重新运行一次正式求解，以获得可追溯的配置版本记录。",
        )
    return _gate(
        "ok",
        "配置版本",
        "配置版本一致",
        "当前基础数据和规则配置与本批次求解时记录的版本一致。",
        "继续进行课表抽查、诊断报告复核和发布审批。",
    )


def _solution_gate(status: dict[str, Any], outcome: dict[str, Any], categories: dict[str, Any]) -> dict[str, Any]:
    solver_status = str(outcome.get("solver_status") or "")
    solution_count = outcome.get("solution_count")
    if solver_status == "INFEASIBLE":
        return _gate("error", "可行解", "规则组合不可行", "求解器证明当前规则组合不可行。", "回到冲突检测或规则总览，放宽互斥硬约束。", blocking=True)
    if isinstance(solution_count, int):
        if solution_count > 0:
            if _is_diagnostic_trial(status):
                return _gate(
                    "ok",
                    "候选解",
                    f"诊断试跑捕获 {solution_count} 个候选解",
                    "临时放宽规则后可以形成候选方案，说明正式规则的阻断项需要继续排查。",
                    "用候选解定位无解来源，不进入发布流程；修复阻断项后重新正式求解。",
                )
            return _gate("ok", "可行解", f"捕获 {solution_count} 个可行解", "结果批次包含可用于复核的候选课表。", "发布前查看诊断报告和关键班级/教师课表。")
        if str(status.get("status") or "") == "completed" and categories.get("schedule", 0) <= 0:
            return _gate("error", "可行解", "未捕获可行解", "求解完成但没有导出可交付课表。", "延长求解时间、降低硬约束强度，或先处理求解前校验项。", blocking=True)
    if str(status.get("status") or "") == "completed" and categories.get("schedule", 0) > 0:
        return _gate("warning", "可行解", "可行解数量未写入日志", "已发现课表文件，但日志或摘要里没有明确可行解数量。", "下载结果包并人工复核课表文件完整性。")
    return _gate("warning", "可行解", "尚无可行解证据", "当前批次没有明确的可行解数量或课表文件。", "等待求解完成，或查看日志确认是否已找到可行解。")


def _solver_quality_gate(outcome: dict[str, Any], status: dict[str, Any]) -> dict[str, Any] | None:
    if str(outcome.get("solver_status") or "").strip().upper() != "FEASIBLE":
        return None
    quality = solver_quality_fields(
        outcome.get("objective_value", outcome.get("best_objective")),
        outcome.get("best_bound"),
        solver_status=str(outcome.get("solver_status") or ""),
    )
    if str(_first_present(outcome.get("optimality_status"), quality.get("optimality_status")) or "") != "gap_remaining":
        return None
    objective = _first_present(outcome.get("objective_value"), outcome.get("best_objective"))
    bound = outcome.get("best_bound")
    gap = _first_present(outcome.get("objective_gap"), quality.get("objective_gap"))
    gap_percent = _first_present(outcome.get("gap_percent"), quality.get("gap_percent"))
    detail = f"当前求解只证明有可行解，尚未证明全局最优；目标值 {objective}，最优界 {bound}，缺口 {gap}"
    if gap_percent is not None:
        detail += f"，约 {float(gap_percent):.2f}%"
    proof = _stronger_bound_proof(status, current_gap=gap)
    if proof:
        candidate_objective = proof.get("candidate_objective_value")
        if candidate_objective not in (None, "") and str(candidate_objective) != str(objective):
            detail += f"；推荐正式候选目标值 {candidate_objective}"
        detail += (
            f"；同配置复跑提供更强证明界 {proof['proof_best_bound']}，"
            f"按该界重算缺口 {proof['candidate_objective_gap']}"
        )
        if proof.get("candidate_gap_percent") is not None:
            detail += f"，约 {float(proof['candidate_gap_percent']):.2f}%"
    gate = _gate(
        "warning",
        "求解质量",
        "可行解尚未证明全局最优",
        detail + "。",
        "商业发布前建议延长时限、设置更严格 gap 阈值或用固定随机种子复跑；若接受当前候选解，请在学校既有 OA 或线下流程中记录该最优性缺口。",
    )
    gate["decision_points"] = [
        {"label": "目标值", "value": _decision_value(objective), "tone": "warning"},
        {"label": "最优界", "value": _decision_value(bound), "tone": "warning"},
        {"label": "最优性缺口", "value": _gap_decision_value(gap, gap_percent), "tone": "warning"},
        {"label": "建议动作", "value": "延长时限或收紧 gap 阈值复跑；若接受当前候选解，请在既有 OA 或线下流程中记录最优性缺口。", "tone": "warning"},
    ]
    if proof:
        gate["decision_points"].insert(
            3,
            {
                "label": "更强证明",
                "value": (
                    (f"候选目标值 {proof['candidate_objective_value']} / " if proof.get("candidate_objective_value") not in (None, "") and str(proof.get("candidate_objective_value")) != str(objective) else "")
                    + f"最优界 {proof['proof_best_bound']} / "
                    + f"重算缺口 {_gap_decision_value(proof.get('candidate_objective_gap'), proof.get('candidate_gap_percent'))}"
                ),
                "tone": "warning",
            },
        )
    return gate


def _stronger_bound_proof(status: dict[str, Any], *, current_gap: Any) -> dict[str, Any]:
    recommended = status.get("recommended_formal_candidate") if isinstance(status.get("recommended_formal_candidate"), dict) else {}
    evidence = recommended.get("quality_evidence") if isinstance(recommended.get("quality_evidence"), dict) else {}
    summary = evidence.get("summary") if isinstance(evidence.get("summary"), dict) else {}
    if str(summary.get("status") or "") != "bound_improved":
        return {}
    proof_run_id = str(summary.get("proof_run_id") or "").strip()
    proof_best_bound = summary.get("proof_best_bound")
    candidate_gap = summary.get("candidate_objective_gap")
    if not proof_run_id or proof_best_bound in (None, "") or candidate_gap in (None, ""):
        return {}
    current_gap_num = _float_or_none(current_gap)
    candidate_gap_num = _float_or_none(candidate_gap)
    if current_gap_num is not None and candidate_gap_num is not None and candidate_gap_num >= current_gap_num:
        return {}
    evidence_candidate = evidence.get("candidate") if isinstance(evidence.get("candidate"), dict) else {}
    recommended_candidate = recommended.get("candidate") if isinstance(recommended.get("candidate"), dict) else {}
    return {
        "proof_run_id": proof_run_id,
        "proof_best_bound": proof_best_bound,
        "candidate_objective_value": _first_present(
            summary.get("candidate_objective_value"),
            recommended_candidate.get("objective_value"),
            evidence_candidate.get("objective_value"),
        ),
        "candidate_objective_gap": candidate_gap,
        "candidate_gap_percent": summary.get("candidate_gap_percent"),
    }


def _float_or_none(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _decision_value(value: Any) -> str:
    if value is None or value == "":
        return "未记录"
    return str(value)


def _gap_decision_value(gap: Any, gap_percent: Any) -> str:
    if gap is None or gap == "":
        return "未记录"
    if gap_percent is None or gap_percent == "":
        return str(gap)
    try:
        return f"{gap}（约 {float(gap_percent):.2f}%）"
    except (TypeError, ValueError):
        return f"{gap}（{gap_percent}%）"


def _schedule_file_gate(status: dict[str, Any], categories: dict[str, Any]) -> dict[str, Any]:
    count = int(categories.get("schedule", 0) or 0)
    if count:
        if _is_diagnostic_trial(status):
            return _gate(
                "ok",
                "候选课表",
                f"已发现 {count} 个候选课表文件",
                "结果中包含诊断试跑生成的候选课表，只能用于定位无解来源和排障参考。",
                "不要作为正式课表交付；修复阻断项后用正式规则重新求解。",
            )
        return _gate(
            "ok",
            "课表文件",
            f"已发现 {count} 个课表文件",
            "结果中包含可用于下载前抽查和归档的课表类文件。",
            "优先打开班级课表、教师课表和汇总课表抽查，并结合发布评估决定是否交付。",
        )
    return _gate("error", "课表文件", "缺少课表文件", "结果批次没有识别到班级、教师或汇总课表。", "重新求解并确认导出配置，不能只用日志作为交付物。", blocking=True)


def _diagnostic_file_gate(status: dict[str, Any], categories: dict[str, Any]) -> dict[str, Any]:
    count = int(categories.get("diagnostic", 0) or 0)
    if count:
        if _is_diagnostic_trial(status):
            return _gate("ok", "诊断材料", f"已发现 {count} 个诊断材料", "结果中包含解释诊断试跑、约束放宽和候选解差异的材料。", "优先查看多解诊断、约束审计和配置差异，用于确定正式规则要补数据还是放宽。")
        return _gate("ok", "诊断材料", f"已发现 {count} 个诊断材料", "结果中包含约束审计、诊断或风险解释材料。", "下载前至少查看多解诊断和约束审计。")
    return _gate("warning", "诊断材料", "缺少诊断材料", "没有识别到诊断或审计文件，难以解释为什么这份课表可交付。", "重新打包结果，或检查诊断导出开关。")


def _checkin_policy_gate(status: dict[str, Any], files: list[Any]) -> dict[str, Any] | None:
    if str(status.get("mode") or "") not in {"night", "joint"}:
        return None
    policy = _checkin_policy_from_effective_config(files)
    violation_summary = _checkin_same_day_violations(files)
    mode = str(policy.get("mode") or "").strip().lower()
    weight = policy.get("weight")
    if not mode and not violation_summary:
        return None

    if violation_summary:
        count = int(violation_summary.get("count") or 0)
        penalty = float(violation_summary.get("penalty") or 0)
        teachers = str(violation_summary.get("teachers") or "")
        item_label = str(violation_summary.get("item_label") or "")
        solution_context = _checkin_solution_context(files, violation_summary)
        detail = f"本批次存在 {count} 条晚查寝教师当天无晚自习记录，总罚分 {penalty:g}"
        if teachers:
            detail += f"，涉及教师：{teachers}"
        if item_label:
            detail += f"；明细：{item_label}"
        if mode:
            detail += f"；当前策略：{_checkin_policy_label(mode)}"
        if weight not in (None, ""):
            detail += f"，权重 {weight}"
        if solution_context.get("detail"):
            detail += f"。{solution_context['detail']}"
        suggestion = _checkin_violation_suggestion(solution_context, violation_summary)
        gate = _gate(
            "warning",
            "晚查寝策略",
            "存在当天无晚自习查寝安排",
            detail + "。",
            suggestion,
        )
        ledger = solution_context.get("supply_ledger") if isinstance(solution_context.get("supply_ledger"), list) else []
        if ledger:
            gate["checkin_supply_ledger"] = ledger
        if violation_summary.get("source_event_log"):
            gate["evidence_source"] = str(violation_summary.get("source_event_log") or "")
        gate["remediation_options"] = _checkin_violation_remediations(solution_context, violation_summary)
        requirements = _checkin_candidate_data_requirements(solution_context.get("candidate_suggestions"))
        if requirements:
            gate["candidate_data_requirements"] = requirements
        return gate

    if mode == "soft":
        detail = "本批次把“晚查寝教师当天必须有晚自习课”作为软约束处理"
        if weight not in (None, ""):
            detail += f"，权重 {weight}"
        return _gate(
            "warning",
            "晚查寝策略",
            "晚查寝当天有课为软约束",
            detail + "。",
            "下载前查看诊断报告，确认是否存在当天无课查寝安排；如学校要求严格同日有课，应改回硬约束后重新求解。",
        )
    if mode == "off":
        return _gate(
            "warning",
            "晚查寝策略",
            "晚查寝当天有课检查已关闭",
            "本批次没有检查查寝教师当天是否有晚自习课。",
            "只有学校明确允许查寝教师无晚自习也参与值守时才可发布；否则改为硬约束或软约束后重新求解。",
        )
    return None


def _package_gate(status: dict[str, Any], categories: dict[str, Any]) -> dict[str, Any]:
    package = str(status.get("package_file") or "")
    package_exists = bool(package and Path(package).exists())
    count = int(categories.get("package", 0) or 0)
    if package_exists or count:
        if _is_diagnostic_trial(status):
            return _gate(
                "ok",
                "排障包",
                "排障包可用",
                "可一次性归档候选课表、诊断、日志和配置差异，仅用于排障讨论和历史留痕。",
                "修复发布阻断项并重新正式求解后，再下载正式结果包交付。",
            )
        return _gate(
            "ok",
            "结果包",
            "结果包可用",
            "可一次性归档课表、诊断、日志和配置差异。",
            "下载前先查看发布评估；只有可发布批次才能作为正式交付包。",
        )
    return _gate("warning", "结果包", "结果包尚未生成", "文件可逐个下载，但缺少归档包不利于追责和交付。", "点击下载完整结果包生成归档。")


def _readiness_gate(readiness: dict[str, Any]) -> dict[str, Any] | None:
    summary = readiness.get("summary") if isinstance(readiness, dict) else {}
    if not isinstance(summary, dict) or not summary:
        return None
    if summary.get("can_publish") is False:
        errors = int(summary.get("errors") or 0)
        blocking_items = [
            item
            for item in (readiness.get("items") or [])
            if isinstance(item, dict) and item.get("severity") == "error"
        ]
        titles = "；".join(str(item.get("title") or "错误级风险") for item in blocking_items[:3])
        first_suggestion = next((str(item.get("suggestion") or "").strip() for item in blocking_items if str(item.get("suggestion") or "").strip()), "")
        detail = f"当前配置仍有 {errors} 个错误级风险" + (f"：{titles}。" if titles else "。")
        return _gate(
            "error",
            "当前配置",
            "发布校验未通过",
            detail,
            first_suggestion or "先处理求解前校验或教务风险，再发布最新课表。",
            blocking=True,
        )
    if int(summary.get("warnings") or 0) > 0:
        return _gate("warning", "当前配置", "当前配置存在风险提示", str(summary.get("message") or "当前配置仍有警告。"), "下载前查看风险提示，必要时调整配置后重新求解。")
    return _gate("ok", "当前配置", "当前配置可发布", "基础数据、规则冲突和教务风险未发现发布阻断项。", "可以进入课表抽查和结果下载。")


def _solve_outcome(status: dict[str, Any], files: list[Any]) -> dict[str, Any]:
    log = str(status.get("log_tail") or "")
    summary = _read_run_summary(files)
    overview = _read_solver_overview(files)
    solver_status = (
        _first_present(status.get("solver_status"), overview.get("solver_status"))
        or _first_match(log, r"求解状态：([A-Z_]+)")
        or _first_match(log, r"status\s*=\s*([A-Z_]+)")
    )
    solution_raw = _first_present(status.get("solution_count")) or _first_match(log, r"回调捕获解数量：(\d+)") or _first_match(log, r"top_count=(\d+)")
    if solution_raw is None:
        for key in ("top_exported", "kept_solutions", "captured"):
            value = summary.get(key)
            if isinstance(value, int):
                solution_raw = str(value)
                break
    objective_value = _first_present(status.get("objective_value"), overview.get("objective_value"), status.get("best_objective"), summary.get("best_objective"))
    best_bound = _first_present(status.get("best_bound"), overview.get("best_bound"))
    quality = solver_quality_fields(objective_value, best_bound, solver_status=str(solver_status or ""))
    best = _best_objective_label(status.get("best_objective"), overview.get("objective_value"), summary.get("best_objective"), _first_match(log, r"best=([^,\n]+)"))
    return {
        "solver_status": solver_status or "",
        "solution_count": int(solution_raw) if str(solution_raw or "").isdigit() else None,
        "best_objective": best,
        "objective_value": objective_value,
        "best_bound": best_bound,
        "objective_gap": _first_present(status.get("objective_gap"), quality.get("objective_gap")),
        "gap_percent": _first_present(status.get("gap_percent"), quality.get("gap_percent")),
        "optimality_status": _first_present(status.get("optimality_status"), quality.get("optimality_status")),
    }


def _first_present(*values: Any) -> Any:
    for value in values:
        if value is None:
            continue
        if isinstance(value, str) and not value.strip():
            continue
        return value
    return None


def _read_run_summary(files: list[Any]) -> dict[str, Any]:
    for file in files:
        if not isinstance(file, dict):
            continue
        label = str(file.get("label") or file.get("path") or "")
        if not label.replace("\\", "/").endswith("run_summary.json"):
            continue
        path = Path(str(file.get("path") or ""))
        if not path.exists():
            continue
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            return data if isinstance(data, dict) else {}
        except Exception:
            return {}
    return {}


def _read_solver_overview(files: list[Any]) -> dict[str, Any]:
    candidates: list[Path] = []
    for file in files:
        if not isinstance(file, dict):
            continue
        label = str(file.get("label") or file.get("path") or "").replace("\\", "/")
        if not label.endswith("final_solver_overview.json"):
            continue
        path = Path(str(file.get("path") or ""))
        if path.exists() and path.is_file():
            candidates.append(path)
    if not candidates:
        return {}
    latest = max(candidates, key=lambda path: path.stat().st_mtime)
    try:
        data = json.loads(latest.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _checkin_policy_from_effective_config(files: list[Any]) -> dict[str, Any]:
    cfg = _read_effective_config(files)
    checkin = cfg.get("checkin") if isinstance(cfg.get("checkin"), dict) else {}
    raw_mode = checkin.get("require_teacher_has_class_that_day_mode") if isinstance(checkin, dict) else None
    if raw_mode is not None:
        mode = str(raw_mode or "").strip().lower()
        if mode in {"hard", "soft", "off"}:
            return {"mode": mode, "weight": checkin.get("w_require_teacher_has_class_that_day")}
    if isinstance(checkin, dict) and "require_teacher_has_class_that_day" in checkin:
        enabled = _truthy(checkin.get("require_teacher_has_class_that_day"), default=True)
        return {"mode": "hard" if enabled else "off", "weight": checkin.get("w_require_teacher_has_class_that_day")}
    return {}


def _read_effective_config(files: list[Any]) -> dict[str, Any]:
    for file in files:
        if not isinstance(file, dict):
            continue
        label = str(file.get("label") or file.get("path") or "").replace("\\", "/")
        if not label.endswith("effective_config.yaml"):
            continue
        path = Path(str(file.get("path") or ""))
        if not path.exists() or not path.is_file():
            continue
        try:
            data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
            return data if isinstance(data, dict) else {}
        except Exception:
            return {}
    return {}


def _checkin_same_day_violations(files: list[Any]) -> dict[str, Any]:
    event_log_paths = _preferred_checkin_event_log_paths(files)
    best: dict[str, Any] = {}
    for path in event_log_paths:
        if not path.exists() or not path.is_file():
            continue
        try:
            df = pd.read_csv(path, encoding="utf-8-sig")
        except Exception:
            continue
        if "constraint_id" not in df.columns:
            continue
        hits = df[df["constraint_id"].astype(str) == "night_checkin_need_same_day_class"]
        if hits.empty:
            continue
        penalty = float(hits["penalty"].fillna(0).sum()) if "penalty" in hits.columns else 0.0
        teachers = []
        if "teacher_name" in hits.columns:
            teachers = [str(x).strip() for x in hits["teacher_name"].dropna().tolist() if str(x).strip()]
        rows = []
        for row in hits.to_dict("records"):
            rows.append(
                {
                    "teacher": str(row.get("teacher_name") or "").strip(),
                    "day": str(row.get("day") or "").strip(),
                    "slot": str(row.get("period") or "").strip(),
                    "penalty": float(row.get("penalty") or 0),
                }
            )
        current = {
            "count": int(len(hits)),
            "penalty": penalty,
            "teachers": "、".join(sorted(set(teachers))[:8]),
            "rows": rows,
            "source_event_log": str(path),
            "item_label": "；".join(
                f"{row['teacher']}/{row['day']}/{row['slot']}"
                for row in rows[:8]
                if row.get("teacher") or row.get("day") or row.get("slot")
            ),
        }
        if int(current["count"]) > int(best.get("count") or 0):
            best = current
    return best


def _preferred_checkin_event_log_paths(files: list[Any]) -> list[Path]:
    all_logs = _event_log_paths(files)
    preferred: list[Path] = []
    best_meta = _best_solution_meta(files)

    for raw in _event_logs_from_best_meta(best_meta):
        _append_existing_path(preferred, raw)

    seq_id = _meta_seq_id(best_meta)
    if seq_id is not None:
        seq_token = f"sol_{seq_id:04d}"
        snapshot_token = f"seq{seq_id:04d}"
        for path in all_logs:
            text = str(path).replace("\\", "/")
            if seq_token in text or snapshot_token in text:
                _append_existing_path(preferred, path)

    return preferred or all_logs


def _event_log_paths(files: list[Any]) -> list[Path]:
    paths: list[Path] = []
    for file in files:
        if not isinstance(file, dict):
            continue
        label = str(file.get("label") or file.get("path") or "").replace("\\", "/")
        if not label.endswith("event_log.csv"):
            continue
        _append_existing_path(paths, file.get("path"))
    return paths


def _best_solution_meta(files: list[Any]) -> dict[str, Any]:
    candidates: list[tuple[int, Path]] = []
    for file in files:
        if not isinstance(file, dict):
            continue
        label = str(file.get("label") or file.get("path") or "").replace("\\", "/")
        path = Path(str(file.get("path") or ""))
        if not path.exists() or not path.is_file():
            continue
        if not label.endswith("最终全局最优解_meta.json"):
            continue
        score = 0 if "/best/" in label else 10
        candidates.append((score, path))
    for _, path in sorted(candidates, key=lambda item: (item[0], str(item[1]))):
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if isinstance(data, dict):
            return data
    return {}


def _event_logs_from_best_meta(meta: dict[str, Any]) -> list[Path]:
    paths: list[Path] = []
    if not isinstance(meta, dict) or not meta:
        return paths
    _append_existing_path(paths, meta.get("event_log_file"))
    source_solution = str(meta.get("source_solution_file") or "").strip()
    if source_solution:
        _append_existing_path(paths, Path(source_solution).with_name("event_log.csv"))
    return paths


def _meta_seq_id(meta: dict[str, Any]) -> int | None:
    if not isinstance(meta, dict):
        return None
    value = meta.get("seq_id")
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _append_existing_path(paths: list[Path], raw: Any) -> None:
    text = str(raw or "").strip()
    if not text:
        return
    path = Path(text)
    if not path.exists() or not path.is_file():
        return
    if path not in paths:
        paths.append(path)


def _checkin_solution_context(files: list[Any], violation_summary: dict[str, Any]) -> dict[str, Any]:
    rows = violation_summary.get("rows") if isinstance(violation_summary.get("rows"), list) else []
    if not rows:
        return {}
    workbook = _find_best_schedule_workbook(files)
    if workbook is None:
        return {}
    try:
        by_gender = _head_evening_teachers_by_day(workbook)
    except Exception:
        return {}

    missing_days: dict[str, set[str]] = {"男": set(), "女": set()}
    available: dict[str, list[str]] = {}
    for row in rows:
        if not isinstance(row, dict):
            continue
        day = str(row.get("day") or "").strip()
        slot = str(row.get("slot") or "")
        gender = "女" if "女" in slot else "男" if "男" in slot else ""
        if not day or not gender:
            continue
        names = sorted(by_gender.get(gender, {}).get(day, set()))
        available[f"{gender}:{day}"] = names
        if not names:
            missing_days[gender].add(day)

    details: list[str] = []
    for gender in ("男", "女"):
        days = _ordered_days(missing_days.get(gender, set()))
        if days:
            details.append(f"当前班主任课表中，{'、'.join(days)}没有{gender}班主任晚自习可直接承接{gender}寝晚查")
    candidate_suggestions = _checkin_candidate_suggestions(files, workbook, missing_days)
    supply_ledger = _checkin_supply_ledger(rows, by_gender, candidate_suggestions)
    if candidate_suggestions:
        ready = [item for item in candidate_suggestions if item.get("status") == "ready_to_add"]
        pending = [item for item in candidate_suggestions if item.get("status") == "needs_confirmation"]
        if ready:
            details.append(
                "系统找到可优先补充的同日晚自习候选："
                + "；".join(f"{item.get('day')}{item.get('required_gender')}寝 {item.get('teacher')}" for item in ready[:4])
            )
        elif pending:
            examples = "；".join(f"{item.get('day')} {item.get('teacher')}" for item in pending[:6])
            details.append(
                f"同日晚自习表中有待确认名单：{examples}；这些不是已确认可用人力，"
                "当前问题应按候选供给不足处理，优先引入可承担查寝的新教师、宿管或行政值守人员"
            )
    if not details:
        examples = []
        for key, names in available.items():
            if names:
                gender, day = key.split(":", 1)
                examples.append(f"{day}{gender}候选：{'、'.join(names[:4])}")
        if examples:
            details.append("当前结果存在同日晚自习候选，可优先尝试人工换班：" + "；".join(examples[:3]))
    return {
        "detail": "；".join(details),
        "missing_days_by_gender": {gender: _ordered_days(days) for gender, days in missing_days.items() if days},
        "available": available,
        "candidate_suggestions": candidate_suggestions,
        "supply_ledger": supply_ledger,
    }


def _find_best_schedule_workbook(files: list[Any]) -> Path | None:
    candidates: list[tuple[int, Path]] = []
    for file in files:
        if not isinstance(file, dict):
            continue
        label = str(file.get("label") or file.get("path") or "").replace("\\", "/")
        path = Path(str(file.get("path") or ""))
        if not path.exists() or not path.is_file() or path.suffix.lower() not in {".xlsx", ".xlsm"}:
            continue
        if "诊断" in label or "diagnostic" in label.lower() or "audit" in label.lower():
            continue
        score = 50
        if label.endswith("/best/最终全局最优解.xlsx") or label.endswith("\\best\\最终全局最优解.xlsx"):
            score = 0
        elif "最终全局最优解" in label:
            score = 5
        elif "课表" in label:
            score = 20
        candidates.append((score, path))
    if not candidates:
        return None
    return sorted(candidates, key=lambda item: (item[0], str(item[1])))[0][1]


def _head_evening_teachers_by_day(workbook_path: Path) -> dict[str, dict[str, set[str]]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    if "班主任课表" not in wb.sheetnames:
        return {}
    ws = wb["班主任课表"]
    by_gender: dict[str, dict[str, set[str]]] = {"男": {}, "女": {}}
    current_gender = ""
    day_columns: dict[int, str] = {}
    for row in ws.iter_rows(values_only=True):
        first = str(row[0] or "").strip() if row else ""
        if first == "男班主任课表":
            current_gender = "男"
            day_columns = {}
            continue
        if first == "女班主任课表":
            current_gender = "女"
            day_columns = {}
            continue
        if first == "节次":
            day_columns = {
                idx: str(value).strip()
                for idx, value in enumerate(row)
                if idx > 0 and str(value or "").strip().startswith("星期")
            }
            continue
        if current_gender and first in {"晚自习1", "晚自习2"}:
            for idx, day in day_columns.items():
                names = _split_teacher_names(row[idx] if idx < len(row) else "")
                if names:
                    by_gender.setdefault(current_gender, {}).setdefault(day, set()).update(names)
    return by_gender


def _checkin_candidate_suggestions(
    files: list[Any],
    workbook_path: Path,
    missing_days: dict[str, set[str]],
) -> list[dict[str, str]]:
    if not any(missing_days.get(gender) for gender in ("男", "女")):
        return []
    evening_by_day = _night_evening_teachers_by_day(workbook_path)
    if not evening_by_day:
        return []
    gender_map = _teacher_gender_map(files)
    suggestions: list[dict[str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    for required_gender in ("男", "女"):
        for day in _ordered_days(missing_days.get(required_gender, set())):
            candidates: list[tuple[int, str, str]] = []
            for teacher in sorted(evening_by_day.get(day, set())):
                known_gender = gender_map.get(teacher, "")
                if known_gender and known_gender != required_gender:
                    continue
                status = "ready_to_add" if known_gender == required_gender else "needs_confirmation"
                rank = 0 if status == "ready_to_add" else 1
                candidates.append((rank, teacher, status))
            for _, teacher, status in sorted(candidates, key=lambda item: (item[0], item[1]))[:6]:
                key = (teacher, day, required_gender)
                if key in seen:
                    continue
                seen.add(key)
                known_gender = gender_map.get(teacher, "")
                if status == "ready_to_add":
                    reason = f"{teacher}在{day}已有晚自习，教师定位表性别为{required_gender}，可由教务确认后加入额外晚查寝候选名单。"
                    confidence = "高"
                else:
                    reason = f"{teacher}在{day}已有晚自习；教师定位表未记录可用于查寝的性别信息，若确认为{required_gender}且具备值守资格，可加入额外晚查寝候选名单。"
                    confidence = "待确认"
                suggestions.append(
                    {
                        "teacher": teacher,
                        "day": day,
                        "required_gender": required_gender,
                        "known_gender": known_gender,
                        "status": status,
                        "confidence": confidence,
                        "reason": reason,
                        "config_hint": json.dumps(
                            {"name": teacher, "gender": required_gender, "days": [day]},
                            ensure_ascii=False,
                        ),
                    }
                )
    return suggestions[:12]


def _checkin_supply_ledger(
    violation_rows: list[Any],
    by_gender: dict[str, dict[str, set[str]]],
    candidate_suggestions: list[dict[str, str]],
) -> list[dict[str, Any]]:
    grouped: dict[tuple[str, str], dict[str, Any]] = {}
    for row in violation_rows:
        if not isinstance(row, dict):
            continue
        day = str(row.get("day") or "").strip()
        slot = str(row.get("slot") or "")
        gender = "女" if "女" in slot else "男" if "男" in slot else ""
        if not day or not gender:
            continue
        item = grouped.setdefault(
            (day, gender),
            {
                "day": day,
                "gender": gender,
                "exception_count": 0,
                "assigned_teachers": set(),
                "slots": set(),
                "penalty_sum": 0.0,
            },
        )
        item["exception_count"] += 1
        item["penalty_sum"] += float(row.get("penalty") or 0)
        teacher = str(row.get("teacher") or "").strip()
        if teacher:
            item["assigned_teachers"].add(teacher)
        if slot:
            item["slots"].add(slot)

    suggestions_by_key: dict[tuple[str, str], list[dict[str, str]]] = {}
    for suggestion in candidate_suggestions:
        if not isinstance(suggestion, dict):
            continue
        key = (
            str(suggestion.get("day") or "").strip(),
            str(suggestion.get("required_gender") or "").strip(),
        )
        if key[0] and key[1]:
            suggestions_by_key.setdefault(key, []).append(suggestion)

    ledger: list[dict[str, Any]] = []
    for day, gender in sorted(grouped, key=lambda key: (_day_order_index(key[0]), key[1])):
        item = grouped[(day, gender)]
        direct_candidates = sorted(by_gender.get(gender, {}).get(day, set()))
        suggestions = suggestions_by_key.get((day, gender), [])
        ready = [x for x in suggestions if x.get("status") == "ready_to_add"]
        pending = [x for x in suggestions if x.get("status") == "needs_confirmation"]
        exception_count = int(item.get("exception_count") or 0)
        direct_count = len(direct_candidates)
        shortage = max(0, exception_count - direct_count)
        if direct_count >= exception_count and exception_count:
            tone = "ok"
            status = "可调整同日班主任"
        elif ready:
            tone = "warning"
            status = "有已确认候选可补充"
        elif pending:
            tone = "error"
            status = "现有人力不足，需引入新候选"
        else:
            tone = "error"
            status = "现有人力不足，需引入新候选"
        ledger.append(
            {
                "day": day,
                "gender": gender,
                "duty": f"{gender}寝晚查",
                "exception_count": exception_count,
                "penalty_sum": round(float(item.get("penalty_sum") or 0.0), 2),
                "potential_penalty_reduction": round(float(item.get("penalty_sum") or 0.0), 2),
                "assigned_teachers": sorted(item.get("assigned_teachers") or []),
                "slots": sorted(item.get("slots") or []),
                "direct_candidate_count": direct_count,
                "direct_candidates": direct_candidates[:8],
                "suggested_candidate_count": len(suggestions),
                "ready_candidate_count": len(ready),
                "pending_candidate_count": len(pending),
                "suggested_candidates": [str(x.get("teacher") or "") for x in suggestions[:8] if str(x.get("teacher") or "").strip()],
                "shortage": shortage,
                "status": status,
                "tone": tone,
                "capacity_note": (
                    "待确认名单不能视为已确认可用人力；建议引入可负责查寝的新教师、宿管或行政值守人员。"
                    if shortage
                    else ""
                ),
            }
        )
    return ledger


def _night_evening_teachers_by_day(workbook_path: Path) -> dict[str, set[str]]:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    if "晚自习课表" not in wb.sheetnames:
        return {}
    ws = wb["晚自习课表"]
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        return {}
    day_columns = {
        idx: _normalize_day_label(str(value or "").strip())
        for idx, value in enumerate(header)
        if idx > 0 and _normalize_day_label(str(value or "").strip())
    }
    by_day: dict[str, set[str]] = {}
    for row in rows:
        for idx, day in day_columns.items():
            if idx >= len(row):
                continue
            for teacher in _teachers_from_evening_cell(row[idx]):
                by_day.setdefault(day, set()).add(teacher)
    return by_day


def _teachers_from_evening_cell(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    teachers: list[str] = []
    for part in re.split(r"[\n；;]+", text):
        piece = part.strip()
        if not piece:
            continue
        teacher = piece.rsplit("-", 1)[-1].strip() if "-" in piece else piece
        teachers.extend(_split_teacher_names(teacher))
    return teachers


def _teacher_gender_map(files: list[Any]) -> dict[str, str]:
    cfg = _read_effective_config(files)
    gender_map: dict[str, str] = {}
    teacher_table = cfg.get("teacher_table") if isinstance(cfg.get("teacher_table"), dict) else {}
    columns = teacher_table.get("columns") if isinstance(teacher_table.get("columns"), dict) else {}
    head_col = str(columns.get("head") or "班主任")
    gender_col = str(columns.get("head_gender") or "班主任性别")
    rows = (((cfg.get("web_tables") or {}) if isinstance(cfg.get("web_tables"), dict) else {}).get("teacher_subjects") or [])
    if isinstance(rows, list) and rows:
        for row in rows:
            if isinstance(row, dict):
                _add_teacher_gender(gender_map, row.get(head_col), row.get(gender_col))
    else:
        path = _existing_path(teacher_table.get("path"))
        if path is not None:
            try:
                df = pd.read_excel(path, sheet_name=teacher_table.get("sheet_name", 0))
                for row in df.fillna("").to_dict("records"):
                    _add_teacher_gender(gender_map, row.get(head_col), row.get(gender_col))
            except Exception:
                pass

    checkin = cfg.get("checkin") if isinstance(cfg.get("checkin"), dict) else {}
    extra_heads = checkin.get("extra_heads") if isinstance(checkin.get("extra_heads"), list) else []
    for item in extra_heads:
        if isinstance(item, dict):
            _add_teacher_gender(gender_map, item.get("name"), item.get("gender"))
    return gender_map


def _add_teacher_gender(gender_map: dict[str, str], name_raw: Any, gender_raw: Any) -> None:
    name = str(name_raw or "").strip()
    gender = str(gender_raw or "").strip()
    if name and gender in {"男", "女"}:
        gender_map[name] = gender


def _existing_path(raw: Any) -> Path | None:
    text = str(raw or "").strip()
    if not text:
        return None
    path = Path(text)
    if path.exists():
        return path
    for base in (Path.cwd(), Path.cwd() / "scheduler"):
        candidate = base / text
        if candidate.exists():
            return candidate
    return None


def _normalize_day_label(value: str) -> str:
    text = re.sub(r"\.\d+$", "", value.strip())
    aliases = {
        "周一": "星期一",
        "周二": "星期二",
        "周三": "星期三",
        "周四": "星期四",
        "周五": "星期五",
        "周六": "星期六",
        "周日": "星期日",
        "星期天": "星期日",
    }
    return aliases.get(text, text if text.startswith("星期") else "")


def _split_teacher_names(value: Any) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    return [part.strip() for part in re.split(r"[，,、;；\s]+", text) if part.strip()]


def _ordered_days(days: set[str] | list[str]) -> list[str]:
    day_set = set(days)
    known = ordered_days(day_set)
    return known + sorted(day for day in day_set if day not in known)


def _day_order_index(day: str) -> int:
    return day_rank(day, default=len(DAY_ORDER))


def _checkin_violation_suggestion(solution_context: dict[str, Any], violation_summary: dict[str, Any]) -> str:
    missing = solution_context.get("missing_days_by_gender") if isinstance(solution_context, dict) else {}
    if isinstance(missing, dict) and missing:
        parts = []
        for gender in ("男", "女"):
            days = missing.get(gender) if isinstance(missing.get(gender), list) else []
            if days:
                parts.append(f"{'、'.join(days)}{gender}寝晚查")
        if parts:
            current_label = _checkin_current_assignment_label(violation_summary.get("rows"))
            current_text = f"当前{current_label}安排" if current_label else "当前例外安排"
            return (
                f"这是现有可用人力不足，不建议继续围绕{current_text}反复优化；"
                "优先为"
                + "、".join(parts)
                + "引入可负责查寝的新教师、宿管或行政值守人员，确认性别、资格和值守口径后写入额外候选，再重新求解。"
            )
    if solution_context.get("detail"):
        return "优先按同日晚自习候选做人工换班；若保持当前安排，则由教务在发布前确认专门返校、宿管代查或人工调整方案。"
    return "发布前由教务确认这些教师是否可接受专门返校、宿管代查或人工调整；确认后再交付正式课表。"


def _checkin_violation_remediations(solution_context: dict[str, Any], violation_summary: dict[str, Any]) -> list[dict[str, Any]]:
    rows = violation_summary.get("rows") if isinstance(violation_summary.get("rows"), list) else []
    missing = solution_context.get("missing_days_by_gender") if isinstance(solution_context, dict) else {}
    count = int(violation_summary.get("count") or len(rows) or 0)
    penalty = float(violation_summary.get("penalty") or 0.0)
    item_label = str(violation_summary.get("item_label") or "")
    teachers = str(violation_summary.get("teachers") or "")
    current_label = _checkin_current_assignment_label(rows)
    missing_label = _checkin_missing_label(missing)
    issue_label = missing_label or item_label or "当前晚查寝例外"
    candidate_suggestions = (
        solution_context.get("candidate_suggestions") if isinstance(solution_context.get("candidate_suggestions"), list) else []
    )
    candidate_hint = _checkin_candidate_hint(candidate_suggestions)
    candidate_detail = (
        f"为{issue_label}引入可负责查寝的新教师、宿管或行政值守人员；确认后写入额外候选并重新运行联合求解。"
        if missing_label
        else "核对 event_log 中的晚查寝例外，补充已确认可用的同日候选或改派到已有同日候选后重新求解。"
    )
    options: list[dict[str, Any]] = [
        {
            "id": "publish.checkin.add_same_day_candidate",
            "title": "引入新晚查候选",
            "detail": candidate_detail,
            "risk": "这是候选供给不足问题；待确认名单不能直接视为可用人力，新增候选需确认性别、值守资格、周次数上限和到校事实。",
            "action_type": "manual",
            "manual_hint": (
                "优先引入已确认可负责查寝的新教师、宿管或行政值守人员，并在规则总览维护额外晚查寝候选名单；"
                f"后续不必继续围绕{current_label or issue_label}当前安排反复优化。"
                + (f" 系统候选提示：{candidate_hint}" if candidate_hint else "")
            ),
            "candidate_suggestions": candidate_suggestions,
            "quick_actions": [
                {
                    "id": "open-extra-heads-rule",
                    "label": "维护额外候选人",
                    "action": "open_rule_editor",
                    "view": "rules",
                    "rule_id": "duty.night_checkin.extra_heads",
                    "search": "额外晚查寝候选人",
                },
                {"id": "open-teachers", "label": "打开教师定位", "action": "activate_view", "view": "teachers"},
                {"id": "open-solve", "label": "重新求解", "action": "activate_view", "view": "solve"},
            ],
            "decision_points": _checkin_remediation_decision_points(count, teachers, missing_label, item_label, penalty),
            "checklist": [
                "确认这是可用人力不足，不把待确认名单当作已确认候选。",
                "引入可承担对应寝室查寝的新教师、宿管或行政值守人员。",
                "确认候选人性别、查寝资格、到校事实与寝室类型匹配。",
                "确认补充后不突破个人周值班次数、禁排和公平性规则。",
                "重新运行联合求解并检查发布评估是否清除该复核项。",
            ],
        },
        {
            "id": "publish.checkin.adjust_evening_assignment",
            "title": "调整班主任晚自习",
            "detail": f"把{issue_label}对应日期的晚自习安排调整给可承担查寝的班主任，使晚自习与晚查寝发生在同一天。",
            "risk": "调晚自习会影响班级看班、教师个人禁排和跨日公平性，需要同步复核课表与值班表。",
            "action_type": "manual",
            "manual_hint": "在教务工作台或源数据中调整班主任晚自习后重新求解，避免只手工改最终 Excel。",
            "quick_actions": [
                {"id": "open-academic", "label": "打开教务工作台", "action": "activate_view", "view": "academic"},
                {"id": "open-rules", "label": "查看规则", "action": "activate_view", "view": "rules"},
            ],
            "decision_points": _checkin_remediation_decision_points(count, teachers, missing_label, item_label, penalty),
            "checklist": [
                "先确认需要调整的是星期几、男寝还是女寝。",
                "选择同日可承担查寝且无禁排冲突的班主任。",
                "调整后重新求解，不直接发布手工局部改表。",
                "复核班级晚自习覆盖、查寝覆盖和个人周次数。",
            ],
        },
        {
            "id": "publish.checkin.record_external_exception",
            "title": "外部记录保留例外",
            "detail": f"若学校允许专门返校、宿管代查或行政兜底，可保留这 {count} 条例外，并在学校既有 OA 或线下流程中记录处理口径。",
            "risk": "保留例外意味着最终课表与严格同日有晚自习规则不一致，发布后可能产生到校责任和考勤解释风险。",
            "action_type": "external_record",
            "manual_hint": "系统不再提供内置确认表单；请在学校既有 OA 或线下流程中记录责任人、替代方式、考勤口径和说明。",
            "quick_actions": [
                {"id": "open-results", "label": "回到结果页", "action": "activate_view", "view": "results"},
            ],
            "decision_points": _checkin_remediation_decision_points(count, teachers, missing_label, item_label, penalty),
            "checklist": [
                "确认学校制度允许非同日晚自习教师承担晚查寝。",
                "确认具体责任人、替代方式和考勤口径。",
                "在学校既有 OA 或线下流程中记录例外原因和人工处理方案。",
                "归档 event_log 与结果包，便于后续追溯。",
            ],
        },
    ]
    return options


def _checkin_candidate_hint(candidates: Any) -> str:
    if not isinstance(candidates, list) or not candidates:
        return ""
    ready = [item for item in candidates if isinstance(item, dict) and item.get("status") == "ready_to_add"]
    pending = [item for item in candidates if isinstance(item, dict) and item.get("status") == "needs_confirmation"]
    source = ready or pending
    names = []
    for item in source[:6]:
        teacher = str(item.get("teacher") or "").strip()
        day = str(item.get("day") or "").strip()
        gender = str(item.get("required_gender") or "").strip()
        if teacher:
            names.append(f"{teacher}({day}{gender})" if day or gender else teacher)
    if not names:
        return ""
    suffix = "可优先确认并补充。" if ready else "需先确认性别、查寝资格和学校值守口径。"
    return "、".join(names) + suffix


def _checkin_candidate_data_requirements(candidates: Any) -> list[dict[str, Any]]:
    if not isinstance(candidates, list):
        return []
    rows: list[dict[str, Any]] = []
    for item in candidates:
        if not isinstance(item, dict) or item.get("status") != "needs_confirmation":
            continue
        teacher = str(item.get("teacher") or "").strip()
        day = str(item.get("day") or "").strip()
        required_gender = str(item.get("required_gender") or "").strip()
        known_gender = str(item.get("known_gender") or "").strip()
        if not teacher:
            continue
        missing_fields = []
        if required_gender in {"男", "女"} and known_gender not in {"男", "女"}:
            missing_fields.append("性别")
        missing_fields.extend(["查寝资格", "值守口径"])
        rows.append(
            {
                "teacher": teacher,
                "day": day,
                "required_gender": required_gender,
                "known_gender": known_gender,
                "missing_fields": missing_fields,
                "source": "教师定位表或人工确认材料",
                "action": "补齐后再确认是否加入额外晚查寝候选名单",
            }
        )
    return rows[:12]


def _checkin_missing_label(missing: Any) -> str:
    if not isinstance(missing, dict):
        return ""
    parts: list[str] = []
    for gender in ("男", "女"):
        days = missing.get(gender) if isinstance(missing.get(gender), list) else []
        if days:
            parts.append(f"{'、'.join(str(day) for day in days)}{gender}寝")
    return "、".join(parts)


def _checkin_current_assignment_label(rows: Any) -> str:
    if not isinstance(rows, list):
        return ""
    labels: list[str] = []
    for row in rows[:6]:
        if not isinstance(row, dict):
            continue
        teacher = str(row.get("teacher") or "").strip()
        day = str(row.get("day") or "").strip()
        slot = str(row.get("slot") or "").strip()
        parts = [part for part in (teacher, day, slot) if part]
        if parts:
            labels.append("/".join(parts))
    return "、".join(labels)


def _checkin_remediation_decision_points(
    count: int,
    teachers: str,
    missing_label: str,
    item_label: str,
    penalty: float = 0.0,
) -> list[dict[str, str]]:
    points = [{"label": "例外数量", "value": f"{count} 条", "tone": "warning"}]
    if penalty > 0:
        points.append({"label": "可减少罚分", "value": f"{penalty:g}", "tone": "warning"})
    if teachers:
        points.append({"label": "涉及教师", "value": teachers, "tone": "warning"})
    if missing_label:
        points.append({"label": "候选缺口", "value": missing_label, "tone": "error"})
    elif item_label:
        points.append({"label": "例外明细", "value": item_label, "tone": "warning"})
    return points


def _truthy(value: Any, *, default: bool) -> bool:
    if value is None:
        return default
    if isinstance(value, str):
        return value.strip().lower() not in {"0", "false", "no", "off", "否", "关闭"}
    return bool(value)


def _checkin_policy_label(mode: str) -> str:
    return {"hard": "硬约束", "soft": "软约束", "off": "关闭"}.get(str(mode or "").lower(), str(mode or "未设置"))


def _category_counts(files: list[Any], package_file: Any = None) -> dict[str, Any]:
    counts = {"schedule": 0, "diagnostic": 0, "package": 0, "log": 0, "other": 0}
    for file in files:
        if not isinstance(file, dict):
            continue
        counts[_categorize_file(str(file.get("label") or file.get("path") or ""))] += 1
    if package_file and counts["package"] == 0:
        counts["package"] = 1
    counts["total"] = sum(counts.values())
    return counts


def _categorize_file(source: str) -> str:
    normalized = source.replace("\\", "/")
    filename = normalized.rsplit("/", 1)[-1].lower()
    relative = normalized.lower()
    is_sheet = filename.endswith((".xlsx", ".xls", ".csv"))
    if filename.endswith(".zip") or "result_package" in filename:
        return "package"
    if any(token in relative for token in ("诊断", "审计", "diagnostic", "audit", "checklist", "penalty", "conflict", "violation")):
        return "diagnostic"
    if (
        "课表" in relative
        or (is_sheet and any(token in relative for token in ("最优解", "全局最优", "正式版")))
        or (is_sheet and any(token in filename for token in ("schedule", "timetable", "teacher", "class")))
    ):
        return "schedule"
    if filename.endswith((".log", ".json", ".txt")) or any(token in filename for token in ("meta", "summary", "config_diff")):
        return "log"
    return "other"


def _is_diagnostic_trial(status: dict[str, Any]) -> bool:
    return str(status.get("run_purpose") or "") == "diagnostic_trial"


def _summary(gates: list[dict[str, Any]]) -> dict[str, Any]:
    errors = sum(1 for gate in gates if gate.get("severity") == "error")
    warnings = sum(1 for gate in gates if gate.get("severity") == "warning")
    blocking = sum(1 for gate in gates if gate.get("blocking"))
    if errors:
        status = "blocked"
        label = "不可发布"
        message = f"有 {errors} 个发布阻断项。"
    elif warnings:
        status = "ready"
        label = "可发布（含风险提示）"
        message = f"没有硬阻断；保留 {warnings} 个风险提示供下载前查看。"
    else:
        status = "ready"
        label = "可进入发布"
        message = "结果文件、诊断材料和当前配置均未发现发布阻断项。"
    return {
        "status": status,
        "status_label": label,
        "message": message,
        "errors": errors,
        "warnings": warnings,
        "blocking": blocking,
        "can_publish": errors == 0,
    }


def _next_actions(gates: list[dict[str, Any]]) -> list[dict[str, str]]:
    ordered = sorted(gates, key=lambda gate: -SEVERITY_ORDER.get(str(gate.get("severity") or ""), 0))
    return [
        {
            "title": str(gate.get("title") or ""),
            "domain": str(gate.get("domain") or ""),
            "severity": str(gate.get("severity") or ""),
            "suggestion": str(gate.get("suggestion") or ""),
        }
        for gate in ordered
        if gate.get("severity") in {"error", "warning"}
    ][:5]


def _gate(
    severity: str,
    domain: str,
    title: str,
    detail: str,
    suggestion: str,
    *,
    blocking: bool = False,
) -> dict[str, Any]:
    return {
        "severity": severity,
        "domain": domain,
        "title": title,
        "detail": detail,
        "suggestion": suggestion,
        "blocking": bool(blocking),
    }


def _with_gate_key(gate: dict[str, Any]) -> dict[str, Any]:
    gate = dict(gate)
    gate["key"] = _gate_key(gate)
    return gate


def _gate_key(gate: dict[str, Any]) -> str:
    domain = str(gate.get("domain") or "").strip()
    title = str(gate.get("title") or "").strip()
    return f"{domain}|{title}"


def _first_match(text: str, pattern: str) -> str | None:
    match = re.search(pattern, text)
    return match.group(1).strip() if match else None


def _short_hash(value: Any) -> str:
    if isinstance(value, dict):
        raw = str(value.get("hash") or "")
    else:
        raw = str(value or "")
    return raw[:12] if raw else ""


def _best_objective_label(*candidates: Any) -> str:
    for candidate in candidates:
        text = str(candidate if candidate is not None else "").strip()
        if not text:
            continue
        lower = text.replace("\\", "/").lower()
        if lower.endswith((".xlsx", ".xls", ".csv", ".json", ".zip")) or "/" in lower or ":" in text:
            continue
        return text
    return ""
