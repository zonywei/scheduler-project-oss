# -*- coding: utf-8 -*-
from __future__ import annotations

import copy
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from scheduler.app.config_service import PROJECT_ROOT
from scheduler.output.constants import DAYS


SCHEDULE_REPAIR_SCHEMA = "scheduler.leave_log_local_timetable_repair.v1"
LOCAL_REPAIR_DIR_NAME = "local_repairs"
MAX_TARGET_LESSONS = 12
MAX_CANDIDATES = 16


class LocalTimetableRepairError(ValueError):
    pass


@dataclass(frozen=True)
class CellLocation:
    sheet: str
    row: int
    col: int
    role: str


@dataclass
class ScheduleCell:
    key: str
    class_name: str
    day: str
    slot: str
    value: str
    locations: list[CellLocation]


def preview_leave_substitution_repair(
    *,
    status: dict[str, Any],
    academic_payload: dict[str, Any],
    teacher_rows: list[dict[str, Any]],
    request: dict[str, Any] | None = None,
    project_root: Path = PROJECT_ROOT,
) -> dict[str, Any]:
    return _build_leave_repair(
        status=status,
        academic_payload=academic_payload,
        teacher_rows=teacher_rows,
        request=request or {},
        project_root=project_root,
        apply_changes=False,
    )


def apply_leave_substitution_repair(
    *,
    status: dict[str, Any],
    academic_payload: dict[str, Any],
    teacher_rows: list[dict[str, Any]],
    request: dict[str, Any] | None = None,
    output_dir: Path | None = None,
    project_root: Path = PROJECT_ROOT,
) -> dict[str, Any]:
    return _build_leave_repair(
        status=status,
        academic_payload=academic_payload,
        teacher_rows=teacher_rows,
        request=request or {},
        project_root=project_root,
        output_dir=output_dir,
        apply_changes=True,
    )


def _build_leave_repair(
    *,
    status: dict[str, Any],
    academic_payload: dict[str, Any],
    teacher_rows: list[dict[str, Any]],
    request: dict[str, Any],
    project_root: Path,
    apply_changes: bool,
    output_dir: Path | None = None,
) -> dict[str, Any]:
    workbook_path = _select_schedule_workbook(status, request, project_root)
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        cells = _scan_schedule_cells(workbook)
        if not cells:
            raise LocalTimetableRepairError("未在当前课表中识别到可编辑的班级课表网格。")
        leave = _resolve_leave_request(academic_payload, request)
        plans, candidate_options, after_values = _plan_leave_repair(cells, leave, teacher_rows, request)
        payload = _repair_payload(workbook_path, leave, cells, plans, candidate_options, after_values)
        if apply_changes:
            if not plans:
                raise LocalTimetableRepairError(payload["summary"]["message"])
            output_path = _write_adjusted_workbook(workbook, workbook_path, cells, plans, output_dir)
            payload["applied"] = True
            payload["output_file"] = str(output_path)
            payload["record"] = _repair_record(payload, output_path)
        return payload
    finally:
        workbook.close()


def _select_schedule_workbook(status: dict[str, Any], request: dict[str, Any], project_root: Path) -> Path:
    explicit = str(request.get("workbook_path") or "").strip()
    if explicit:
        path = Path(explicit).resolve()
        _assert_project_file(path, project_root)
        if not path.exists():
            raise LocalTimetableRepairError(f"指定课表文件不存在：{path}")
        return path

    candidates: list[dict[str, Any]] = []
    for key, label in (
        ("local_timetable_adjustment", "active_manual_adjustment"),
        ("local_timetable_repair", "active_local_repair"),
    ):
        local = status.get(key) if isinstance(status.get(key), dict) else {}
        active = str(local.get("active_file") or "").strip()
        if active:
            candidates.append({"path": active, "label": label})
    availability = status.get("result_availability") if isinstance(status.get("result_availability"), dict) else {}
    primary = availability.get("primary_schedule_files") if isinstance(availability.get("primary_schedule_files"), list) else []
    candidates.extend(item for item in primary if isinstance(item, dict))
    files = status.get("files") if isinstance(status.get("files"), list) else []
    candidates.extend(item for item in files if isinstance(item, dict))

    schedule_paths: dict[str, Path] = {}
    for item in candidates:
        raw = str(item.get("path") or "").strip()
        if not raw:
            continue
        path = Path(raw).resolve()
        if not _is_schedule_workbook(path):
            continue
        try:
            _assert_project_file(path, project_root)
        except LocalTimetableRepairError:
            continue
        if path.exists() and path.is_file():
            schedule_paths[str(path)] = path

    ranked = sorted(schedule_paths.values(), key=_workbook_rank)
    if not ranked:
        raise LocalTimetableRepairError("未找到可用于局部调课的课表文件。请先完成一次可行求解并生成全局课表。")
    return ranked[0]


def _assert_project_file(path: Path, project_root: Path) -> None:
    root = project_root.resolve()
    target = path.resolve()
    if target != root and root not in target.parents:
        raise LocalTimetableRepairError("课表文件必须位于当前项目目录内。")


def _is_schedule_workbook(path: Path) -> bool:
    source = str(path).replace("\\", "/").lower()
    name = path.name.lower()
    if not name.endswith((".xlsx", ".xlsm")):
        return False
    if any(token in source for token in ("诊断", "审计", "diagnostic", "audit", "penalty", "conflict", "violation")):
        return False
    return "课表" in source or "最优解" in source or any(token in name for token in ("schedule", "timetable", "class", "teacher"))


def _workbook_rank(path: Path) -> tuple[int, int, str]:
    source = str(path).replace("\\", "/")
    lower = source.lower()
    rank = 50
    if "课表微调" in source or "微调" in source:
        rank -= 60
    if "局部调课" in source or "调整后" in source:
        rank -= 50
    if "/best/" in lower or "最终" in source or "最优解" in source:
        rank -= 30
    if "正式版" in source:
        rank -= 20
    if "top50" in lower:
        rank += 20
    try:
        size = int(path.stat().st_size)
    except OSError:
        size = 0
    return rank, -size, source


def _scan_schedule_cells(workbook: Any) -> dict[str, ScheduleCell]:
    cells: dict[str, ScheduleCell] = {}
    for ws in workbook.worksheets:
        _scan_simple_class_sheet(ws, cells)
        _scan_wide_class_sheet(ws, cells)
    return cells


def _scan_simple_class_sheet(ws: Any, cells: dict[str, ScheduleCell]) -> None:
    header = _find_day_header(ws)
    if not header:
        return
    header_row, day_cols = header
    first_header = _cell_text(ws.cell(header_row, 1).value)
    if first_header == "班级":
        return
    class_name = _class_name_from_sheet(ws.title)
    if not class_name:
        return
    max_row = int(ws.max_row or 0)
    for row_idx in range(header_row + 1, max_row + 1):
        slot = _normalize_slot(_cell_text(ws.cell(row_idx, 1).value))
        if not slot:
            continue
        for day, col in day_cols:
            _add_cell_location(
                cells,
                class_name=class_name,
                day=day,
                slot=slot,
                value=_cell_text(ws.cell(row_idx, col).value),
                location=CellLocation(ws.title, row_idx, col, "class_sheet"),
            )


def _scan_wide_class_sheet(ws: Any, cells: dict[str, ScheduleCell]) -> None:
    max_col = int(ws.max_column or 0)
    if max_col < 2:
        return
    headers: list[tuple[int, str, str]] = []
    for col in range(2, max_col + 1):
        parsed = _parse_day_slot_header(_cell_text(ws.cell(1, col).value))
        if parsed:
            day, slot = parsed
            headers.append((col, day, slot))
    if len(headers) < 2:
        return
    for row_idx in range(2, int(ws.max_row or 0) + 1):
        class_name = _class_name_from_sheet(_cell_text(ws.cell(row_idx, 1).value))
        if not class_name:
            continue
        for col, day, slot in headers:
            _add_cell_location(
                cells,
                class_name=class_name,
                day=day,
                slot=slot,
                value=_cell_text(ws.cell(row_idx, col).value),
                location=CellLocation(ws.title, row_idx, col, "wide_sheet"),
            )


def _add_cell_location(
    cells: dict[str, ScheduleCell],
    *,
    class_name: str,
    day: str,
    slot: str,
    value: str,
    location: CellLocation,
) -> None:
    key = _cell_key(class_name, day, slot)
    current = cells.get(key)
    if current is None:
        cells[key] = ScheduleCell(key, class_name, day, slot, value, [location])
        return
    current.locations.append(location)
    if not current.value and value:
        current.value = value
    if location.role == "class_sheet" and value:
        current.value = value


def _find_day_header(ws: Any) -> tuple[int, list[tuple[str, int]]] | None:
    max_row = min(int(ws.max_row or 0), 8)
    max_col = min(int(ws.max_column or 0), 24)
    for row_idx in range(1, max_row + 1):
        cols: list[tuple[str, int]] = []
        for col in range(1, max_col + 1):
            day = _normalize_day(_cell_text(ws.cell(row_idx, col).value))
            if day in DAYS:
                cols.append((day, col))
        if len(cols) >= 3:
            return row_idx, cols
    return None


def _parse_day_slot_header(text: str) -> tuple[str, str] | None:
    value = _cell_text(text).replace("周", "星期").replace("-", "_")
    match = re.match(r"^(星期[一二三四五六日])(?:[_\s.]+(.+))?$", value)
    if not match:
        return None
    day = match.group(1)
    slot = _normalize_slot(match.group(2) or "晚自习1")
    return day, slot


def _resolve_leave_request(academic_payload: dict[str, Any], request: dict[str, Any]) -> dict[str, Any]:
    data = academic_payload.get("data") if isinstance(academic_payload.get("data"), dict) else academic_payload
    tables = data.get("tables") if isinstance(data, dict) else {}
    leave_rows = tables.get("teacher_leave") if isinstance(tables, dict) and isinstance(tables.get("teacher_leave"), list) else []
    row: dict[str, Any] = {}
    if isinstance(request.get("leave"), dict):
        row = copy.deepcopy(request["leave"])
    else:
        raw_index = request.get("leave_index")
        if raw_index is not None and str(raw_index).strip() != "":
            idx = int(float(raw_index))
            if idx < 0 or idx >= len(leave_rows):
                raise LocalTimetableRepairError("请假日志行不存在，请刷新教务工作台后重试。")
            row = copy.deepcopy(leave_rows[idx]) if isinstance(leave_rows[idx], dict) else {}
    for key in ("教师", "星期", "时段", "状态", "原因", "开始日期", "结束日期"):
        if str(request.get(key) or "").strip():
            row[key] = request.get(key)
    teacher = _cell_text(request.get("teacher") or row.get("教师"))
    day = _normalize_day(_cell_text(request.get("day") or row.get("星期")))
    slot = _normalize_slot(_cell_text(request.get("slot") or row.get("时段") or "全天"))
    if not teacher:
        raise LocalTimetableRepairError("请先选择或填写请假教师。")
    if not day:
        raise LocalTimetableRepairError("请假日志需要包含星期，才能定位课表。")
    return {
        "teacher": teacher,
        "day": day,
        "slot": slot or "全天",
        "status": _cell_text(row.get("状态")),
        "reason": _cell_text(row.get("原因")),
        "row": row,
        "class_name": _class_name_from_sheet(_cell_text(request.get("class_name") or request.get("班级"))),
        "subject": _cell_text(request.get("subject") or request.get("学科")),
        "substitute_teacher": _cell_text(request.get("substitute_teacher") or request.get("代课教师")),
    }


def _plan_leave_repair(
    cells: dict[str, ScheduleCell],
    leave: dict[str, Any],
    teacher_rows: list[dict[str, Any]],
    request: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, str]]:
    before_values = {key: cell.value for key, cell in cells.items()}
    working = dict(before_values)
    targets = _target_leave_cells(cells, working, leave)
    if not targets:
        return [], [], working
    assignments = _teacher_assignments(teacher_rows)
    plans: list[dict[str, Any]] = []
    candidate_matrix: list[dict[str, Any]] = []
    used_targets: set[str] = set()
    for target in targets[:MAX_TARGET_LESSONS]:
        if target.key in used_targets:
            continue
        candidates = _candidate_substitutes(target, leave, assignments)
        options = _candidate_options_for_target(target, leave, candidates, cells, working)
        feasible = [option["_plan"] for option in options if option.get("_plan")]
        best = _best_plan_from_options(feasible)
        public_options: list[dict[str, Any]] = []
        for idx, option in enumerate(_rank_candidate_options(options, best), start=1):
            item = {key: value for key, value in option.items() if key != "_plan"}
            item["rank"] = idx
            item["recommended"] = bool(best and item.get("plan_id") == best.get("id"))
            public_options.append(item)
        candidate_matrix.append(
            {
                "target": _public_target(target, working.get(target.key, target.value)),
                "options": public_options,
            }
        )
        if not best:
            plans.append(_unavailable_plan(target, leave, candidates))
            continue
        plans.append(best)
        used_targets.add(target.key)
        for change in best["changes"]:
            working[change["key"]] = change["after"]
    return plans, candidate_matrix, working


def _target_leave_cells(cells: dict[str, ScheduleCell], working: dict[str, str], leave: dict[str, Any]) -> list[ScheduleCell]:
    teacher = str(leave.get("teacher") or "")
    day = str(leave.get("day") or "")
    target_class = str(leave.get("class_name") or "")
    target_slot = str(leave.get("slot") or "")
    out: list[ScheduleCell] = []
    for cell in cells.values():
        if day and cell.day != day:
            continue
        if target_class and cell.class_name != target_class:
            continue
        if not _slot_in_leave_window(cell.slot, target_slot):
            continue
        if teacher not in _extract_teachers(working.get(cell.key, cell.value)):
            continue
        out.append(cell)
    return sorted(out, key=lambda cell: (_slot_rank(cell.slot), cell.class_name))


def _candidate_substitutes(
    target: ScheduleCell,
    leave: dict[str, Any],
    assignments: list[dict[str, str]],
) -> list[dict[str, Any]]:
    requested = str(leave.get("substitute_teacher") or "").strip()
    target_subject = str(leave.get("subject") or "").strip() or _extract_subject(target.value)
    target_grade = _grade_key(target.class_name)
    if requested:
        return [{"teacher": requested, "subject": target_subject, "grade_match": None, "source": "教务指定代课教师", "penalty": 0}]

    candidates: list[dict[str, Any]] = []
    for item in assignments:
        teacher = item.get("teacher") or ""
        if not teacher or teacher == leave.get("teacher"):
            continue
        if target_subject and item.get("subject") != target_subject:
            continue
        same_grade = bool(target_grade and item.get("grade") == target_grade)
        candidates.append(
            {
                "teacher": teacher,
                "subject": item.get("subject") or target_subject,
                "class": item.get("class") or "",
                "grade_match": same_grade,
                "source": "同年级同学科" if same_grade else "非同年级同学科",
                "penalty": 0 if same_grade else 20,
            }
        )
    if not candidates and target_subject:
        for item in assignments:
            teacher = item.get("teacher") or ""
            if teacher and teacher != leave.get("teacher"):
                candidates.append(
                    {
                        "teacher": teacher,
                        "subject": item.get("subject") or "",
                        "class": item.get("class") or "",
                        "grade_match": False,
                        "source": "学科不完全匹配候选",
                        "penalty": 80,
                    }
                )
    return _dedupe_candidates(candidates)[:MAX_CANDIDATES]


def _candidate_options_for_target(
    target: ScheduleCell,
    leave: dict[str, Any],
    candidates: list[dict[str, Any]],
    cells: dict[str, ScheduleCell],
    working: dict[str, str],
) -> list[dict[str, Any]]:
    options: list[dict[str, Any]] = []
    for candidate in candidates:
        plan = _plan_candidate_for_target(target, leave, candidate, cells, working)
        if plan:
            impact = plan.get("impact") or {}
            options.append(
                {
                    "teacher": candidate.get("teacher") or "",
                    "source": candidate.get("source") or "",
                    "subject": candidate.get("subject") or "",
                    "class_name": candidate.get("class") or "",
                    "grade_match": candidate.get("grade_match"),
                    "status": "feasible",
                    "status_label": "可执行",
                    "plan_id": plan.get("id") or "",
                    "impact": {
                        "score": int(impact.get("score") or 0),
                        "involved_teacher_count": int(impact.get("involved_teacher_count") or 0),
                        "changed_cell_count": int(impact.get("changed_cell_count") or 0),
                        "moved_lesson_count": int(impact.get("moved_lesson_count") or 0),
                    },
                    "explanation": plan.get("explanation") or "",
                    "blockers": [],
                    "_plan": plan,
                }
            )
        else:
            options.append(_blocked_candidate_option(target, leave, candidate, cells, working))
    return options


def _best_plan_from_options(plans: list[dict[str, Any]]) -> dict[str, Any] | None:
    if not plans:
        return None
    return sorted(
        plans,
        key=lambda item: (
            item["impact"]["involved_teacher_count"],
            item["impact"]["changed_cell_count"],
            item["impact"]["score"],
        ),
    )[0]


def _rank_candidate_options(options: list[dict[str, Any]], best: dict[str, Any] | None) -> list[dict[str, Any]]:
    best_id = str(best.get("id") or "") if best else ""
    return sorted(
        options,
        key=lambda item: (
            0 if best_id and item.get("plan_id") == best_id else 1,
            0 if item.get("status") == "feasible" else 1,
            int((item.get("impact") or {}).get("involved_teacher_count") or 99),
            int((item.get("impact") or {}).get("changed_cell_count") or 99),
            int((item.get("impact") or {}).get("score") or 9999),
            str(item.get("teacher") or ""),
        ),
    )


def _blocked_candidate_option(
    target: ScheduleCell,
    leave: dict[str, Any],
    candidate: dict[str, Any],
    cells: dict[str, ScheduleCell],
    working: dict[str, str],
) -> dict[str, Any]:
    original_teacher = str(leave.get("teacher") or "")
    substitute = str(candidate.get("teacher") or "").strip()
    blockers: list[str] = []
    conflicts: list[ScheduleCell] = []
    if not substitute:
        blockers.append("候选教师为空。")
    elif substitute == original_teacher:
        blockers.append("代课教师不能是请假教师本人。")
    else:
        target_value = working.get(target.key, target.value)
        target_after = _replace_teacher(target_value, original_teacher, substitute, _extract_subject(target_value) or candidate.get("subject"))
        if not target_after or target_after == target_value:
            blockers.append("无法把原课格替换为该候选教师。")
        conflicts = _teacher_cells_at(substitute, target.day, target.slot, cells, working, ignore_keys={target.key})
        for conflict in conflicts:
            destination = _find_move_destination(conflict, substitute, cells, working, ignore_keys={target.key, conflict.key})
            if not destination:
                blockers.append(
                    f"{conflict.class_name} {conflict.day} {conflict.slot} 的 {working.get(conflict.key, conflict.value)} 无可移动空课/自习课格。"
                )
    if not blockers:
        blockers.append("该候选暂时无法形成局部最小影响方案。")
    involved = sorted({name for name in (original_teacher, substitute) if name})
    return {
        "teacher": substitute,
        "source": candidate.get("source") or "",
        "subject": candidate.get("subject") or "",
        "class_name": candidate.get("class") or "",
        "grade_match": candidate.get("grade_match"),
        "status": "blocked",
        "status_label": "不可局部处理",
        "plan_id": "",
        "impact": {
            "score": 9999,
            "involved_teacher_count": len(involved),
            "changed_cell_count": 0,
            "moved_lesson_count": len(conflicts),
        },
        "explanation": "；".join(blockers[:3]),
        "blockers": blockers[:5],
    }


def _plan_candidate_for_target(
    target: ScheduleCell,
    leave: dict[str, Any],
    candidate: dict[str, Any],
    cells: dict[str, ScheduleCell],
    working: dict[str, str],
) -> dict[str, Any] | None:
    original_teacher = str(leave.get("teacher") or "")
    substitute = str(candidate.get("teacher") or "").strip()
    if not substitute or substitute == original_teacher:
        return None
    target_value = working.get(target.key, target.value)
    target_after = _replace_teacher(target_value, original_teacher, substitute, _extract_subject(target_value) or candidate.get("subject"))
    if not target_after or target_after == target_value:
        return None

    changes: list[dict[str, Any]] = []
    moved_lessons: list[dict[str, Any]] = []
    temp = dict(working)
    conflicts = _teacher_cells_at(substitute, target.day, target.slot, cells, temp, ignore_keys={target.key})
    score = int(candidate.get("penalty") or 0)
    for conflict in conflicts:
        destination = _find_move_destination(conflict, substitute, cells, temp, ignore_keys={target.key, conflict.key})
        if not destination:
            return None
        conflict_value = temp.get(conflict.key, conflict.value)
        dest_before = temp.get(destination["cell"].key, destination["cell"].value)
        changes.append(_change(conflict, conflict_value, "", "移动代课教师原有冲突课"))
        changes.append(_change(destination["cell"], dest_before, conflict_value, "承接代课教师原有冲突课"))
        temp[conflict.key] = ""
        temp[destination["cell"].key] = conflict_value
        score += int(destination["score"])
        moved_lessons.append(
            {
                "class_name": conflict.class_name,
                "subject": _extract_subject(conflict_value),
                "from": {"day": conflict.day, "slot": conflict.slot},
                "to": {"day": destination["cell"].day, "slot": destination["cell"].slot},
                "reason": destination["reason"],
            }
        )
    changes.append(_change(target, target_value, target_after, "代请假教师上课"))
    temp[target.key] = target_after
    involved = sorted({original_teacher, substitute})
    score += len(involved) * 100 + len(changes) * 10
    return {
        "id": f"repair_{_safe_id(target.class_name)}_{_safe_id(target.day)}_{_safe_id(target.slot)}_{_safe_id(substitute)}",
        "status": "feasible",
        "title": f"{target.class_name} {target.day} {target.slot} 请假代课局部最小影响调课",
        "target": {
            "class_name": target.class_name,
            "day": target.day,
            "slot": target.slot,
            "before": target_value,
            "after": target_after,
        },
        "substitute": {
            "teacher": substitute,
            "source": candidate.get("source") or "",
            "grade_match": candidate.get("grade_match"),
        },
        "moved_lessons": moved_lessons,
        "changes": changes,
        "impact": {
            "score": score,
            "involved_teachers": involved,
            "involved_teacher_count": len(involved),
            "changed_cell_count": len(changes),
            "moved_lesson_count": len(moved_lessons),
        },
        "explanation": _plan_explanation(substitute, moved_lessons),
    }


def _find_move_destination(
    source: ScheduleCell,
    teacher: str,
    cells: dict[str, ScheduleCell],
    working: dict[str, str],
    *,
    ignore_keys: set[str],
) -> dict[str, Any] | None:
    candidates: list[dict[str, Any]] = []
    for cell in cells.values():
        if cell.class_name != source.class_name or cell.key == source.key or cell.key in ignore_keys:
            continue
        value = working.get(cell.key, cell.value)
        dest_type = _destination_type(value)
        if not dest_type:
            continue
        if _teacher_busy_at(teacher, cell.day, cell.slot, cells, working, ignore_keys={source.key, cell.key}):
            continue
        same_day_penalty = 0 if cell.day == source.day else 12
        distance = abs(_slot_rank(cell.slot) - _slot_rank(source.slot))
        candidates.append(
            {
                "cell": cell,
                "score": same_day_penalty + distance + dest_type["penalty"],
                "reason": dest_type["reason"] if cell.day == source.day else f"{dest_type['reason']}，跨日移动",
            }
        )
    return sorted(candidates, key=lambda item: (item["score"], DAYS.index(item["cell"].day) if item["cell"].day in DAYS else 99, _slot_rank(item["cell"].slot)))[0] if candidates else None


def _destination_type(value: str) -> dict[str, Any] | None:
    text = str(value or "").strip()
    if not text:
        return {"penalty": 0, "reason": "使用空课格"}
    if text in {"自习", "自主学习", "空", "空课"}:
        return {"penalty": 4, "reason": "使用自习课格"}
    return None


def _unavailable_plan(target: ScheduleCell, leave: dict[str, Any], candidates: list[dict[str, Any]]) -> dict[str, Any]:
    return {
        "id": f"repair_unavailable_{_safe_id(target.class_name)}_{_safe_id(target.day)}_{_safe_id(target.slot)}",
        "status": "unavailable",
        "title": f"{target.class_name} {target.day} {target.slot} 暂无局部调课方案",
        "target": {"class_name": target.class_name, "day": target.day, "slot": target.slot, "before": target.value, "after": target.value},
        "substitute": {"teacher": str(leave.get("substitute_teacher") or ""), "source": ""},
        "candidate_count": len(candidates),
        "changes": [],
        "impact": {"score": 9999, "involved_teachers": [str(leave.get("teacher") or "")], "involved_teacher_count": 1, "changed_cell_count": 0},
        "explanation": "未找到可用代课教师，或代课教师同节次冲突课无法移动到空课/自习课格。",
    }


def _change(cell: ScheduleCell, before: str, after: str, reason: str) -> dict[str, Any]:
    return {
        "key": cell.key,
        "class_name": cell.class_name,
        "day": cell.day,
        "slot": cell.slot,
        "before": before,
        "after": after,
        "reason": reason,
        "locations": [{"sheet": loc.sheet, "row": loc.row, "col": loc.col, "role": loc.role} for loc in cell.locations],
    }


def _public_target(target: ScheduleCell, value: str) -> dict[str, str]:
    return {
        "class_name": target.class_name,
        "day": target.day,
        "slot": target.slot,
        "before": value,
        "subject": _extract_subject(value),
    }


def _repair_payload(
    workbook_path: Path,
    leave: dict[str, Any],
    cells: dict[str, ScheduleCell],
    plans: list[dict[str, Any]],
    candidate_options: list[dict[str, Any]],
    after_values: dict[str, str],
) -> dict[str, Any]:
    feasible = [plan for plan in plans if plan.get("status") == "feasible"]
    before_values = {key: cell.value for key, cell in cells.items()}
    involved = sorted({teacher for plan in feasible for teacher in plan.get("impact", {}).get("involved_teachers", []) if teacher})
    changed_cells = sum(int(plan.get("impact", {}).get("changed_cell_count") or 0) for plan in feasible)
    message = (
        f"已生成 {len(feasible)} 个局部调课方案，涉及 {len(involved)} 位教师、{changed_cells} 个课格。"
        if feasible
        else "未找到可执行的局部调课方案；请换代课教师或人工扩大调整范围。"
    )
    return {
        "schema_version": SCHEDULE_REPAIR_SCHEMA,
        "applied": False,
        "workbook": {"path": str(workbook_path), "label": workbook_path.name},
        "leave": leave,
        "summary": {
            "status": "ready" if feasible else "blocked",
            "message": message,
            "target_lesson_count": len(plans),
            "feasible_plan_count": len(feasible),
            "involved_teacher_count": len(involved),
            "changed_cell_count": changed_cells,
            "involved_teachers": involved,
            "candidate_option_count": sum(len(group.get("options") or []) for group in candidate_options),
            "feasible_candidate_option_count": sum(
                1
                for group in candidate_options
                for option in group.get("options") or []
                if option.get("status") == "feasible"
            ),
        },
        "candidate_options": candidate_options,
        "plans": plans,
        "affected_teacher_schedules": _affected_teacher_schedules(involved, cells, before_values, after_values),
    }


def _write_adjusted_workbook(
    workbook: Any,
    source_path: Path,
    cells: dict[str, ScheduleCell],
    plans: list[dict[str, Any]],
    output_dir: Path | None,
    suffix: str = "局部调课",
) -> Path:
    out_dir = (output_dir or source_path.parent / LOCAL_REPAIR_DIR_NAME).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = re.sub(r"[\\/:*?\"<>|]+", "_", source_path.stem)
    safe_suffix = re.sub(r"[\\/:*?\"<>|]+", "_", suffix or "局部调课")
    out_path = out_dir / f"{stem}_{safe_suffix}_{stamp}.xlsx"
    for plan in plans:
        if plan.get("status") != "feasible":
            continue
        for change in plan.get("changes", []) or []:
            key = str(change.get("key") or "")
            after = str(change.get("after") or "")
            cell = cells.get(key)
            if cell is None:
                continue
            for loc in cell.locations:
                workbook[loc.sheet].cell(loc.row, loc.col).value = after or None
    workbook.save(out_path)
    return out_path


def _repair_record(payload: dict[str, Any], output_path: Path) -> dict[str, Any]:
    return {
        "schema_version": SCHEDULE_REPAIR_SCHEMA,
        "applied_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "active_file": str(output_path),
        "summary": copy.deepcopy(payload.get("summary") or {}),
        "leave": copy.deepcopy(payload.get("leave") or {}),
        "candidate_options": copy.deepcopy(payload.get("candidate_options") or []),
        "plans": copy.deepcopy(payload.get("plans") or []),
        "affected_teacher_schedules": copy.deepcopy(payload.get("affected_teacher_schedules") or []),
    }


def _affected_teacher_schedules(
    teachers: list[str],
    cells: dict[str, ScheduleCell],
    before_values: dict[str, str],
    after_values: dict[str, str],
) -> list[dict[str, Any]]:
    return [
        {
            "teacher": teacher,
            "before": _teacher_schedule_rows(teacher, cells, before_values),
            "after": _teacher_schedule_rows(teacher, cells, after_values),
            "changed_slots": _teacher_changed_slots(teacher, cells, before_values, after_values),
        }
        for teacher in teachers
    ]


def _teacher_schedule_rows(teacher: str, cells: dict[str, ScheduleCell], values: dict[str, str]) -> list[dict[str, Any]]:
    by_slot: dict[str, dict[str, list[str]]] = defaultdict(lambda: {day: [] for day in DAYS})
    for cell in cells.values():
        value = values.get(cell.key, cell.value)
        if teacher not in _extract_teachers(value):
            continue
        label = f"{cell.class_name} {_extract_subject(value)}".strip()
        by_slot[cell.slot][cell.day].append(label)
    rows = []
    for slot in sorted(by_slot, key=_slot_rank):
        rows.append({"slot": slot, "values": {day: "；".join(by_slot[slot].get(day, [])) for day in DAYS}})
    return rows


def _teacher_changed_slots(
    teacher: str,
    cells: dict[str, ScheduleCell],
    before_values: dict[str, str],
    after_values: dict[str, str],
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for day in DAYS:
        slots = sorted({cell.slot for cell in cells.values() if cell.day == day}, key=_slot_rank)
        for slot in slots:
            before = _teacher_slot_label(teacher, day, slot, cells, before_values)
            after = _teacher_slot_label(teacher, day, slot, cells, after_values)
            if before != after:
                out.append({"day": day, "slot": slot, "before": before, "after": after})
    return out


def _teacher_slot_label(teacher: str, day: str, slot: str, cells: dict[str, ScheduleCell], values: dict[str, str]) -> str:
    labels = []
    for cell in cells.values():
        if cell.day != day or cell.slot != slot:
            continue
        value = values.get(cell.key, cell.value)
        if teacher in _extract_teachers(value):
            labels.append(f"{cell.class_name} {_extract_subject(value)}".strip())
    return "；".join(labels)


def _teacher_cells_at(
    teacher: str,
    day: str,
    slot: str,
    cells: dict[str, ScheduleCell],
    values: dict[str, str],
    *,
    ignore_keys: set[str] | None = None,
) -> list[ScheduleCell]:
    ignored = ignore_keys or set()
    return [
        cell
        for cell in cells.values()
        if cell.key not in ignored
        and cell.day == day
        and cell.slot == slot
        and teacher in _extract_teachers(values.get(cell.key, cell.value))
    ]


def _teacher_busy_at(
    teacher: str,
    day: str,
    slot: str,
    cells: dict[str, ScheduleCell],
    values: dict[str, str],
    *,
    ignore_keys: set[str] | None = None,
) -> bool:
    return bool(_teacher_cells_at(teacher, day, slot, cells, values, ignore_keys=ignore_keys))


def _teacher_assignments(rows: list[dict[str, Any]]) -> list[dict[str, str]]:
    assignments: list[dict[str, str]] = []
    excluded = {"row_index", "班级", "班主任", "班主任性别"}
    for row in rows:
        if not isinstance(row, dict):
            continue
        class_name = _class_name_from_sheet(row.get("班级"))
        grade = _grade_key(class_name)
        for key, value in row.items():
            subject = str(key)
            if subject in excluded:
                continue
            teacher = _cell_text(value)
            if teacher:
                assignments.append({"teacher": teacher, "subject": subject, "class": class_name, "grade": grade})
    return assignments


def _dedupe_candidates(candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[str] = set()
    out: list[dict[str, Any]] = []
    for item in sorted(candidates, key=lambda row: (int(row.get("penalty") or 0), str(row.get("teacher") or ""))):
        teacher = str(item.get("teacher") or "")
        if not teacher or teacher in seen:
            continue
        seen.add(teacher)
        out.append(item)
    return out


def _plan_explanation(substitute: str, moved_lessons: list[dict[str, Any]]) -> str:
    if not moved_lessons:
        return f"{substitute} 在请假节次没有原课，直接代课，只改 1 个课格。"
    moved = "；".join(
        f"{item['class_name']} {item['subject']} 从 {item['from']['day']} {item['from']['slot']} 移到 {item['to']['day']} {item['to']['slot']}"
        for item in moved_lessons[:4]
    )
    return f"{substitute} 原课与请假课重合，系统仅移动该教师自己的冲突课：{moved}。"


def _slot_in_leave_window(slot: str, leave_slot: str) -> bool:
    text = _normalize_slot(leave_slot)
    if not text or _is_all_day(text):
        return True
    return _normalize_slot(slot) == text


def _is_all_day(slot: str) -> bool:
    return str(slot or "").strip() in {"全天", "全日", "上午", "下午", "白天"}


def _replace_teacher(value: str, old_teacher: str, new_teacher: str, subject: Any = "") -> str:
    text = str(value or "").strip()
    subject_text = str(subject or "").strip() or _extract_subject(text)
    if old_teacher and old_teacher in _extract_teachers(text):
        pattern = rf"([（(])\s*{re.escape(old_teacher)}\s*([）)])"
        replaced = re.sub(pattern, rf"\1{new_teacher}\2", text)
        if replaced != text:
            return replaced
    return f"{subject_text}({new_teacher})" if subject_text else new_teacher


def _extract_teachers(text: str) -> list[str]:
    value = str(text or "").strip()
    teachers = [match.strip() for match in re.findall(r"[（(]([^()（）]+)[）)]", value) if match.strip()]
    if teachers:
        return list(dict.fromkeys(teachers))
    if "-" in value:
        tail = value.rsplit("-", 1)[-1].strip()
        if tail and len(tail) <= 8:
            return [tail]
    return []


def _extract_subject(text: str) -> str:
    value = str(text or "").strip()
    value = re.sub(r"[（(][^()（）]+[）)]", "", value).strip()
    if "-" in value:
        value = value.split("-", 1)[0].strip()
    return value


def _normalize_day(value: Any) -> str:
    text = _cell_text(value).replace("周", "星期")
    mapping = {
        "星期天": "星期日",
        "周天": "星期日",
        "周日": "星期日",
        "周一": "星期一",
        "周二": "星期二",
        "周三": "星期三",
        "周四": "星期四",
        "周五": "星期五",
        "周六": "星期六",
    }
    return mapping.get(text, text)


def _normalize_slot(value: Any) -> str:
    text = _cell_text(value)
    replacements = {
        "早自习": "早自习1",
        "早读": "早自习1",
        "第一节": "上午1",
        "第二节": "上午2",
        "第三节": "上午3",
        "第四节": "上午4",
        "第五节": "下午1",
        "第六节": "下午2",
        "第七节": "下午3",
        "第八节": "下午4",
    }
    return replacements.get(text, text)


def _slot_rank(slot: str) -> int:
    order = {
        "早自习1": 1,
        "早自习": 1,
        "上午1": 10,
        "上午2": 20,
        "上午3": 30,
        "上午4": 40,
        "下午1": 50,
        "下午2": 60,
        "下午3": 70,
        "下午4": 80,
        "晚自习1": 90,
        "晚自习2": 100,
    }
    return order.get(str(slot or ""), 500)


def _grade_key(class_name: str) -> str:
    text = str(class_name or "").strip()
    for pattern in (r"(初[一二三123])", r"(高[一二三123])", r"([七八九]年级)", r"([一二三四五六]年级)", r"(\d+年级)"):
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    prefix = re.match(r"^(\d+)", text)
    return prefix.group(1) if prefix else ""


def _class_name_from_sheet(value: Any) -> str:
    text = _cell_text(value)
    text = re.sub(r"\s*课表\s*$", "", text)
    return text


def _cell_key(class_name: str, day: str, slot: str) -> str:
    return f"{class_name}|{day}|{slot}"


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ").strip()
    return re.sub(r"\s+", " ", text)


def _safe_id(value: str) -> str:
    text = re.sub(r"\W+", "_", str(value or ""), flags=re.UNICODE).strip("_")
    return text or "item"
