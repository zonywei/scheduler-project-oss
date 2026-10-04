# -*- coding: utf-8 -*-
from __future__ import annotations

import json
import os
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook

ROOT_DIR = Path(__file__).resolve().parents[2]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from scheduler.output_paths import project_output_dir


def _assert_equal(actual, expected, label: str) -> None:
    if actual != expected:
        raise AssertionError(f"{label} mismatch: {actual!r} != {expected!r}")


def main() -> None:
    out_dir = project_output_dir() / "meta" / "encoding_regression"
    out_dir.mkdir(parents=True, exist_ok=True)

    payload = {
        "title": "诊断报告",
        "teacher": "指定教师",
        "message": "周日晚自习禁排",
    }

    txt_path = out_dir / "diagnostic_report.txt"
    txt_path.write_text("诊断报告：中文写入校验", encoding="utf-8")
    _assert_equal(txt_path.read_text(encoding="utf-8"), "诊断报告：中文写入校验", "txt")

    json_path = out_dir / "diagnostic_report.json"
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    _assert_equal(json.loads(json_path.read_text(encoding="utf-8")), payload, "json")

    yaml_path = out_dir / "diagnostic_report.yaml"
    yaml_path.write_text(yaml.safe_dump(payload, allow_unicode=True, sort_keys=False), encoding="utf-8")
    _assert_equal(yaml.safe_load(yaml_path.read_text(encoding="utf-8")), payload, "yaml")

    xlsx_path = out_dir / "中文编码校验.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "中文诊断"
    ws["A1"] = "晚自习排课"
    ws["B1"] = "周日晚自习禁排"
    wb.save(xlsx_path)

    wb2 = load_workbook(xlsx_path)
    ws2 = wb2["中文诊断"]
    _assert_equal(ws2["A1"].value, "晚自习排课", "excel:A1")
    _assert_equal(ws2["B1"].value, "周日晚自习禁排", "excel:B1")

    stdout_encoding = (sys.stdout.encoding or "").lower()
    print(f"[encoding-regression] stdout_encoding={sys.stdout.encoding}")
    print(f"[encoding-regression] written_dir={out_dir}")
    if os.name == "nt" and "utf" not in stdout_encoding:
        print("[encoding-regression] 建议先执行: chcp 65001")
        print("[encoding-regression] 建议再执行: set PYTHONUTF8=1")
    print("[encoding-regression] PASS")

if __name__ == "__main__":
    raise SystemExit(main())

