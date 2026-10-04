from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook


REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))


from scheduler.app.result_preview import build_result_preview
from scheduler.app.rule_explain import build_rule_explain


def _status_for(path: Path, label: str = "outputs/课表及值班安排.xlsx") -> dict:
    return {
        "status": "completed",
        "files": [
            {
                "label": label,
                "path": str(path),
                "size": path.stat().st_size,
                "modified_at": "2026-05-06 08:00:00",
            }
        ],
    }


def _file_item(path: Path, label: str) -> dict:
    return {
        "label": label,
        "path": str(path),
        "size": path.stat().st_size,
        "modified_at": "2026-05-06 08:00:00",
    }


def test_result_preview_extracts_class_grid_and_derived_teacher_view(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "初一1班"
    ws.append(["", "星期一", "星期二", "星期三", "星期四", "星期五"])
    ws.append(["上午1", "语文(教师A)", "数学(教师B)", "", "", ""])
    ws.append(["上午2", "", "英语(教师C)", "", "", ""])
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)

    assert preview["summary"]["status"] == "ready"
    assert preview["summary"]["class_views"] == 1
    assert preview["summary"]["teacher_views"] == 3
    class_view = preview["class_views"][0]
    assert class_view["name"] == "初一1班"
    assert class_view["rows"][0]["values"]["星期一"] == "语文(教师A)"
    teacher_names = {view["name"] for view in preview["teacher_views"]}
    assert {"教师A", "教师B", "教师C"} <= teacher_names
    assert preview["quality"]["summary"]["teacher_count"] == 3
    assert preview["quality"]["summary"]["empty_cells"] > 0
    assert any(issue["title"] == "班级课表存在空课格" for issue in preview["quality"]["issues"])
    stats = preview["statistics"]
    assert stats["schema_version"] == "scheduler.result_schedule_statistics.v1"
    assert stats["summary"]["status"] == "ready"
    assert stats["summary"]["classes"] == 1
    assert stats["summary"]["teachers"] == 3
    assert stats["summary"]["filled_cells"] == 3
    assert stats["summary"]["top_subject"].startswith("语文")
    assert stats["day_distribution"][0] == {"day": "星期一", "count": 1, "share": 33.3}
    assert {"subject": "数学", "count": 1, "share": 33.3} in stats["subject_distribution"]
    assert stats["class_day_load"]["max"] == 2
    assert stats["teacher_load"]["top"][0]["count"] == 1
    assert any(item["teacher"] == "教师A" and item["day"] == "星期一" and item["count"] == 1 for item in stats["teacher_day_load"]["top"])
    assert stats["teacher_consecutive_load"]["max_streak"] == 0
    assert any(item["title"] == "存在空课格需要确认" for item in stats["insights"])


def test_result_preview_explains_best_solution_rule_penalties(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "初一1班"
    ws.append(["", "星期一", "星期二", "星期三"])
    ws.append(["上午1", "语文(教师A)", "数学(教师B)", "英语(教师C)"])
    wb.save(workbook_path)

    stale_log = tmp_path / "pool_cache" / "sol_0010" / "event_log.csv"
    stale_log.parent.mkdir(parents=True)
    stale_log.write_text(
        "constraint_id,constraint_name,teacher_name,unit_penalty,count_value,penalty,mode\n"
        "stale_rule,旧候选解规则,旧教师,100,1,100,soft\n",
        encoding="utf-8",
    )
    best_log = tmp_path / "pool_cache" / "sol_0035" / "event_log.csv"
    best_log.parent.mkdir(parents=True)
    best_log.write_text(
        "constraint_id,constraint_name,teacher_name,day,period,class_name,subject,unit_penalty,count_value,penalty,mode,description\n"
        "lang_tue_fri_pm3_penalty,语文外语周二至周五下午3惩罚,教师A,星期二,下午3,1班,语文,10,2,20,soft,语文下午偏好未满足\n"
        "lang_tue_fri_pm3_penalty,语文外语周二至周五下午3惩罚,教师B,星期三,下午3,2班,英语,10,1,10,soft,外语下午偏好未满足\n"
        "day_night_link_w1,晚自习需当天下午有课,教师C,星期五,,3班,,5,1,5,soft,当日有晚自习但当天下午无课\n"
        "lang_tue_fri_am1_reward,语文外语周二至周五上午1奖励,教师A,星期二,上午1,1班,语文,-2,1,-2,soft,奖励\n",
        encoding="utf-8",
    )
    best_meta = tmp_path / "best" / "最终全局最优解_meta.json"
    best_meta.parent.mkdir(parents=True)
    best_meta.write_text(
        json.dumps(
            {
                "seq_id": 35,
                "event_log_file": str(best_log),
                "source_solution_file": str(best_log.with_name("最终全局最优解.xlsx")),
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    status = _status_for(workbook_path)
    status["objective_value"] = 33
    status["files"].extend(
        [
            _file_item(stale_log, "outputs/web_runs/run_x/solutions/run_x/pool_cache/sol_0010/event_log.csv"),
            _file_item(best_log, "outputs/web_runs/run_x/solutions/run_x/pool_cache/sol_0035/event_log.csv"),
            _file_item(best_meta, "outputs/web_runs/run_x/solutions/run_x/best/最终全局最优解_meta.json"),
        ]
    )

    preview = build_result_preview(status, project_root=tmp_path)

    explain = preview["statistics"]["rule_explain"]
    assert explain["schema_version"] == "scheduler.result_rule_explain.v1"
    assert explain["summary"]["status"] == "warning"
    assert explain["summary"]["total_penalty"] == 35
    assert explain["summary"]["reward_credit"] == 2
    assert explain["summary"]["net_event_penalty"] == 33
    assert explain["summary"]["solver_objective_value"] == 33
    assert explain["summary"]["objective_explain_delta"] == 0
    assert explain["summary"]["objective_reconciliation_status"] == "matched"
    assert explain["summary"]["rule_count"] == 2
    assert explain["summary"]["source"] == str(best_log.resolve())
    assert explain["top"][0]["rule_id"] == "lang_tue_fri_pm3_penalty"
    assert explain["top"][0]["penalty_sum"] == 30
    assert explain["top"][0]["share"] == 85.7
    assert explain["top"][0]["top_offenders"][0] == {"entity": "教师A", "penalty_sum": 20.0}
    assert explain["top"][0]["potential_penalty_reduction"] == 30
    assert explain["top"][0]["example_events"][0] == {
        "teacher": "教师A",
        "day": "星期二",
        "period": "下午3",
        "class_name": "1班",
        "subject": "语文",
        "penalty": 20.0,
        "unit_penalty": 10.0,
        "count": 2.0,
        "description": "语文下午偏好未满足",
    }
    assert "学科" in explain["top"][0]["action"]
    assert any(item["rule_id"] == "day_night_link_w1" and "白天-晚自习" in item["action"] for item in explain["top"])
    assert "旧候选解规则" not in json.dumps(explain, ensure_ascii=False)


def test_rule_explain_deduplicates_legacy_multi_class_rollup(tmp_path: Path) -> None:
    event_log = tmp_path / "event_log.csv"
    event_log.write_text(
        "constraint_id,constraint_name,teacher_name,unit_penalty,count_value,penalty,mode\n"
        "multi_class_halfday,多班教师半天偏好,教师A,300,2,600,soft\n"
        "multi_class_halfday_split,多班教师半天拆分,教师A,300,2,600,soft\n"
        "day_night_link_w1,晚自习需当天下午有课,教师C,100,1,100,soft\n",
        encoding="utf-8",
    )

    explain = build_rule_explain(
        {
            "objective_value": 700,
            "files": [_file_item(event_log, "outputs/web_runs/run_x/event_log.csv")],
        },
        project_root=tmp_path,
    )

    assert explain["summary"]["total_penalty"] == 700
    assert explain["summary"]["net_event_penalty"] == 700
    assert explain["summary"]["objective_reconciliation_status"] == "matched"
    rule_ids = {item["rule_id"] for item in explain["top"]}
    assert "multi_class_halfday_split" in rule_ids
    assert "multi_class_halfday" not in rule_ids


def test_rule_explain_ignores_legacy_nonobjective_head_duty_trigger(tmp_path: Path) -> None:
    event_log = tmp_path / "event_log.csv"
    event_log.write_text(
        "constraint_id,constraint_name,teacher_name,unit_penalty,count_value,penalty,mode\n"
        "head_duty_pm1_trigger,Head PM1 standalone penalty,教师A,1000,1,1000,soft\n"
        "day_night_link_w1,晚自习需当天下午有课,教师C,100,1,100,soft\n",
        encoding="utf-8",
    )

    explain = build_rule_explain(
        {
            "objective_value": 100,
            "files": [_file_item(event_log, "outputs/web_runs/run_x/event_log.csv")],
        },
        project_root=tmp_path,
    )

    assert explain["summary"]["total_penalty"] == 100
    assert explain["summary"]["objective_reconciliation_status"] == "matched"
    assert [item["rule_id"] for item in explain["top"]] == ["day_night_link_w1"]


def test_rule_explain_counts_current_head_duty_pm1_penalty(tmp_path: Path) -> None:
    event_log = tmp_path / "event_log.csv"
    event_log.write_text(
        "constraint_id,constraint_name,teacher_name,day,period,unit_penalty,count_value,penalty,mode\n"
        "head_duty_pm1_without_duty,班主任PM1无同日值班惩罚,教师A,星期六,下午1,1000,1,1000,soft\n"
        "day_night_link_w1,晚自习需当天下午有课,教师C,星期一,,100,1,100,soft\n",
        encoding="utf-8",
    )

    explain = build_rule_explain(
        {
            "objective_value": 1100,
            "files": [_file_item(event_log, "outputs/web_runs/run_x/event_log.csv")],
        },
        project_root=tmp_path,
    )

    assert explain["summary"]["total_penalty"] == 1100
    assert explain["summary"]["objective_reconciliation_status"] == "matched"
    assert {item["rule_id"] for item in explain["top"]} == {
        "head_duty_pm1_without_duty",
        "day_night_link_w1",
    }


def test_result_preview_quality_detects_teacher_same_slot_conflict(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "初一1班"
    ws1.append(["", "星期一", "星期二", "星期三"])
    ws1.append(["上午1", "语文(教师A)", "", ""])
    ws2 = wb.create_sheet("初一2班")
    ws2.append(["", "星期一", "星期二", "星期三"])
    ws2.append(["上午1", "数学(教师A)", "", ""])
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)

    assert preview["quality"]["summary"]["status"] == "blocked"
    assert preview["quality"]["summary"]["teacher_conflicts"] == 1
    assert any(issue["title"] == "教师同节次出现在多个班级" for issue in preview["quality"]["issues"])


def test_result_preview_ignores_globally_empty_teaching_slots(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "初一1班"
    ws1.append(["", "星期一", "星期二", "星期三"])
    ws1.append(["上午1", "语文(教师A)", "", "英语(教师C)"])
    ws2 = wb.create_sheet("初一2班")
    ws2.append(["", "星期一", "星期二", "星期三"])
    ws2.append(["上午1", "数学(教师B)", "", "历史(教师D)"])
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)
    summary = preview["quality"]["summary"]
    titles = {issue["title"] for issue in preview["quality"]["issues"]}

    assert summary["empty_cells"] == 0
    assert summary["inactive_slot_count"] == 1
    assert summary["inactive_cells_ignored"] == 2
    assert "已忽略全校非开课空白时段" in titles
    assert "班级课表存在空课格" not in titles


def test_result_preview_still_warns_when_one_class_has_empty_slot(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "初一1班"
    ws1.append(["", "星期一", "星期二", "星期三"])
    ws1.append(["上午1", "", "", "英语(教师C)"])
    ws2 = wb.create_sheet("初一2班")
    ws2.append(["", "星期一", "星期二", "星期三"])
    ws2.append(["上午1", "数学(教师B)", "", "历史(教师D)"])
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)
    summary = preview["quality"]["summary"]
    titles = {issue["title"] for issue in preview["quality"]["issues"]}

    assert summary["empty_cells"] == 1
    assert summary["inactive_cells_ignored"] == 2
    assert "班级课表存在空课格" in titles
    assert "已忽略全校非开课空白时段" in titles


def test_result_preview_treats_load_spread_as_explanatory_review_hint(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "初一1班"
    ws.append(["", "星期一", "星期二", "星期三"])
    ws.append(["上午1", "语文(教师A)", "语文(教师A)", "语文(教师A)"])
    ws.append(["上午2", "语文(教师A)", "语文(教师A)", "数学(教师B)"])
    ws.append(["下午1", "语文(教师A)", "语文(教师A)", "英语(教师C)"])
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)
    summary = preview["quality"]["summary"]
    load_issue = next(
        issue for issue in preview["quality"]["issues"]
        if issue["title"] == "教师课时负荷需按任课量复核"
    )

    assert summary["status"] == "ok"
    assert summary["warnings"] == 0
    assert summary["teacher_load_spread"] == 6
    assert summary["teacher_load_median"] == 1
    assert load_issue["severity"] == "info"
    assert "不等同于学校正式工作量口径" in load_issue["detail"]
    assert preview["quality"]["teacher_load_low"][0] == {"teacher": "教师B", "count": 1}
    assert preview["statistics"]["teacher_day_load"]["max"] == 3
    assert preview["statistics"]["teacher_consecutive_load"]["max_streak"] == 0
    assert any(item["title"] == "教师负荷存在可复查差异" for item in preview["statistics"]["insights"])


def test_result_preview_flags_teacher_day_load_concentration(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "初一1班"
    ws.append(["", "星期一", "星期二", "星期三"])
    ws.append(["上午1", "语文(教师A)", "数学(教师B)", "英语(教师C)"])
    ws.append(["上午2", "语文(教师A)", "数学(教师B)", "英语(教师C)"])
    ws.append(["下午1", "语文(教师A)", "数学(教师B)", "英语(教师C)"])
    ws.append(["下午2", "语文(教师A)", "数学(教师B)", "英语(教师C)"])
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)
    teacher_day = preview["statistics"]["teacher_day_load"]

    assert teacher_day["max"] == 4
    assert any(
        item["teacher"] == "教师A" and item["day"] == "星期一" and item["count"] == 4
        for item in teacher_day["top"]
    )
    assert preview["statistics"]["teacher_consecutive_load"]["max_streak"] == 0
    assert any(item["title"] == "教师单日课量偏高" for item in preview["statistics"]["insights"])


def test_result_preview_suggests_teacher_day_load_balance(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "初一1班"
    ws.append(["", "星期一", "星期二", "星期三"])
    ws.append(["早自习1", "", "", ""])
    ws.append(["上午1", "语文(教师A)", "", ""])
    ws.append(["上午3", "数学(教师A)", "", ""])
    ws.append(["下午1", "英语(教师A)", "", ""])
    ws.append(["下午3", "物理(教师A)", "", ""])
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)
    teacher_day = preview["statistics"]["teacher_day_load"]

    assert teacher_day["max"] == 4
    suggestion = teacher_day["top"][0]["suggestion"]
    assert suggestion["kind"] == "teacher_day_balance"
    assert suggestion["teacher"] == "教师A"
    assert suggestion["class_name"] == "初一1班"
    assert suggestion["from_day"] == "星期一"
    assert suggestion["from_slot"] == "下午3"
    assert suggestion["to_day"] == "星期二"
    assert suggestion["to_slot"] == "下午3"
    assert suggestion["original_subject"] == "物理"
    assert suggestion["original_teacher"] == "教师A"
    assert suggestion["target_kind"] == "empty"
    assert suggestion["reason"] == "移动到教师低负荷日同班常规空课格"
    queue = preview["statistics"]["adjustment_suggestions"]
    assert queue["count"] == 1
    assert queue["top"][0]["source"] == "teacher_day_load"
    assert queue["top"][0]["source_label"] == "教师日负荷均衡"
    assert queue["top"][0]["suggestion"]["kind"] == "teacher_day_balance"


def test_result_preview_flags_teacher_consecutive_streak(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "初一1班"
    ws.append(["", "星期一", "星期二", "星期三"])
    ws.append(["上午1", "语文(教师A)", "数学(教师B)", "英语(教师C)"])
    ws.append(["上午2", "语文(教师A)", "", ""])
    ws.append(["上午3", "语文(教师A)", "", ""])
    ws.append(["下午1", "数学(教师B)", "", ""])
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)
    consecutive = preview["statistics"]["teacher_consecutive_load"]

    assert consecutive["max_streak"] == 3
    assert consecutive["top"][0]["teacher"] == "教师A"
    assert consecutive["top"][0]["day"] == "星期一"
    assert consecutive["top"][0]["slot_text"] == "上午1 至 上午3"
    assert consecutive["top"][0]["classes"] == ["初一1班"]
    suggestion = consecutive["top"][0]["suggestion"]
    assert suggestion["kind"] == "consecutive_break"
    assert suggestion["class_name"] == "初一1班"
    assert suggestion["from_day"] == "星期一"
    assert suggestion["from_slot"] == "上午2"
    assert suggestion["to_day"] == "星期二"
    assert suggestion["to_slot"] == "上午2"
    assert suggestion["original_subject"] == "语文"
    assert suggestion["original_teacher"] == "教师A"
    queue = preview["statistics"]["adjustment_suggestions"]
    assert queue["count"] >= 1
    assert queue["top"][0]["source"] == "teacher_consecutive_load"
    assert queue["top"][0]["suggestion"]["kind"] == "consecutive_break"
    assert any(item["title"] == "教师连续课风险" for item in preview["statistics"]["insights"])


def test_result_preview_keeps_consecutive_break_for_regular_lessons_in_regular_slots(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "初一1班"
    ws.append(["", "星期一", "星期二", "星期三"])
    ws.append(["早自习1", "", "", ""])
    ws.append(["上午1", "语文(教师A)", "数学(教师B)", "英语(教师C)"])
    ws.append(["上午2", "语文(教师A)", "", ""])
    ws.append(["上午3", "语文(教师A)", "", ""])
    ws.append(["下午1", "", "", ""])
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)
    suggestion = preview["statistics"]["teacher_consecutive_load"]["top"][0]["suggestion"]

    assert suggestion["kind"] == "consecutive_break"
    assert suggestion["from_slot"] == "上午2"
    assert suggestion["to_slot"] == "下午1"
    assert suggestion["target_kind"] == "empty"
    assert suggestion["reason"] == "移动到同班常规空课格"


def test_result_preview_flags_class_day_load_and_suggests_balance(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "初一1班"
    ws.append(["", "星期一", "星期二", "星期三"])
    ws.append(["早自习1", "", "", ""])
    ws.append(["上午1", "语文(教师A)", "自习", ""])
    ws.append(["上午2", "数学(教师B)", "", ""])
    ws.append(["上午3", "英语(教师C)", "", ""])
    ws.append(["下午1", "历史(教师D)", "", ""])
    ws.append(["下午2", "物理(教师H)", "", ""])
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)
    class_day = preview["statistics"]["class_day_load"]

    assert class_day["max"] == 5
    assert class_day["top"][0]["class_name"] == "初一1班"
    assert class_day["top"][0]["day"] == "星期一"
    assert class_day["top"][0]["count"] == 5
    suggestion = class_day["top"][0]["suggestion"]
    assert suggestion["kind"] == "class_day_balance"
    assert suggestion["class_name"] == "初一1班"
    assert suggestion["from_day"] == "星期一"
    assert suggestion["from_slot"] == "下午2"
    assert suggestion["to_day"] == "星期二"
    assert suggestion["to_slot"] == "下午2"
    assert suggestion["original_subject"] == "物理"
    assert suggestion["original_teacher"] == "教师H"
    assert suggestion["target_kind"] == "empty"
    assert suggestion["reason"] == "移动到低负荷日常规空课格"
    queue = preview["statistics"]["adjustment_suggestions"]
    assert queue["count"] == 1
    assert queue["top"][0]["source"] == "class_day_load"
    assert queue["top"][0]["suggestion"]["kind"] == "class_day_balance"
    assert queue["top"][0]["metric_value"] == "5 节"
    assert any(item["title"] == "班级单日课量偏高" for item in preview["statistics"]["insights"])


def test_result_preview_extracts_teacher_block_sheet(tmp_path: Path) -> None:
    workbook_path = tmp_path / "课表及值班安排.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "老师课表"
    ws["B1"] = "教师A课表"
    ws["D3"] = "周一"
    ws["F3"] = "周二"
    ws["H3"] = "周三"
    ws["J3"] = "周四"
    ws["L3"] = "周五"
    ws["B5"] = "上午1"
    ws["D5"] = "初一1班 语文"
    ws["F5"] = "初一2班 语文"
    wb.save(workbook_path)

    preview = build_result_preview(_status_for(workbook_path), project_root=tmp_path)

    assert preview["summary"]["status"] == "ready"
    assert preview["summary"]["teacher_views"] == 1
    teacher_view = preview["teacher_views"][0]
    assert teacher_view["name"] == "教师A"
    assert teacher_view["rows"][0]["values"]["星期一"] == "初一1班 语文"


def test_result_preview_returns_empty_without_schedule_workbook(tmp_path: Path) -> None:
    diagnostic = tmp_path / "多解诊断报告.xlsx"
    wb = Workbook()
    wb.save(diagnostic)

    preview = build_result_preview(_status_for(diagnostic, label="outputs/多解诊断报告.xlsx"), project_root=tmp_path)

    assert preview["summary"]["status"] == "empty"
    assert preview["summary"]["schedule_files"] == 0
    assert preview["class_views"] == []
    assert preview["statistics"]["summary"]["status"] == "empty"


def test_result_preview_prefers_formal_best_workbook_over_top50(tmp_path: Path) -> None:
    top_path = tmp_path / "top50_excels" / "白天课表_汇总版_top01.xlsx"
    best_path = tmp_path / "best" / "最终全局最优解_正式版.xlsx"
    top_path.parent.mkdir()
    best_path.parent.mkdir()

    top_wb = Workbook()
    top_wb.active.title = "top候选"
    top_wb.active.append(["只是候选"])
    top_wb.save(top_path)

    best_wb = Workbook()
    ws = best_wb.active
    ws.title = "1班"
    ws.append(["", "星期一", "星期二", "星期三"])
    ws.append(["上午1", "语文(教师A)", "", ""])
    best_wb.save(best_path)

    preview = build_result_preview(
        {
            "status": "completed",
            "files": [
                _file_item(top_path, "outputs/top50_excels/白天课表_汇总版_top01.xlsx"),
                _file_item(best_path, "outputs/best/最终全局最优解_正式版.xlsx"),
            ],
        },
        project_root=tmp_path,
    )

    assert preview["summary"]["status"] == "ready"
    assert preview["summary"]["workbooks"] == 1
    assert preview["workbooks"][0]["label"] == "outputs/best/最终全局最优解_正式版.xlsx"
    assert preview["class_views"][0]["name"] == "1班"


def test_result_preview_uses_recommended_historical_candidate_when_latest_regresses(tmp_path: Path) -> None:
    current_path = tmp_path / "outputs" / "web_runs" / "run_current" / "best" / "最终全局最优解_正式版.xlsx"
    prior_path = tmp_path / "outputs" / "web_runs" / "run_prior" / "best" / "最终全局最优解_正式版.xlsx"
    current_path.parent.mkdir(parents=True)
    prior_path.parent.mkdir(parents=True)

    for path, value in ((current_path, "退化批次(教师A)"), (prior_path, "历史最好(教师B)")):
        wb = Workbook()
        ws = wb.active
        ws.title = "1班"
        ws.append(["", "星期一", "星期二", "星期三"])
        ws.append(["上午1", value, "", ""])
        wb.save(path)

    prior_status = {
        "run_id": "run_prior",
        "run_dir": str(prior_path.parents[1]),
        "status": "completed",
        "files": [_file_item(prior_path, "outputs/web_runs/run_prior/best/最终全局最优解_正式版.xlsx")],
    }
    prior_status_path = prior_path.parents[1] / "status.json"
    prior_status_path.write_text(json.dumps(prior_status, ensure_ascii=False), encoding="utf-8")
    current_status = {
        "run_id": "run_current",
        "run_dir": str(current_path.parents[1]),
        "status": "completed",
        "files": [_file_item(current_path, "outputs/web_runs/run_current/best/最终全局最优解_正式版.xlsx")],
        "recommended_formal_candidate": {
            "summary": {
                "status": "historical_best",
                "selected_run_id": "run_prior",
                "current_run_id": "run_current",
            },
            "candidate": {
                "run_id": "run_prior",
                "status_file": str(prior_status_path),
                "run_dir": str(prior_path.parents[1]),
            },
        },
    }

    preview = build_result_preview(current_status, project_root=tmp_path)

    assert preview["summary"]["source_kind"] == "recommended_formal_candidate"
    assert preview["summary"]["source_run_id"] == "run_prior"
    assert preview["summary"]["latest_run_id"] == "run_current"
    assert preview["workbooks"][0]["label"] == "outputs/web_runs/run_prior/best/最终全局最优解_正式版.xlsx"
    assert preview["class_views"][0]["rows"][0]["values"]["星期一"] == "历史最好(教师B)"


def test_result_preview_prefers_active_manual_adjustment_workbook(tmp_path: Path) -> None:
    original_path = tmp_path / "最终全局最优解_正式版.xlsx"
    adjusted_path = tmp_path / "最终全局最优解_正式版_课表微调.xlsx"
    for path, value in ((original_path, "语文(教师A)"), (adjusted_path, "数学(教师B)")):
        wb = Workbook()
        ws = wb.active
        ws.title = "1班"
        ws.append(["", "星期一", "星期二", "星期三"])
        ws.append(["上午1", value, "", ""])
        wb.save(path)

    preview = build_result_preview(
        {
            "status": "completed",
            "local_timetable_adjustment": {"active_file": str(adjusted_path)},
            "files": [
                _file_item(original_path, "outputs/best/最终全局最优解_正式版.xlsx"),
                _file_item(adjusted_path, "outputs/web_runs/run/local_adjustments/最终全局最优解_正式版_课表微调.xlsx"),
            ],
        },
        project_root=tmp_path,
    )

    assert preview["summary"]["status"] == "ready"
    assert preview["workbooks"][0]["label"] == "active_manual_adjustment"
    assert preview["class_views"][0]["rows"][0]["values"]["星期一"] == "数学(教师B)"


def test_result_preview_builds_schedule_version_timeline_from_local_adjustments(tmp_path: Path) -> None:
    original_path = tmp_path / "最终全局最优解_正式版.xlsx"
    adjusted_path = tmp_path / "最终全局最优解_正式版_课表微调.xlsx"
    for path, value in ((original_path, "语文(教师A)"), (adjusted_path, "数学(教师B)")):
        wb = Workbook()
        ws = wb.active
        ws.title = "1班"
        ws.append(["", "星期一", "星期二", "星期三"])
        ws.append(["上午1", value, "", ""])
        wb.save(path)

    preview = build_result_preview(
        {
            "status": "completed",
            "local_timetable_adjustment": {"active_file": str(adjusted_path)},
            "local_timetable_adjustments": [
                {
                    "active_file": str(adjusted_path),
                    "applied_at": "2026-05-07T08:20:00",
                    "summary": {
                        "message": "1班 星期一 上午1 已完成人工微调。",
                        "changed_cell_count": 2,
                        "involved_teachers": ["教师A", "教师B"],
                        "affected_classes": ["1班"],
                    },
                    "plan": {
                        "changes": [
                            {
                                "class_name": "1班",
                                "day": "星期一",
                                "slot": "上午1",
                                "before": "语文(教师A)",
                                "after": "数学(教师B)",
                                "reason": "人工微调",
                            }
                        ]
                    },
                }
            ],
            "files": [
                _file_item(original_path, "outputs/best/最终全局最优解_正式版.xlsx"),
            ],
        },
        project_root=tmp_path,
    )

    versions = preview["versions"]
    assert versions["schema_version"] == "scheduler.result_schedule_versions.v1"
    assert versions["summary"]["version_count"] == 2
    assert versions["summary"]["adjustment_count"] == 1
    assert versions["summary"]["changed_cell_count"] == 2
    assert versions["summary"]["active_kind"] == "课表微调"
    assert versions["items"][0]["kind"] == "base"
    adjustment = versions["items"][1]
    assert adjustment["active"] is True
    assert adjustment["summary"]["involved_teacher_count"] == 2
    assert adjustment["summary"]["affected_classes"] == ["1班"]
    assert adjustment["changes"][0]["before"] == "语文(教师A)"
    assert adjustment["changes"][0]["after"] == "数学(教师B)"
