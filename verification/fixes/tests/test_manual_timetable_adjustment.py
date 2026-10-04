from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook


REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))


from scheduler.app.local_timetable_adjustment import (
    apply_manual_timetable_adjustment,
    preview_manual_timetable_adjustment,
)


def _write_schedule(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "1班"
    ws.append([None, "星期一", "星期二", "星期三"])
    ws.append(["上午1", "语文(教师A)", "", ""])
    ws.append(["上午2", "数学(教师B)", "", ""])
    ws.append(["上午3", "", "", ""])

    ws2 = wb.create_sheet("2班")
    ws2.append([None, "星期一", "星期二", "星期三"])
    ws2.append(["上午1", "", "", ""])
    ws2.append(["上午2", "语文(教师C)", "", ""])
    ws2.append(["上午3", "", "", ""])

    wide = wb.create_sheet("白天课表_宽表")
    wide.append(["班级", "星期一_上午1", "星期一_上午2", "星期一_上午3"])
    wide.append(["1班", "语文(教师A)", "数学(教师B)", ""])
    wide.append(["2班", "", "语文(教师C)", ""])
    wb.save(path)


def test_manual_timetable_adjustment_swaps_occupied_slot_and_updates_global_schedule(tmp_path: Path) -> None:
    schedule = tmp_path / "最终全局最优解_正式版.xlsx"
    _write_schedule(schedule)
    status = {"files": [{"path": str(schedule), "label": "outputs/最终全局最优解_正式版.xlsx"}]}
    academic = {
        "data": {
            "tables": {
                "timetable_changes": [
                    {
                        "变更标题": "1班语文数学换课",
                        "班级": "1班",
                        "原星期": "星期一",
                        "原节次": "上午1",
                        "原学科": "语文",
                        "原教师": "教师A",
                        "新星期": "星期一",
                        "新节次": "上午2",
                        "新教师": "教师A",
                        "状态": "已通过",
                    }
                ]
            }
        }
    }

    preview = preview_manual_timetable_adjustment(
        status=status,
        academic_payload=academic,
        request={"change_index": 0},
        project_root=tmp_path,
    )

    assert preview["schema_version"] == "scheduler.manual_timetable_adjustment.v1"
    assert preview["summary"]["status"] == "ready"
    assert preview["summary"]["changed_cell_count"] == 2
    assert preview["summary"]["involved_teachers"] == ["教师A", "教师B"]
    assert "互换" in preview["plan"]["explanation"]
    class_schedule = preview["affected_class_schedules"][0]
    assert class_schedule["changed_slots"] == [
        {"day": "星期一", "slot": "上午1", "before": "语文(教师A)", "after": "数学(教师B)"},
        {"day": "星期一", "slot": "上午2", "before": "数学(教师B)", "after": "语文(教师A)"},
    ]

    applied = apply_manual_timetable_adjustment(
        status=status,
        academic_payload=academic,
        request={"change_index": 0},
        output_dir=tmp_path / "local_adjustments",
        project_root=tmp_path,
    )

    output = Path(applied["output_file"])
    assert output.exists()
    assert "课表微调" in output.name
    wb = load_workbook(output, data_only=True)
    try:
        assert wb["1班"].cell(2, 2).value == "数学(教师B)"
        assert wb["1班"].cell(3, 2).value == "语文(教师A)"
        assert wb["白天课表_宽表"].cell(2, 2).value == "数学(教师B)"
        assert wb["白天课表_宽表"].cell(2, 3).value == "语文(教师A)"
    finally:
        wb.close()


def test_manual_timetable_adjustment_blocks_teacher_conflict_and_suggests_slots(tmp_path: Path) -> None:
    schedule = tmp_path / "最终全局最优解_正式版.xlsx"
    _write_schedule(schedule)
    status = {"files": [{"path": str(schedule), "label": "outputs/最终全局最优解_正式版.xlsx"}]}

    preview = preview_manual_timetable_adjustment(
        status=status,
        academic_payload={"data": {"tables": {"timetable_changes": []}}},
        request={
            "class_name": "1班",
            "from_day": "星期一",
            "from_slot": "上午1",
            "to_day": "星期一",
            "to_slot": "上午2",
            "new_teacher": "教师C",
        },
        project_root=tmp_path,
    )

    assert preview["summary"]["status"] == "blocked"
    assert "新教师在目标节次已有课或任务" in preview["summary"]["message"]
    assert {"class_name": "1班", "day": "星期一", "slot": "上午3", "reason": "使用空课格", "score": 20} in preview["plan"]["alternatives"]
