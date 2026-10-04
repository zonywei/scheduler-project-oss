from __future__ import annotations

import ast
from pathlib import Path

from openpyxl import load_workbook

from scheduler.output.excel_writer import DEFAULT_XLSX_ENGINE
from scheduler.output.night_exporter import export_result_xlsx


REPO_ROOT = Path(__file__).resolve().parents[3]


class _FakeSolver:
    def Value(self, var: object) -> int:
        return int(var)


def _parse_python(path: Path) -> ast.AST:
    return ast.parse(path.read_text(encoding="utf-8-sig"), filename=str(path))


def test_project_default_excel_engine_is_openpyxl() -> None:
    assert DEFAULT_XLSX_ENGINE == "openpyxl"


def test_scheduler_code_does_not_request_xlsxwriter_engine() -> None:
    offenders: list[str] = []
    for path in (REPO_ROOT / "scheduler").rglob("*.py"):
        tree = _parse_python(path)
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            for kw in node.keywords:
                if kw.arg == "engine" and isinstance(kw.value, ast.Constant) and kw.value.value == "xlsxwriter":
                    offenders.append(str(path.relative_to(REPO_ROOT)))
    assert not offenders, f"scheduler runtime must not use xlsxwriter engine: {offenders}"


def test_pandas_excelwriter_is_wrapped_by_project_policy() -> None:
    offenders: list[str] = []
    for path in (REPO_ROOT / "scheduler").rglob("*.py"):
        if path == REPO_ROOT / "scheduler" / "output" / "excel_writer.py":
            continue
        tree = _parse_python(path)
        for node in ast.walk(tree):
            if not isinstance(node, ast.Call):
                continue
            func = node.func
            if (
                isinstance(func, ast.Attribute)
                and func.attr == "ExcelWriter"
                and isinstance(func.value, ast.Name)
                and func.value.id == "pd"
            ):
                offenders.append(str(path.relative_to(REPO_ROOT)))
    assert not offenders, f"use create_excel_writer instead of pd.ExcelWriter: {offenders}"


def test_night_exporter_writes_openpyxl_readable_workbook(tmp_path: Path) -> None:
    out_path = tmp_path / "night.xlsx"
    days = ["星期一"]
    periods = [1, 2]
    classes = ["1班"]
    cst = {("1班", "语文"): "教师A"}
    vars_ = {
        "y": {
            ("1班", "语文", "星期一", 1): 1,
            ("1班", "语文", "星期一", 2): 0,
        },
        "checkin_m": {("男班主任", "星期一"): 1},
        "checkin_f": {("女班主任", "星期一"): 1},
    }

    export_result_xlsx(
        out_path,
        _FakeSolver(),
        vars_,
        classes,
        cst,
        days,
        periods,
        ["男班主任"],
        ["女班主任"],
    )

    workbook = load_workbook(out_path, read_only=True)
    assert workbook.sheetnames == ["排课", "晚查寝安排"]
    schedule = workbook["排课"]
    assert schedule["A1"].value == "班级"
    assert schedule["B2"].value == "语文-教师A"
    checkin = workbook["晚查寝安排"]
    assert checkin["B2"].value == "男班主任"
    assert checkin["C2"].value == "女班主任"
