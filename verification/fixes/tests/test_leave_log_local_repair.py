from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook


REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))


from scheduler.app.local_timetable_repair import (
    apply_leave_substitution_repair,
    preview_leave_substitution_repair,
)


def _write_schedule(path: Path) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "1班"
    ws1.append([None, "星期一", "星期二", "星期三"])
    ws1.append(["上午1", "语文(教师A)", "数学(教师D)", ""])
    ws1.append(["上午2", "", "", ""])
    ws1.append(["下午1", "", "", ""])

    ws2 = wb.create_sheet("2班")
    ws2.append([None, "星期一", "星期二", "星期三"])
    ws2.append(["上午1", "语文(教师B)", "", ""])
    ws2.append(["上午2", "", "", ""])
    ws2.append(["下午1", "", "", ""])

    wide = wb.create_sheet("白天课表_宽表")
    wide.append(["班级", "星期一_上午1", "星期一_上午2", "星期二_上午1"])
    wide.append(["1班", "语文(教师A)", "", "数学(教师D)"])
    wide.append(["2班", "语文(教师B)", "", ""])
    wb.save(path)


def test_leave_log_local_repair_moves_substitute_conflict_with_minimal_teacher_impact(tmp_path: Path) -> None:
    schedule = tmp_path / "最终全局最优解_正式版.xlsx"
    _write_schedule(schedule)
    status = {"files": [{"path": str(schedule), "label": "outputs/最终全局最优解_正式版.xlsx"}]}
    academic = {
        "data": {
            "tables": {
                "teacher_leave": [
                    {"教师": "教师A", "星期": "星期一", "时段": "上午1", "状态": "已审批"},
                ]
            }
        }
    }
    teacher_rows = [
        {"班级": "1班", "语文": "教师A", "班主任": "", "班主任性别": ""},
        {"班级": "2班", "语文": "教师B", "班主任": "", "班主任性别": ""},
    ]

    preview = preview_leave_substitution_repair(
        status=status,
        academic_payload=academic,
        teacher_rows=teacher_rows,
        request={"leave_index": 0, "substitute_teacher": "教师B"},
        project_root=tmp_path,
    )

    assert preview["summary"]["status"] == "ready"
    assert preview["summary"]["involved_teachers"] == ["教师A", "教师B"]
    assert preview["summary"]["changed_cell_count"] == 3
    assert preview["summary"]["candidate_option_count"] == 1
    assert preview["summary"]["feasible_candidate_option_count"] == 1
    option = preview["candidate_options"][0]["options"][0]
    assert option["teacher"] == "教师B"
    assert option["status"] == "feasible"
    assert option["recommended"] is True
    assert option["impact"]["moved_lesson_count"] == 1
    plan = preview["plans"][0]
    assert plan["impact"]["moved_lesson_count"] == 1
    assert "仅移动该教师自己的冲突课" in plan["explanation"]
    assert any(change["class_name"] == "1班" and change["after"] == "语文(教师B)" for change in plan["changes"])
    assert any(change["class_name"] == "2班" and change["slot"] == "上午2" and change["after"] == "语文(教师B)" for change in plan["changes"])
    schedules = {item["teacher"]: item for item in preview["affected_teacher_schedules"]}
    assert schedules["教师B"]["changed_slots"] == [
        {"day": "星期一", "slot": "上午1", "before": "2班 语文", "after": "1班 语文"},
        {"day": "星期一", "slot": "上午2", "before": "", "after": "2班 语文"},
    ]

    applied = apply_leave_substitution_repair(
        status=status,
        academic_payload=academic,
        teacher_rows=teacher_rows,
        request={"leave_index": 0, "substitute_teacher": "教师B"},
        output_dir=tmp_path / "local_repairs",
        project_root=tmp_path,
    )

    output = Path(applied["output_file"])
    assert output.exists()
    wb = load_workbook(output, data_only=True)
    try:
        assert wb["1班"].cell(2, 2).value == "语文(教师B)"
        assert wb["2班"].cell(2, 2).value is None
        assert wb["2班"].cell(3, 2).value == "语文(教师B)"
        assert wb["白天课表_宽表"].cell(2, 2).value == "语文(教师B)"
        assert wb["白天课表_宽表"].cell(3, 2).value is None
        assert wb["白天课表_宽表"].cell(3, 3).value == "语文(教师B)"
    finally:
        wb.close()


def test_leave_log_local_repair_exposes_candidate_impact_matrix_with_blockers(tmp_path: Path) -> None:
    schedule = tmp_path / "最终全局最优解_正式版.xlsx"
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "初一1班"
    ws1.append([None, "星期一", "星期二", "星期三"])
    ws1.append(["上午1", "语文(教师A)", "", ""])
    ws1.append(["上午2", "", "", ""])

    ws2 = wb.create_sheet("初一2班")
    ws2.append([None, "星期一", "星期二", "星期三"])
    ws2.append(["上午1", "语文(教师B)", "", ""])
    ws2.append(["上午2", "", "", ""])

    ws3 = wb.create_sheet("初二1班")
    ws3.append([None, "星期一", "星期二", "星期三"])
    ws3.append(["上午1", "语文(教师C)", "历史(教师G)", "地理(钱老师)"])
    ws3.append(["上午2", "数学(教师D)", "物理(教师H)", "英语(吴老师)"])
    wb.save(schedule)

    status = {"files": [{"path": str(schedule), "label": "outputs/最终全局最优解_正式版.xlsx"}]}
    academic = {
        "data": {
            "tables": {
                "teacher_leave": [
                    {"教师": "教师A", "星期": "星期一", "时段": "上午1", "状态": "已确认"},
                ]
            }
        }
    }
    teacher_rows = [
        {"班级": "初一1班", "语文": "教师A", "班主任": "", "班主任性别": ""},
        {"班级": "初一2班", "语文": "教师B", "班主任": "", "班主任性别": ""},
        {"班级": "初二1班", "语文": "教师C", "班主任": "", "班主任性别": ""},
    ]

    preview = preview_leave_substitution_repair(
        status=status,
        academic_payload=academic,
        teacher_rows=teacher_rows,
        request={"leave_index": 0},
        project_root=tmp_path,
    )

    options = preview["candidate_options"][0]["options"]
    assert [item["teacher"] for item in options] == ["教师B", "教师C"]
    assert options[0]["recommended"] is True
    assert options[0]["source"] == "同年级同学科"
    assert options[0]["status"] == "feasible"
    assert options[1]["source"] == "非同年级同学科"
    assert options[1]["status"] == "blocked"
    assert "无可移动空课/自习课格" in options[1]["explanation"]
    assert preview["summary"]["candidate_option_count"] == 2
    assert preview["summary"]["feasible_candidate_option_count"] == 1
