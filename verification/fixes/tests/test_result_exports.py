from __future__ import annotations

from pathlib import Path

from io import BytesIO

from openpyxl import Workbook, load_workbook

from scheduler.app.result_exports import build_result_pdf, build_result_xlsx, resolve_result_excel, result_export_name


def _schedule_workbook(tmp_path: Path) -> Path:
    target = tmp_path / "正式课表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "总表"
    sheet.append(["班级", "星期一-上午1", "星期二-上午1", "星期三-上午1"])
    sheet.append(["高一1班", "数学(样例教师A)", "语文(样例教师B)", "英语(样例教师C)"])
    workbook.save(target)
    return target


def test_formal_publish_exports_excel_and_pdf(tmp_path: Path) -> None:
    workbook = _schedule_workbook(tmp_path)
    status = {
        "status": "completed",
        "files": [{
            "path": str(workbook),
            "name": workbook.name,
            "label": "正式课表",
            "category": "schedule",
        }],
    }

    assert resolve_result_excel(status, project_root=tmp_path) == workbook.resolve()
    pdf = build_result_pdf(status, project_root=tmp_path)
    assert pdf.startswith(b"%PDF")
    assert len(pdf) > 1_000
    assert result_export_name("pdf").endswith(".pdf")
    assert result_export_name("xlsx").endswith(".xlsx")


def test_formal_publish_builds_focused_class_and_teacher_workbooks(tmp_path: Path) -> None:
    workbook = _schedule_workbook(tmp_path)
    status = {
        "status": "completed",
        "files": [{"path": str(workbook), "name": workbook.name, "label": "正式课表", "category": "schedule"}],
    }

    class_bytes = build_result_xlsx(status, project_root=tmp_path, view="class")
    teacher_bytes = build_result_xlsx(status, project_root=tmp_path, view="teacher")
    assert class_bytes.startswith(b"PK")
    assert teacher_bytes.startswith(b"PK")
    assert load_workbook(BytesIO(class_bytes)).sheetnames
    assert load_workbook(BytesIO(teacher_bytes)).sheetnames
    assert result_export_name("xlsx", "class") == "课有序_班级课表.xlsx"
    assert result_export_name("pdf", "teacher") == "课有序_教师课表.pdf"
