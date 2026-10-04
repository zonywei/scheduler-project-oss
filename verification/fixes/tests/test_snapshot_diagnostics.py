from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[3]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))


from scheduler.snapshot_diagnostics import SHEET_SOLVER_OVERVIEW, build_multi_solution_diagnostic


def test_multi_solution_diagnostic_handles_global_only_event_log(tmp_path: Path) -> None:
    out_dir = tmp_path / "run"
    out_dir.mkdir()
    pd.DataFrame(
        [
            {
                "constraint_id": "night_global_soft_probe",
                "constraint_name": "全局软约束",
                "constraint_category": "preference",
                "mode": "soft",
                "teacher_name": "GLOBAL",
                "day": "星期一",
                "period": "晚1",
                "penalty": 0,
                "unit_penalty": 0,
                "count_value": 0,
            }
        ]
    ).to_csv(out_dir / "event_log.csv", index=False, encoding="utf-8-sig")

    report = build_multi_solution_diagnostic(out_dir)

    assert report.exists()
    workbook = load_workbook(report, read_only=True)
    try:
        assert SHEET_SOLVER_OVERVIEW in workbook.sheetnames
        assert workbook[SHEET_SOLVER_OVERVIEW].max_row >= 1
    finally:
        workbook.close()
